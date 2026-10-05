#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""
analyze_channel_variants_v3.py

Hurtig analyse af kanalvarianter i EPGoal.

Laeser kun:
  1) C:\EPGoal\output_epgshare\epgoal.xml
  2) C:\EPGoal\channel_priority_v2.xlsx

Opretter:
  C:\EPGoal\output_epgshare\channel_variant_analysis_v3.xlsx

Scriptet aendrer ikke epgoal.xml eller channel_priority_v2.xlsx.
"""
from __future__ import annotations

import argparse
import re
import sys
import time
import unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Optional

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    print("FEJL: openpyxl mangler. Koer: py -m pip install openpyxl", file=sys.stderr)
    raise

DEFAULT_XML = Path(r"C:\EPGoal\output_epgshare\epgoal.xml")
DEFAULT_MAPPING = Path(r"C:\EPGoal\channel_priority_v2.xlsx")
DEFAULT_OUTPUT = Path(r"C:\EPGoal\output_epgshare\channel_variant_analysis_v3.xlsx")

VARIANT_ORDER = {"NORMAL": 0, "HD": 1, "FHD": 2, "UHD": 3}
QUALITY_RE = re.compile(
    r"(?i)(?:^|[\s._+()\-\[\]])(4k|uhd|fhd|full\s*hd|hd|sd)(?=$|[\s._+()\-\[\]])"
)
COUNTRY_RE = re.compile(
    r"(?i)(?:^|[\s._+()\-\[\]])(dk|da|denmark|danmark)(?=$|[\s._+()\-\[\]])"
)


@dataclass(frozen=True)
class MappingRow:
    excel_row: int
    sheet: str
    channel_name: str
    stream_id: str
    output_id: str
    variant: str
    base_key: str
    source_values: str


def clean_text(value: object) -> str:
    return " ".join(str(value or "").strip().split())


def ascii_fold(value: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(ch)
    )


def normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", ascii_fold(clean_text(value)).lower())


def detect_variant(*values: str) -> str:
    text = " ".join(clean_text(v) for v in values).lower()
    if re.search(r"(?:^|[\s._+()\-\[\]])(?:4k|uhd)(?=$|[\s._+()\-\[\]])", text):
        return "UHD"
    if re.search(r"(?:^|[\s._+()\-\[\]])(?:fhd|full\s*hd)(?=$|[\s._+()\-\[\]])", text):
        return "FHD"
    if re.search(r"(?:^|[\s._+()\-\[\]])hd(?=$|[\s._+()\-\[\]])", text):
        return "HD"
    return "NORMAL"


def make_base_key(*values: str) -> str:
    text = next((clean_text(v) for v in values if clean_text(v)), "")
    text = ascii_fold(text).lower()
    text = QUALITY_RE.sub(" ", text)
    text = COUNTRY_RE.sub(" ", text)
    text = re.sub(
        r"\([^)]*(?:dk|da|denmark|danmark|sd|hd|fhd|uhd|4k)[^)]*\)",
        " ", text, flags=re.I,
    )
    return re.sub(r"[^a-z0-9]+", "", text)


def local_tag(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def human_number(value: int) -> str:
    return f"{value:,}".replace(",", ".")


def elapsed(started: float) -> str:
    seconds = int(time.monotonic() - started)
    minutes, seconds = divmod(seconds, 60)
    hours, minutes = divmod(minutes, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"
    return f"{minutes:02d}:{seconds:02d}"


def print_stage(text: str) -> None:
    print(text, flush=True)


def read_epgoal_xml(path: Path, progress_every: int = 250_000):
    """Laes kanaler og programmer i samme iterparse-gennemloeb."""
    channel_names: dict[str, str] = {}
    programme_counts: Counter[str] = Counter()
    programmes_seen = 0
    started = time.monotonic()

    size_mb = path.stat().st_size / (1024 * 1024)
    print_stage(f"[1/3] Laeser XML: {path}")
    print_stage(f"      Filstoerrelse: {size_mb:,.1f} MB")

    try:
        for _, elem in ET.iterparse(path, events=("end",)):
            tag = local_tag(elem.tag)

            if tag == "channel":
                channel_id = clean_text(elem.attrib.get("id"))
                if channel_id:
                    display_name = ""
                    for child in list(elem):
                        if local_tag(child.tag) == "display-name" and clean_text(child.text):
                            display_name = clean_text(child.text)
                            break
                    channel_names[channel_id] = display_name or channel_id

            elif tag == "programme":
                channel_id = clean_text(elem.attrib.get("channel"))
                if channel_id:
                    programme_counts[channel_id] += 1
                programmes_seen += 1
                if progress_every > 0 and programmes_seen % progress_every == 0:
                    print_stage(
                        f"      Programmer: {human_number(programmes_seen)} | "
                        f"kanal-ID'er: {human_number(len(programme_counts))} | "
                        f"tid: {elapsed(started)}"
                    )

            elem.clear()

    except ET.ParseError as exc:
        raise RuntimeError(f"XML-filen er ugyldig omkring {exc}") from exc

    all_ids = set(channel_names) | set(programme_counts)
    for channel_id in all_ids:
        channel_names.setdefault(channel_id, channel_id)

    print_stage(
        f"      Faerdig: {human_number(len(all_ids))} kanal-ID'er og "
        f"{human_number(sum(programme_counts.values()))} programmer ({elapsed(started)})"
    )
    return channel_names, programme_counts


def find_column(headers: list[str], exact: Iterable[str], contains: Iterable[str] = ()) -> Optional[int]:
    normalized = [normalize_header(h) for h in headers]
    exact_set = {normalize_header(x) for x in exact}
    contains_set = {normalize_header(x) for x in contains}

    for index, header in enumerate(normalized):
        if header in exact_set:
            return index
    for index, header in enumerate(normalized):
        if any(term and term in header for term in contains_set):
            return index
    return None


def read_mapping_xlsx(path: Path) -> tuple[list[MappingRow], list[str]]:
    print_stage(f"[2/3] Laeser mapping: {path}")
    wb = load_workbook(path, read_only=True, data_only=True)
    mappings: list[MappingRow] = []
    warnings: list[str] = []

    for ws in wb.worksheets:
        iterator = ws.iter_rows(values_only=True)
        try:
            first_row = next(iterator)
        except StopIteration:
            continue

        headers = [clean_text(v) for v in first_row]
        name_col = find_column(
            headers,
            exact=("Kanal", "Kanalnavn", "Channel name", "Stream name", "tvg-name", "Navn"),
            contains=("kanalnavn", "channelname", "streamname", "tvgname"),
        )
        stream_id_col = find_column(
            headers,
            exact=("UHF tvg-id", "Stream tvg-id", "M3U tvg-id", "tvg-id", "Stream ID"),
            contains=("uhftvgid", "streamtvgid", "m3utvgid"),
        )
        output_id_col = find_column(
            headers,
            exact=("Output/UHF tvg-id", "Output UHF tvg-id", "Output tvg-id", "EPG tvg-id"),
            contains=("outputuhftvgid", "outputtvgid", "epgtvgid"),
        )
        variant_col = find_column(
            headers,
            exact=("Variant", "Kvalitet", "Quality"),
            contains=("variant", "kvalitet", "quality"),
        )

        if name_col is None and stream_id_col is None and output_id_col is None:
            warnings.append(f"Arket '{ws.title}' blev sprunget over: ingen genkendelige mapping-kolonner.")
            continue

        sheet_count = 0
        for excel_row, row in enumerate(iterator, start=2):
            name = clean_text(row[name_col]) if name_col is not None and name_col < len(row) else ""
            stream_id = clean_text(row[stream_id_col]) if stream_id_col is not None and stream_id_col < len(row) else ""
            output_id = clean_text(row[output_id_col]) if output_id_col is not None and output_id_col < len(row) else ""
            variant_text = clean_text(row[variant_col]) if variant_col is not None and variant_col < len(row) else ""

            if not (name or stream_id or output_id):
                continue

            variant = detect_variant(variant_text, name, stream_id)
            key = make_base_key(name, stream_id, output_id)
            if not key:
                continue

            source_values = " | ".join(
                f"{headers[i] or 'Kolonne ' + str(i + 1)}={clean_text(value)}"
                for i, value in enumerate(row)
                if clean_text(value)
            )
            mappings.append(
                MappingRow(
                    excel_row=excel_row,
                    sheet=ws.title,
                    channel_name=name or stream_id or output_id,
                    stream_id=stream_id,
                    output_id=output_id,
                    variant=variant,
                    base_key=key,
                    source_values=source_values,
                )
            )
            sheet_count += 1

        print_stage(f"      Ark '{ws.title}': {human_number(sheet_count)} mapping-raekker")

    print_stage(f"      Faerdig: {human_number(len(mappings))} mapping-raekker")
    return mappings, warnings


def choose_suggested_id(rows: list[MappingRow], programme_counts: Counter[str]) -> str:
    candidates: list[str] = []
    for row in rows:
        for value in (row.output_id, row.stream_id):
            if value and value not in candidates:
                candidates.append(value)

    candidates.sort(
        key=lambda value: (
            -(1 if programme_counts.get(value, 0) > 0 else 0),
            -programme_counts.get(value, 0),
            len(value),
            value.lower(),
        )
    )
    return candidates[0] if candidates else ""


def analyze(mappings: list[MappingRow], channel_names: dict[str, str], programme_counts: Counter[str]):
    xml_ids_lower = {channel_id.lower(): channel_id for channel_id in channel_names}
    families: dict[str, list[MappingRow]] = defaultdict(list)
    for row in mappings:
        families[row.base_key].append(row)

    detail_rows: list[list[object]] = []
    matrix_rows: list[list[object]] = []
    problem_rows: list[list[object]] = []

    for base_key in sorted(families):
        rows = sorted(
            families[base_key],
            key=lambda r: (VARIANT_ORDER.get(r.variant, 9), r.channel_name.lower(), r.excel_row),
        )
        suggested_id = choose_suggested_id(rows, programme_counts)
        family_effective_ids: set[str] = set()
        variants_present: set[str] = set()
        variants_with_epg: set[str] = set()

        variant_cells: dict[str, dict[str, object]] = {
            variant: {"names": [], "stream_ids": [], "output_ids": [], "effective_ids": [], "programmes": 0}
            for variant in VARIANT_ORDER
        }

        for row in rows:
            requested_id = row.output_id or row.stream_id
            actual_id = xml_ids_lower.get(requested_id.lower(), "") if requested_id else ""
            programme_count = programme_counts.get(actual_id, 0) if actual_id else 0
            epg = "JA" if programme_count > 0 else "NEJ"

            if not requested_id:
                status = "MANGLER TVG-ID"
            elif not actual_id:
                status = "ID FINDES IKKE I XML"
            elif programme_count == 0:
                status = "ID UDEN PROGRAMMER"
            elif row.output_id and row.stream_id and row.output_id.lower() != row.stream_id.lower():
                status = "NORMALISERET"
            else:
                status = "OK"

            variants_present.add(row.variant)
            if epg == "JA":
                variants_with_epg.add(row.variant)
            if requested_id:
                family_effective_ids.add(requested_id.lower())

            cell = variant_cells[row.variant]
            cell["names"].append(row.channel_name)
            if row.stream_id:
                cell["stream_ids"].append(row.stream_id)
            if row.output_id:
                cell["output_ids"].append(row.output_id)
            if requested_id:
                cell["effective_ids"].append(requested_id)
            cell["programmes"] = max(int(cell["programmes"]), programme_count)

            detail = [
                base_key, row.variant, row.channel_name, row.stream_id, row.output_id,
                requested_id, actual_id, channel_names.get(actual_id, ""), programme_count,
                epg, status, suggested_id, row.sheet, row.excel_row, row.source_values,
            ]
            detail_rows.append(detail)
            if status != "OK" and status != "NORMALISERET":
                problem_rows.append(detail)

        missing_variants = [v for v in ("NORMAL", "HD", "FHD") if v not in variants_present]
        missing_epg_variants = [v for v in sorted(variants_present, key=lambda x: VARIANT_ORDER.get(x, 9)) if v not in variants_with_epg]
        inconsistent = len(family_effective_ids) > 1

        family_status_parts: list[str] = []
        if missing_variants:
            family_status_parts.append("MANGLER " + ", ".join(missing_variants))
        if missing_epg_variants:
            family_status_parts.append("UDEN EPG: " + ", ".join(missing_epg_variants))
        if inconsistent:
            family_status_parts.append("FORSKELLIGE TVG-ID'ER")
        family_status = "OK" if not family_status_parts else " | ".join(family_status_parts)

        matrix = [base_key]
        for variant in ("NORMAL", "HD", "FHD", "UHD"):
            cell = variant_cells[variant]
            matrix.extend([
                " | ".join(dict.fromkeys(cell["names"])),
                " | ".join(dict.fromkeys(cell["stream_ids"])),
                " | ".join(dict.fromkeys(cell["output_ids"])),
                " | ".join(dict.fromkeys(cell["effective_ids"])),
                cell["programmes"],
            ])
        matrix.extend([
            suggested_id,
            ", ".join(missing_variants),
            ", ".join(missing_epg_variants),
            "JA" if inconsistent else "NEJ",
            family_status,
        ])
        matrix_rows.append(matrix)

    coverage_rows = [
        [channel_id, channel_names.get(channel_id, channel_id), programme_counts.get(channel_id, 0), "JA" if programme_counts.get(channel_id, 0) > 0 else "NEJ"]
        for channel_id in sorted(channel_names, key=str.lower)
    ]
    return detail_rows, matrix_rows, problem_rows, coverage_rows


def autosize(ws, max_width: int = 70) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for column in range(1, ws.max_column + 1):
        longest = 0
        for row in range(1, min(ws.max_row, 3000) + 1):
            longest = max(longest, len(clean_text(ws.cell(row, column).value)))
        ws.column_dimensions[get_column_letter(column)].width = max(10, min(longest + 2, max_width))
    if ws.max_row:
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[object]]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    if ws.max_row > 1:
        table_name = re.sub(r"[^A-Za-z0-9_]", "", title)[:24] + "Tbl"
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False,
            showLastColumn=False, showRowStripes=True, showColumnStripes=False,
        )
        ws.add_table(table)
    autosize(ws)
    return ws


def color_status_rows(ws, status_column: int) -> None:
    green = PatternFill("solid", fgColor="C6EFCE")
    orange = PatternFill("solid", fgColor="FCE4D6")
    red = PatternFill("solid", fgColor="FFC7CE")
    for row in range(2, ws.max_row + 1):
        status = clean_text(ws.cell(row, status_column).value)
        if status == "OK":
            fill = green
        elif status == "NORMALISERET" or "FORSKELLIGE" in status:
            fill = orange
        else:
            fill = red
        for column in range(1, ws.max_column + 1):
            ws.cell(row, column).fill = fill


def write_report(
    output: Path,
    xml_path: Path,
    mapping_path: Path,
    detail_rows: list[list[object]],
    matrix_rows: list[list[object]],
    problem_rows: list[list[object]],
    coverage_rows: list[list[object]],
    warnings: list[str],
) -> None:
    print_stage(f"[3/3] Skriver rapport: {output}")
    wb = Workbook()
    summary = wb.active
    summary.title = "Oversigt"

    family_ok = sum(1 for row in matrix_rows if row[-1] == "OK")
    missing_hd = sum(1 for row in matrix_rows if "HD" in clean_text(row[-4]).split(", "))
    missing_fhd = sum(1 for row in matrix_rows if "FHD" in clean_text(row[-4]).split(", "))
    without_epg = sum(1 for row in detail_rows if row[9] == "NEJ")
    inconsistent = sum(1 for row in matrix_rows if row[-2] == "JA")

    summary_rows = [
        ("XML", str(xml_path)),
        ("Mapping", str(mapping_path)),
        ("Kanalfamilier", len(matrix_rows)),
        ("Mapping-raekker", len(detail_rows)),
        ("Kanalfamilier OK", family_ok),
        ("Familier der mangler HD", missing_hd),
        ("Familier der mangler FHD", missing_fhd),
        ("Mapping-raekker uden EPG", without_epg),
        ("Familier med forskellige tvg-id'er", inconsistent),
        ("Problemer", len(problem_rows)),
        ("Advarsler", len(warnings)),
    ]
    summary.append(["Noegle", "Vaerdi"])
    for row in summary_rows:
        summary.append(row)
    autosize(summary)

    detail_headers = [
        "Base key", "Variant", "Kanalnavn", "Stream tvg-id", "Output/UHF tvg-id",
        "Anvendt tvg-id", "XML channel-id", "XML display-name", "Programmer", "EPG",
        "Status", "Foreslaaet faelles tvg-id", "Excel-ark", "Excel-raekke", "Kildedata",
    ]
    details_ws = add_sheet(wb, "Alle_mappinger", detail_headers, detail_rows)
    problems_ws = add_sheet(wb, "Mangler_og_fejl", detail_headers, problem_rows)

    matrix_headers = ["Base key"]
    for variant in ("Normal", "HD", "FHD", "UHD"):
        matrix_headers.extend([
            f"{variant} kanalnavn", f"{variant} stream-id", f"{variant} output-id",
            f"{variant} anvendt id", f"{variant} programmer",
        ])
    matrix_headers.extend([
        "Foreslaaet faelles tvg-id", "Manglende varianter", "Varianter uden EPG",
        "Forskellige tvg-id'er", "Familie-status",
    ])
    matrix_ws = add_sheet(wb, "Variant_matrix", matrix_headers, matrix_rows)

    add_sheet(
        wb, "EPG_daekning",
        ["XML channel-id", "Display-name", "Programmer", "Har EPG"],
        coverage_rows,
    )
    add_sheet(wb, "Advarsler", ["Advarsel"], [[w] for w in warnings] or [["Ingen advarsler"]])

    color_status_rows(details_ws, 11)
    color_status_rows(problems_ws, 11)
    color_status_rows(matrix_ws, len(matrix_headers))

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Analyser Normal/HD/FHD/UHD mappings mod den faerdige epgoal.xml."
    )
    parser.add_argument("--xml", type=Path, default=DEFAULT_XML, help="Sti til epgoal.xml")
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING, help="Sti til channel_priority_v2.xlsx")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Sti til Excel-rapport")
    parser.add_argument(
        "--progress-every", type=int, default=250_000,
        help="Vis XML-status for hvert antal programme-elementer. Brug 0 for ingen mellemstatus.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    started = time.monotonic()

    print("=== analyze_channel_variants_v3 ===", flush=True)
    print(f"XML     : {args.xml}", flush=True)
    print(f"Mapping : {args.mapping}", flush=True)
    print(f"Rapport : {args.output}", flush=True)
    print(flush=True)

    if not args.xml.is_file():
        print(f"FEJL: XML-filen findes ikke: {args.xml}", file=sys.stderr)
        return 2
    if not args.mapping.is_file():
        print(f"FEJL: Mapping-filen findes ikke: {args.mapping}", file=sys.stderr)
        return 2
    if args.output.resolve() == args.mapping.resolve():
        print("FEJL: Output maa ikke overskrive mapping-filen.", file=sys.stderr)
        return 2

    try:
        channel_names, programme_counts = read_epgoal_xml(args.xml, args.progress_every)
        mappings, warnings = read_mapping_xlsx(args.mapping)
        if not mappings:
            print("FEJL: Ingen mapping-raekker blev fundet i Excel-filen.", file=sys.stderr)
            if warnings:
                for warning in warnings:
                    print(f"  - {warning}", file=sys.stderr)
            return 3

        detail_rows, matrix_rows, problem_rows, coverage_rows = analyze(
            mappings, channel_names, programme_counts
        )
        write_report(
            args.output, args.xml, args.mapping, detail_rows, matrix_rows,
            problem_rows, coverage_rows, warnings,
        )

    except PermissionError as exc:
        print(f"FEJL: Adgang naegtet. Luk eventuelt Excel-filen og proev igen: {exc}", file=sys.stderr)
        return 4
    except Exception as exc:
        print(f"FEJL: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    print(flush=True)
    print("=== FAERDIG ===", flush=True)
    print(f"XML-kanaler : {human_number(len(channel_names))}", flush=True)
    print(f"Programmer  : {human_number(sum(programme_counts.values()))}", flush=True)
    print(f"Mappinger   : {human_number(len(detail_rows))}", flush=True)
    print(f"Familier    : {human_number(len(matrix_rows))}", flush=True)
    print(f"Problemer   : {human_number(len(problem_rows))}", flush=True)
    print(f"Rapport     : {args.output}", flush=True)
    print(f"Samlet tid  : {elapsed(started)}", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
