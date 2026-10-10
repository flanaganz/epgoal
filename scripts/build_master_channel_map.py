#!/usr/bin/env python3
"""
build_master_channel_map.py

Bygger en UDVIDET udgave af channel_priority.xlsx med separate ID-kolonner
pr. kilde, i stedet for én stor blandet aliaskolonne. Løser problemet hvor
EPG-data findes i XML'en, men under et andet channel-id end det, UHF's
M3U-playliste rent faktisk bruger (tvg-id).

Kilder der sammenholdes (alle matchet via samme normaliserede nøgle):
    - UHF/M3U tvg-id'er   <- din faktiske IPTV-playliste (PÅKRÆVET argument,
                             indeholder login og gemmes derfor ALDRIG i en fil)
    - EPGShare ID'er      <- output_epgshare/epgshare_dk1.xml (hele kilden,
                             ikke kun de allerede filtrerede 52 kanaler)
    - OpenEPG ID'er       <- cachede denmark1-6.xml fra tidligere koerelser
                             (output_epgshare/openepg_raw/), downloades friskt
                             hvis cachen ikke findes
    - BSS ID'er           <- cachet bss_epg.xml (output_epgshare/bss_raw/),
                             downloades friskt via BSS_XMLTV_URL fra .env
                             hvis cachen ikke findes

Output:
    data/channel_priority_v2.xlsx   (ny fil - overskriver IKKE originalen)

Brug:
    python scripts/build_master_channel_map.py "http://n2ip.tv:2095/get.php?username=...&password=...&type=m3u_plus&output=ts"
"""

import os
import re
import sys
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

import requests
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
CONFIG_FILE = ROOT / "config.json"
CHANNEL_PRIORITY = DATA_DIR / "channel_priority.xlsx"
OUTPUT_FILE = DATA_DIR / "channel_priority_v2.xlsx"

OUTPUT_EPGSHARE_DIR = ROOT / "output_epgshare"
EPGSHARE_RAW_XML = OUTPUT_EPGSHARE_DIR / "epgshare_dk1.xml"
OPENEPG_RAW_DIR = OUTPUT_EPGSHARE_DIR / "openepg_raw"
BSS_RAW_FILE = OUTPUT_EPGSHARE_DIR / "bss_raw" / "bss_epg.xml"

EXTINF_ATTR_PATTERN = re.compile(r'([\w-]+)="([^"]*)"')

ENV_FILE = ROOT / ".env"
if ENV_FILE.exists():
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

BSS_XMLTV_URL = os.environ.get("BSS_XMLTV_URL")


def normalize(s: str) -> str:
    """Fjerner alt undtagen a-z0-9, så 'DR1 HD', 'dr1.dk', 'DR1.dk' og
    'DR1 Denmark (DK,DA).dk' alle kan sammenlignes på lige fod."""
    return re.sub(r"[^a-z0-9]", "", (s or "").lower())


def load_json(path: Path, default):
    import json
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return default


# ---------------------------------------------------------------------------
# 1. Læs eksisterende channel_priority.xlsx
# ---------------------------------------------------------------------------

def load_existing_priority():
    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb["Kanal-prioritering"] if "Kanal-prioritering" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]

    kanal_col = headers.index("Kanal")
    key_col = headers.index("Gruppe-nøgle (intern)")
    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")
    artwork_col = headers.index("Artwork ikke pririoteret") if "Artwork ikke pririoteret" in headers else None

    rows = []
    for row in ws.iter_rows(min_row=2):
        kanal = row[kanal_col].value
        if not kanal:
            continue
        key = row[key_col].value or ""
        aliases_raw = row[aliases_col].value or ""
        aliases = [a.strip() for a in str(aliases_raw).split(",") if a.strip()]
        follow = str(row[follow_col].value or "").strip().upper()
        artwork = ""
        if artwork_col is not None:
            artwork = str(row[artwork_col].value or "").strip().upper()

        tokens = {normalize(key)} | {normalize(a) for a in aliases} | {normalize(str(kanal))}
        tokens.discard("")

        rows.append({
            "kanal": str(kanal).strip(),
            "gruppe_noegle": str(key).strip(),
            "foelg": follow,
            "artwork_ikke_prioriteret": artwork,
            "gamle_aliaser": aliases,
            "tokens": tokens,
        })

    return rows


# ---------------------------------------------------------------------------
# 2. M3U -> tvg-id'er (PÅKRÆVET argument, aldrig gemt i en fil)
# ---------------------------------------------------------------------------

def load_m3u_entries(m3u_url_or_path: str):
    """Returnerer liste af (tvg_id, tvg_name, display_name)."""
    if m3u_url_or_path.startswith("http://") or m3u_url_or_path.startswith("https://"):
        print("📥 Henter M3U-playliste ...")
        resp = requests.get(m3u_url_or_path, timeout=120)
        resp.raise_for_status()
        text = resp.text
    else:
        print(f"📄 Læser lokal M3U-fil: {m3u_url_or_path}")
        text = Path(m3u_url_or_path).read_text(encoding="utf-8", errors="ignore")

    entries = []
    for line in text.splitlines():
        if not line.startswith("#EXTINF"):
            continue
        attrs = dict(EXTINF_ATTR_PATTERN.findall(line))
        tvg_id = attrs.get("tvg-id", "").strip()
        tvg_name = attrs.get("tvg-name", "").strip()
        display_name = line.rsplit(",", 1)[-1].strip() if "," in line else ""
        entries.append((tvg_id, tvg_name, display_name))

    print(f"✅ {len(entries):,} #EXTINF-poster fundet i M3U.\n")
    return entries


# ---------------------------------------------------------------------------
# 3. EPGShare / OpenEPG / BSS -> channel-id'er (+ display-names)
# ---------------------------------------------------------------------------

def load_channel_ids_from_xml(path: Path):
    """Returnerer liste af (channel_id, [display_names])."""
    if not path.exists():
        return []
    root = ET.parse(path).getroot()
    result = []
    for ch in root.findall("channel"):
        cid = ch.get("id", "")
        names = [dn.text or "" for dn in ch.findall("display-name")]
        result.append((cid, names))
    return result


def fetch_openepg_channels():
    """Bruger cachede filer fra openepg_coverage_check.py hvis de findes,
    ellers downloader friskt (uden at skrive cache - det gør det
    dedikerede coverage-script)."""
    if OPENEPG_RAW_DIR.exists():
        cached = sorted(OPENEPG_RAW_DIR.glob("*.xml"))
        if cached:
            print(f"📄 Genbruger {len(cached)} cachede OpenEPG-filer fra {OPENEPG_RAW_DIR}")
            result = []
            for f in cached:
                result.extend(load_channel_ids_from_xml(f))
            return result

    config = load_json(CONFIG_FILE, {})
    sources = config.get("sources", [])
    if not sources:
        print("⚠️  Ingen OpenEPG-cache og ingen kilder i config.json - springer over.")
        return []

    result = []
    for source in sources:
        print(f"📥 Henter OpenEPG-kilde {source['name']} ...")
        resp = requests.get(source["url"], timeout=90)
        resp.raise_for_status()
        root = ET.fromstring(resp.content)
        for ch in root.findall("channel"):
            cid = ch.get("id", "")
            names = [dn.text or "" for dn in ch.findall("display-name")]
            result.append((cid, names))
    return result


def fetch_bss_channels():
    if BSS_RAW_FILE.exists():
        print(f"📄 Genbruger cachet BSS-fil: {BSS_RAW_FILE}")
        return load_channel_ids_from_xml(BSS_RAW_FILE)

    if not BSS_XMLTV_URL:
        print("ℹ️  Ingen BSS-cache og BSS_XMLTV_URL er ikke sat i .env - springer BSS over.")
        return []

    print("📥 Henter BSS xmltv ...")
    resp = requests.get(BSS_XMLTV_URL, timeout=300)
    resp.raise_for_status()
    root = ET.fromstring(resp.content)
    result = []
    for ch in root.findall("channel"):
        cid = ch.get("id", "")
        names = [dn.text or "" for dn in ch.findall("display-name")]
        result.append((cid, names))
    return result


def fetch_epgshare_channels():
    if not EPGSHARE_RAW_XML.exists():
        print(f"⚠️  {EPGSHARE_RAW_XML} findes ikke - kør epgshare_download.py først.")
        return []
    return load_channel_ids_from_xml(EPGSHARE_RAW_XML)


# ---------------------------------------------------------------------------
# 4. Match hver kilde mod hver Følg(X)-kanal
# ---------------------------------------------------------------------------

def match_source_ids(tokens: set, source_entries):
    """source_entries: liste af (id, [display_names]). Returnerer sorteret
    liste af UNIKKE id'er hvis normaliserede form (eller en af dets
    display-names) matcher et af tokens."""
    matched = []
    seen = set()
    for cid, names in source_entries:
        candidates = [cid] + list(names)
        if any(normalize(c) in tokens for c in candidates if c):
            if cid not in seen:
                matched.append(cid)
                seen.add(cid)
    return sorted(matched)


def match_m3u_ids(tokens: set, m3u_entries):
    """Matcher på BÅDE tvg-id og tvg-name/display-name, da M3U-providers er
    inkonsistente med hvad de lægger hvor. Returnerer unikke tvg-id'er."""
    matched = []
    seen = set()
    for tvg_id, tvg_name, display_name in m3u_entries:
        candidates = [tvg_id, tvg_name, display_name]
        if any(normalize(c) in tokens for c in candidates if c):
            key = tvg_id or f"(uden tvg-id: {display_name})"
            if key not in seen:
                matched.append(key)
                seen.add(key)
    return sorted(matched)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        sys.exit(
            "❌ Mangler M3U URL eller filsti.\n"
            "Brug: python build_master_channel_map.py \"http://n2ip.tv:2095/get.php?username=...&password=...&type=m3u_plus&output=ts\"\n"
            "   eller: python build_master_channel_map.py C:\\sti\\til\\tv_channels.m3u8"
        )

    m3u_source = sys.argv[1]

    print("=== Bygger master kanal-mapping ===\n")

    existing_rows = load_existing_priority()
    print(f"📋 {len(existing_rows)} kanaler fundet i channel_priority.xlsx\n")

    m3u_entries = load_m3u_entries(m3u_source)
    epgshare_entries = fetch_epgshare_channels()
    print(f"✅ EPGShare: {len(epgshare_entries):,} <channel> fundet.\n")
    openepg_entries = fetch_openepg_channels()
    print(f"✅ OpenEPG: {len(openepg_entries):,} <channel> fundet.\n")
    bss_entries = fetch_bss_channels()
    print(f"✅ BSS: {len(bss_entries):,} <channel> fundet.\n")

    # --- Byg ny Excel-fil ---
    wb = Workbook()
    ws = wb.active
    ws.title = "Kanal-prioritering v2"

    headers = [
        "Kanal",
        "Gruppe-nøgle (intern)",
        "Følg (X)",
        "Artwork ikke pririoteret",
        "UHF/M3U tvg-id'er",
        "EPGShare ID'er",
        "OpenEPG ID'er",
        "BSS ID'er",
        "Sammenlagte kanal-ID'er (gammel liste, til reference)",
    ]
    ws.append(headers)

    multi_variant_channels = []  # kanaler hvor M3U har >1 tvg-id (HD/FHD-dedup-kandidat)
    no_m3u_match = []

    for row in existing_rows:
        tokens = row["tokens"]

        m3u_ids = match_m3u_ids(tokens, m3u_entries)
        epgshare_ids = match_source_ids(tokens, epgshare_entries)
        openepg_ids = match_source_ids(tokens, openepg_entries)
        bss_ids = match_source_ids(tokens, bss_entries)

        ws.append([
            row["kanal"],
            row["gruppe_noegle"],
            row["foelg"],
            row["artwork_ikke_prioriteret"],
            ", ".join(m3u_ids),
            ", ".join(epgshare_ids),
            ", ".join(openepg_ids),
            ", ".join(bss_ids),
            ", ".join(row["gamle_aliaser"]),
        ])

        if len(m3u_ids) > 1:
            multi_variant_channels.append((row["kanal"], m3u_ids))
        if not m3u_ids and row["foelg"] == "X":
            no_m3u_match.append(row["kanal"])

    # Kolonnebredder, så det er til at læse
    widths = [26, 22, 10, 14, 40, 30, 30, 30, 60]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    print("=" * 90)
    print(f"💾 Gemt: {OUTPUT_FILE}")
    print("=" * 90)

    print(f"\n🔁 Kanaler med FLERE tvg-id-varianter i M3U (HD/FHD-dedup-kandidater): {len(multi_variant_channels)}")
    for kanal, ids in multi_variant_channels[:30]:
        print(f"   {kanal:30s} -> {ids}")
    if len(multi_variant_channels) > 30:
        print(f"   ... og {len(multi_variant_channels) - 30} flere (se Excel-filen)")

    print(f"\n❌ Følg(X)-kanaler UDEN noget match i M3U: {len(no_m3u_match)}")
    for kanal in no_m3u_match:
        print(f"   {kanal}")

    print("\n=== Færdig. Gennemgå channel_priority_v2.xlsx manuelt før den erstatter originalen. ===")


if __name__ == "__main__":
    main()
