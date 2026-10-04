#!/usr/bin/env python3
r"""
build_master_channel_map_v2.py

Henter og gemmer rådata fra:
  1. EPGShare01 XML
  2. OpenEPG Denmark1-6 XML
  3. BSS XMLTV
  4. BSS M3U

Sammenholder ID'er med data/channel_priority.xlsx og genererer:
  data/channel_priority_v2.xlsx

Scriptet kræver ingen kommandolinjeargumenter.

Krav i C:\EPGoal\.env:
  BSS_M3U_URL=http://.../get.php?username=...&password=...&type=m3u_plus&output=ts
  BSS_XMLTV_URL=http://.../xmltv.php?username=...&password=...

Valgfrit i .env:
  EPGSHARE_URL=https://epgshare01.online/epgshare01/epg_ripper_DK1.xml.gz

Originalen data/channel_priority.xlsx overskrives aldrig.
"""

from __future__ import annotations

import gzip
import json
import os
import re
import sys
from pathlib import Path
from xml.etree import ElementTree as ET

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
RAW_DIR = ROOT / "output_epgshare" / "raw_sources"
OPENEPG_DIR = RAW_DIR / "openepg"

CHANNEL_PRIORITY = DATA_DIR / "channel_priority.xlsx"
OUTPUT_FILE = DATA_DIR / "channel_priority_v2.xlsx"
CONFIG_FILE = ROOT / "config.json"
ENV_FILE = ROOT / ".env"

EPGSHARE_RAW_GZ = RAW_DIR / "epgshare_dk1.xml.gz"
EPGSHARE_RAW_XML = RAW_DIR / "epgshare_dk1.xml"
BSS_XMLTV_RAW = RAW_DIR / "bss_epg.xml"
BSS_M3U_RAW = RAW_DIR / "bss_channels.m3u8"

DEFAULT_EPGSHARE_URL = (
    "https://epgshare01.online/epgshare01/epg_ripper_DK1.xml.gz"
)

EXTINF_ATTR_PATTERN = re.compile(r'([\w-]+)="([^"]*)"')
HTTP = requests.Session()


def load_env() -> None:
    if not ENV_FILE.exists():
        return
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(
            key.strip(), value.strip().strip('"').strip("'")
        )


load_env()
BSS_M3U_URL = os.environ.get("BSS_M3U_URL")
BSS_XMLTV_URL = os.environ.get("BSS_XMLTV_URL")
EPGSHARE_URL = os.environ.get("EPGSHARE_URL", DEFAULT_EPGSHARE_URL)


def normalize(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def split_aliases(value) -> list[str]:
    if not value:
        return []
    result = []
    seen = set()
    for part in str(value).split(","):
        alias = part.strip()
        key = alias.lower()
        if alias and key not in seen:
            seen.add(key)
            result.append(alias)
    return result


def load_json(path: Path, default):
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default


def download(url: str, destination: Path, timeout: int) -> bytes:
    destination.parent.mkdir(parents=True, exist_ok=True)
    response = HTTP.get(url, timeout=timeout)
    response.raise_for_status()
    destination.write_bytes(response.content)
    return response.content


def download_with_cache(
    label: str,
    url: str | None,
    destination: Path,
    timeout: int,
) -> bytes:
    if not url:
        if destination.exists():
            print(f"{label}: URL mangler. Genbruger cache: {destination}")
            return destination.read_bytes()
        raise RuntimeError(
            f"{label}: URL mangler, og cache findes ikke: {destination}"
        )

    print(f"Henter {label}...")
    try:
        content = download(url, destination, timeout)
        print(
            f"Gemt {label}: {destination} "
            f"({len(content) / 1024 / 1024:.1f} MB)"
        )
        return content
    except requests.RequestException as exc:
        if destination.exists():
            print(
                f"ADVARSEL: {label} kunne ikke hentes ({exc}). "
                f"Genbruger cache: {destination}"
            )
            return destination.read_bytes()
        raise


def load_priority_rows() -> list[dict]:
    if not CHANNEL_PRIORITY.exists():
        sys.exit(f"FEJL: Mangler {CHANNEL_PRIORITY}")

    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    sheet_name = "Kanal-prioritering"
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]

    required = [
        "Kanal",
        "Gruppe-nøgle (intern)",
        "Sammenlagte kanal-ID'er",
        "Følg (X)",
    ]
    missing = [name for name in required if name not in headers]
    if missing:
        sys.exit(f"FEJL: Manglende kolonner: {missing}")

    canal_idx = headers.index("Kanal")
    group_idx = headers.index("Gruppe-nøgle (intern)")
    alias_idx = headers.index("Sammenlagte kanal-ID'er")
    follow_idx = headers.index("Følg (X)")
    artwork_idx = (
        headers.index("Artwork ikke pririoteret")
        if "Artwork ikke pririoteret" in headers
        else None
    )

    rows = []
    for excel_row in ws.iter_rows(min_row=2):
        canal = excel_row[canal_idx].value
        if not canal:
            continue

        group_key = str(excel_row[group_idx].value or "").strip()
        aliases = split_aliases(excel_row[alias_idx].value)
        follow = str(excel_row[follow_idx].value or "").strip().upper()
        artwork_skip = ""
        if artwork_idx is not None:
            artwork_skip = str(
                excel_row[artwork_idx].value or ""
            ).strip().upper()

        tokens = {normalize(str(canal)), normalize(group_key)}
        tokens.update(normalize(alias) for alias in aliases)
        tokens.discard("")

        rows.append(
            {
                "canal": str(canal).strip(),
                "group_key": group_key,
                "follow": follow,
                "artwork_skip": artwork_skip,
                "aliases": aliases,
                "tokens": tokens,
            }
        )
    return rows


def fetch_epgshare() -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    raw_gz = download_with_cache(
        "EPGShare01",
        EPGSHARE_URL,
        EPGSHARE_RAW_GZ,
        300,
    )
    try:
        xml_data = gzip.decompress(raw_gz)
    except OSError:
        xml_data = raw_gz
    EPGSHARE_RAW_XML.write_bytes(xml_data)
    print(
        f"Udpakket EPGShare01: {EPGSHARE_RAW_XML} "
        f"({len(xml_data) / 1024 / 1024:.1f} MB)"
    )
    return EPGSHARE_RAW_XML


def openepg_sources() -> list[dict]:
    config = load_json(CONFIG_FILE, {})
    sources = config.get("sources", [])
    valid = [
        source
        for source in sources
        if source.get("name") and source.get("url")
    ]
    if valid:
        return valid
    return [
        {
            "name": f"denmark{number}",
            "url": (
                "https://www.open-epg.com/files/"
                f"denmark{number}.xml"
            ),
        }
        for number in range(1, 7)
    ]


def fetch_openepg() -> list[Path]:
    OPENEPG_DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    for source in openepg_sources():
        path = OPENEPG_DIR / f"{source['name']}.xml"
        download_with_cache(
            f"OpenEPG {source['name']}",
            source["url"],
            path,
            180,
        )
        paths.append(path)
    return paths


def fetch_bss_xmltv() -> Path | None:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    if not BSS_XMLTV_URL and not BSS_XMLTV_RAW.exists():
        print(
            "ADVARSEL: BSS_XMLTV_URL mangler i .env, "
            "og BSS XMLTV-cache findes ikke."
        )
        return None

    download_with_cache(
        "BSS XMLTV",
        BSS_XMLTV_URL,
        BSS_XMLTV_RAW,
        420,
    )
    return BSS_XMLTV_RAW


def fetch_bss_m3u() -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    if not BSS_M3U_URL and not BSS_M3U_RAW.exists():
        sys.exit(
            "FEJL: BSS_M3U_URL mangler i C:\\EPGoal\\.env, "
            "og der findes ingen lokal BSS M3U-cache."
        )

    download_with_cache(
        "BSS M3U",
        BSS_M3U_URL,
        BSS_M3U_RAW,
        240,
    )
    return BSS_M3U_RAW


def xml_channel_entries(paths: list[Path]) -> list[tuple[str, list[str]]]:
    entries = []
    for path in paths:
        if not path or not path.exists():
            continue
        root = ET.parse(path).getroot()
        for channel in root.findall("channel"):
            channel_id = (channel.get("id") or "").strip()
            display_names = [
                (node.text or "").strip()
                for node in channel.findall("display-name")
                if (node.text or "").strip()
            ]
            if channel_id:
                entries.append((channel_id, display_names))
    return entries


def m3u_entries(path: Path) -> list[tuple[str, str, str, str]]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    entries = []
    for line in text.splitlines():
        if not line.startswith("#EXTINF"):
            continue
        attrs = dict(EXTINF_ATTR_PATTERN.findall(line))
        tvg_id = attrs.get("tvg-id", "").strip()
        tvg_name = attrs.get("tvg-name", "").strip()
        group_title = attrs.get("group-title", "").strip()
        display_name = line.rsplit(",", 1)[-1].strip() if "," in line else ""
        entries.append((tvg_id, tvg_name, display_name, group_title))
    return entries


def match_xml(tokens: set[str], entries) -> list[str]:
    matches = []
    seen = set()
    for channel_id, display_names in entries:
        source_tokens = {
            normalize(channel_id),
            *(normalize(name) for name in display_names),
        }
        source_tokens.discard("")
        if not tokens.intersection(source_tokens):
            continue
        key = channel_id.lower()
        if key not in seen:
            seen.add(key)
            matches.append(channel_id)
    return sorted(matches, key=str.lower)


def match_m3u(tokens: set[str], entries):
    ids = []
    streams = []
    seen_ids = set()
    seen_streams = set()

    for tvg_id, tvg_name, display_name, group_title in entries:
        source_tokens = {
            normalize(tvg_id),
            normalize(tvg_name),
            normalize(display_name),
        }
        source_tokens.discard("")
        if not tokens.intersection(source_tokens):
            continue

        if tvg_id and tvg_id.lower() not in seen_ids:
            seen_ids.add(tvg_id.lower())
            ids.append(tvg_id)

        label = display_name or tvg_name
        if label:
            stream_value = (
                f"{label} [{group_title}]"
                if group_title
                else label
            )
            if stream_value.lower() not in seen_streams:
                seen_streams.add(stream_value.lower())
                streams.append(stream_value)

    return sorted(ids, key=str.lower), sorted(streams, key=str.lower)


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = [
        28, 24, 10, 24, 42,
        55, 34, 44, 44, 68,
    ]
    for index, width in enumerate(widths, start=1):
        letter = ws.cell(row=1, column=index).column_letter
        ws.column_dimensions[letter].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def create_workbook(
    priority_rows,
    m3u_data,
    epgshare_data,
    openepg_data,
    bss_data,
) -> tuple[int, list[str]]:
    wb = Workbook()
    ws = wb.active
    ws.title = "Kanal-prioritering v2"
    ws.append(
        [
            "Kanal",
            "Gruppe-nøgle (intern)",
            "Følg (X)",
            "Artwork ikke pririoteret",
            "BSS M3U tvg-id'er",
            "BSS M3U stream-navne",
            "EPGShare ID'er",
            "OpenEPG ID'er",
            "BSS XMLTV ID'er",
            "Sammenlagte kanal-ID'er (gammel reference)",
        ]
    )

    multi_id_count = 0
    missing_m3u = []

    for row in priority_rows:
        ids, stream_names = match_m3u(row["tokens"], m3u_data)
        epgshare_ids = match_xml(row["tokens"], epgshare_data)
        openepg_ids = match_xml(row["tokens"], openepg_data)
        bss_ids = match_xml(row["tokens"], bss_data)

        ws.append(
            [
                row["canal"],
                row["group_key"],
                row["follow"],
                row["artwork_skip"],
                ", ".join(ids),
                ", ".join(stream_names),
                ", ".join(epgshare_ids),
                ", ".join(openepg_ids),
                ", ".join(bss_ids),
                ", ".join(row["aliases"]),
            ]
        )

        if len(ids) > 1:
            multi_id_count += 1
        if row["follow"] == "X" and not ids:
            missing_m3u.append(row["canal"])

    style_sheet(ws)

    guide = wb.create_sheet("Vejledning")
    guide.column_dimensions["A"].width = 115
    guide_content = [
        "Formål",
        "Sammenligner tvg-id og channel-id fra BSS M3U, EPGShare, OpenEPG og BSS XMLTV.",
        "Originalen channel_priority.xlsx overskrives ikke.",
        "",
        "Kolonne E",
        "Faktiske tvg-id'er fra BSS M3U.",
        "",
        "Kolonne F",
        "Faktiske streamnavne, inklusive normal/HD/FHD/4K hvor det fremgår.",
        "",
        "Kolonne G-I",
        "Channel-ID'er fra de tre rå EPG-kilder.",
        "",
        "Kontrol",
        "Gennemgå især Følg=X-rækker uden M3U tvg-id og rækker hvor kildernes ID'er ikke matcher.",
    ]
    for value in guide_content:
        guide.append([value])
    guide["A1"].font = Font(bold=True)
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    return multi_id_count, missing_m3u


def main() -> None:
    print("=== build_master_channel_map_v2 ===")
    RAW_DIR.mkdir(parents=True, exist_ok=True)

    priority_rows = load_priority_rows()
    print(f"channel_priority.xlsx: {len(priority_rows)} rækker.")

    epgshare_path = fetch_epgshare()
    openepg_paths = fetch_openepg()
    bss_xmltv_path = fetch_bss_xmltv()
    bss_m3u_path = fetch_bss_m3u()

    print("Indlæser channel-ID'er fra råfiler...")
    epgshare_data = xml_channel_entries([epgshare_path])
    openepg_data = xml_channel_entries(openepg_paths)
    bss_data = xml_channel_entries(
        [bss_xmltv_path] if bss_xmltv_path else []
    )
    m3u_data = m3u_entries(bss_m3u_path)

    print(f"EPGShare channel-poster: {len(epgshare_data):,}")
    print(f"OpenEPG channel-poster: {len(openepg_data):,}")
    print(f"BSS XMLTV channel-poster: {len(bss_data):,}")
    print(f"BSS M3U EXTINF-poster: {len(m3u_data):,}")

    multi_count, missing_m3u = create_workbook(
        priority_rows,
        m3u_data,
        epgshare_data,
        openepg_data,
        bss_data,
    )

    print(f"Gemt: {OUTPUT_FILE}")
    print(f"Råfiler gemt i: {RAW_DIR}")
    print(f"Rækker med flere M3U tvg-id'er: {multi_count}")
    print(f"Følg=X uden fundet M3U tvg-id: {len(missing_m3u)}")
    for canal in missing_m3u:
        print(f"  {canal}")
    print("=== Færdig ===")


if __name__ == "__main__":
    main()
