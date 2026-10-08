#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Analyser streams som mangler EPG-match.

Sammenligner channel_priority_v2.xlsx med epgoal.xml og epgshare_merge_log.json.
Read-only. Ændrer ingen kilder.
"""
from __future__ import annotations
import argparse, csv, json, re, unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE=Path(r"C:\EPGoal\output_epgshare")
PRIORITY=Path(r"C:\EPGoal\data\channel_priority_v2.xlsx")
REPORTS=Path(r"C:\EPGoal\reports")
QUALITY={"SD","HD","FHD","UHD","4K"}

def text(v): return str(v or "").strip()
def split(v): return [x.strip() for x in re.split(r"[,;\n]+",text(v)) if x.strip() and x.strip().upper()!="X"]
def normalized(v):
    s="".join(c for c in unicodedata.normalize("NFKD",text(v).lower()) if not unicodedata.combining(c))
    s=s.replace("+","plus")
    return re.sub(r"[^a-z0-9]","",s)
def quality(name):
    u=text(name).upper()
    if re.search(r"\b(4K|UHD)\b",u): return "UHD"
    if re.search(r"\bFHD\b",u): return "FHD"
    if re.search(r"\bHD\b",u): return "HD"
    if re.search(r"\bSD\b",u): return "SD"
    return "NORMAL"
def canonical_stream_name(name):
    s="".join(c for c in unicodedata.normalize("NFKD",text(name).lower()) if not unicodedata.combining(c))
    s=re.sub(r"\[[^]]*]|\([^)]*\)"," ",s)
    toks=[x for x in re.findall(r"[a-z0-9+]+",s) if x.upper() not in QUALITY and x not in {"dk","denmark","danmark","multi","audio","t","d"}]
    return "".join(toks)

def header_row(ws):
    for r in range(1,min(ws.max_row,40)+1):
        if text(ws.cell(r,1).value)=="Kanal": return r
    raise ValueError("Kolonnen 'Kanal' blev ikke fundet")
def load_priority(path):
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; hr=header_row(ws)
    headers=[text(ws.cell(hr,c).value) for c in range(1,ws.max_column+1)]
    rows=[]
    for r in range(hr+1,ws.max_row+1):
        if not text(ws.cell(r,1).value): continue
        rows.append({headers[c-1]:ws.cell(r,c).value for c in range(1,ws.max_column+1) if headers[c-1]})
    return rows

def load_epgoal(path):
    names=defaultdict(list); counts=Counter()
    if not path.exists(): return names,counts
    for _,e in ET.iterparse(path,events=("end",)):
        tag=e.tag.rsplit("}",1)[-1]
        if tag=="channel":
            cid=text(e.get("id")); names[cid]=[text(x.text) for x in e.findall("./display-name") if text(x.text)]; e.clear()
        elif tag=="programme":
            cid=text(e.get("channel")); counts[cid]+=1; e.clear()
    return names,counts

def load_merge(path):
    by_name={}
    if not path.exists(): return by_name
    data=json.loads(path.read_text(encoding="utf-8-sig"))
    for section in ("kept_as_is","supplemented","no_fallback_available"):
        for row in data.get(section,[]):
            rec=dict(row);rec["section"]=section;by_name[normalized(row.get("kanal"))]=rec
    return by_name

def pair_streams(ids,names):
    if not names: return []
    if not ids: return [("",n,"no-id") for n in names]
    if len(ids)==len(names): return [(ids[i],names[i],"position") for i in range(len(names))]
    if len(ids)==1: return [(ids[0],n,"shared-single-id") for n in names]
    out=[]
    for n in names:
        cn=canonical_stream_name(n); scored=[]
        for sid in ids:
            si=normalized(sid)
            score=0
            if si and si in cn: score=len(si)+20
            elif cn and cn in si: score=len(cn)+10
            scored.append((score,sid))
        best=max(scored,default=(0,""))
        out.append((best[1] if best[0]>0 else "",n,"name-heuristic" if best[0]>0 else "ambiguous"))
    return out

def resolve_epgoal_id(output_id,stream_id,known_ids,ep_names,ep_counts):
    all_ids=set(ep_names)|set(ep_counts);by_norm=defaultdict(list)
    for cid in all_ids: by_norm[normalized(cid)].append(cid)
    candidates=[]
    for value,label in ((output_id,"output"),(stream_id,"stream")):
        if value in all_ids:candidates.append((value,label+"-exact"))
        for cid in by_norm.get(normalized(value),[]):candidates.append((cid,label+"-normalized"))
    for kid in known_ids:
        if kid in all_ids:candidates.append((kid,"priority-exact"))
        for cid in by_norm.get(normalized(kid),[]):candidates.append((cid,"priority-normalized"))
    if not candidates:return "","not-found"
    # output-id har prioritet, dernæst flest programmer.
    candidates=list(dict.fromkeys(candidates))
    candidates.sort(key=lambda x:(0 if x[1].startswith("output") else 1,-ep_counts.get(x[0],0),x[0].lower()))
    return candidates[0]

def classify(stream_id,output_id,resolved_id,programmes,pair_method):
    notes=[]
    if not stream_id:return "MANGLER TVG-ID","Streamen har intet tvg-id i priority-arket"
    if pair_method=="ambiguous":notes.append("ID kan ikke sikkert parres med streamnavnet")
    if not resolved_id:return "ID FINDES IKKE I EPGOAL","Hverken output-, stream- eller kendt kilde-ID findes i epgoal.xml"
    if programmes==0:return "INGEN PROGRAMMER",f"EPGoal-kanalen {resolved_id} har 0 programmer"
    if normalized(stream_id)!=normalized(output_id):
        notes.append(f"stream-ID {stream_id} afviger fra output-ID {output_id}")
        return "TVG-ID AFVIGER", "; ".join(notes)
    return ("KONTROLLÉR PARRING" if notes else "OK"), "; ".join(notes)

def add_sheet(wb,title,headers,rows):
    ws=wb.create_sheet(title[:31]);ws.append(headers)
    for row in rows:ws.append([row.get(h,"") for h in headers])
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;fill=PatternFill("solid",fgColor="1F4E78")
    colors={"OK":"C6EFCE","TVG-ID AFVIGER":"FFEB9C","KONTROLLÉR PARRING":"FFEB9C","INGEN PROGRAMMER":"FFC7CE","MANGLER TVG-ID":"FFC7CE","ID FINDES IKKE I EPGOAL":"F4B084"}
    for c in ws[1]:c.font=Font(color="FFFFFF",bold=True);c.fill=fill;c.alignment=Alignment(wrap_text=True)
    for r in range(2,ws.max_row+1):
        st=text(ws.cell(r,1).value)
        if st in colors:ws.cell(r,1).fill=PatternFill("solid",fgColor=colors[st])
        for c in ws[r]:c.alignment=Alignment(vertical="top",wrap_text=True)
    for col in range(1,ws.max_column+1):
        vals=[len(text(ws.cell(r,col).value)) for r in range(1,min(ws.max_row,300)+1)]
        ws.column_dimensions[get_column_letter(col)].width=min(60,max(11,max(vals,default=9)+2))
    return ws

def main():
    ap=argparse.ArgumentParser(description="Find streams hvis tvg-id ikke matcher epgoal.xml")
    ap.add_argument("--base",type=Path,default=BASE);ap.add_argument("--priority",type=Path,default=PRIORITY);ap.add_argument("--report-dir",type=Path,default=REPORTS)
    a=ap.parse_args();a.report_dir.mkdir(parents=True,exist_ok=True)
    if not a.priority.exists():raise SystemExit(f"FEJL: Mangler {a.priority}")
    print(f"Prioritetsfil: {a.priority}");priority=load_priority(a.priority)
    ep_path=a.base/"epgoal.xml";print(f"EPGoal: {ep_path}");ep_names,ep_counts=load_epgoal(ep_path)
    merge_path=a.base/"epgshare_merge_log.json";merge=load_merge(merge_path);print(f"Merge-log: {merge_path if merge_path.exists() else 'IKKE FUNDET'}")
    results=[];groups=[];statuses=Counter()
    for rec in priority:
        group=text(rec.get("Kanal"));output_id=text(rec.get("Output/UHF tvg-id"));ids=split(rec.get("BSS M3U tvg-id'er"));names=split(rec.get("BSS M3U stream-navne"))
        known=set()
        for col in ("EPGShare ID'er","OpenEPG ID'er","BSS XMLTV ID'er","Output/UHF tvg-id"):
            known.update(split(rec.get(col)))
        merge_rec=merge.get(normalized(group),{});group_stats=Counter()
        for stream_id,stream_name,pair_method in pair_streams(ids,names):
            resolved,match_method=resolve_epgoal_id(output_id,stream_id,known,ep_names,ep_counts);programmes=ep_counts.get(resolved,0)
            status,note=classify(stream_id,output_id,resolved,programmes,pair_method);statuses[status]+=1;group_stats[status]+=1
            merge_programmes=merge_rec.get("programmer",merge_rec.get("programmer_indsat",merge_rec.get("epgshare_programmer",merge_rec.get("epgshare_programmer_foer",""))))
            results.append({
                "Status":status,"Bemærkning":note,"Kanalgruppe":group,"Kvalitet":quality(stream_name),"Streamnavn":stream_name,
                "Stream tvg-id":stream_id,"Output/UHF tvg-id":output_id,"EPGoal kanal-id":resolved,"EPGoal programmer":programmes,
                "EPGoal display-name":", ".join(ep_names.get(resolved,[])),"Matchmetode":match_method,"Parringsmetode":pair_method,
                "Kendte priority-ID'er":", ".join(sorted(known)),"Merge-status":merge_rec.get("section",""),
                "Merge-kilde":merge_rec.get("kilde",merge_rec.get("kilde_brugt","")),"Merge-programmer":merge_programmes,
                "Merge-kilde-ID":merge_rec.get("kilde_channel_id",""),
            })
        severity=next((x for x in ("MANGLER TVG-ID","ID FINDES IKKE I EPGOAL","INGEN PROGRAMMER","TVG-ID AFVIGER","KONTROLLÉR PARRING") if group_stats[x]),"OK")
        groups.append({"Status":severity,"Kanalgruppe":group,"Output/UHF tvg-id":output_id,"EPGoal programmer":ep_counts.get(output_id,0),"Streams":len(names),"Stream-ID'er":len(ids),"Problemer":sum(v for k,v in group_stats.items() if k!="OK"),"Statusfordeling":", ".join(f"{k}={v}" for k,v in group_stats.items())})
    h=["Status","Bemærkning","Kanalgruppe","Kvalitet","Streamnavn","Stream tvg-id","Output/UHF tvg-id","EPGoal kanal-id","EPGoal programmer","EPGoal display-name","Matchmetode","Parringsmetode","Kendte priority-ID'er","Merge-status","Merge-kilde","Merge-programmer","Merge-kilde-ID"]
    gh=["Status","Kanalgruppe","Output/UHF tvg-id","EPGoal programmer","Streams","Stream-ID'er","Problemer","Statusfordeling"]
    problems=[r for r in results if r["Status"]!="OK"]
    focus_names={"discovery science.dk","tv2 news.dk","dr1.dk","dr2.dk","tv2 echo.dk","tv2 fri.dk","6eren.dk","tv2sport.dk","tv2 sport x.dk","tv3+.dk"}
    focus=[r for r in results if r["Kanalgruppe"].lower() in focus_names]
    weak=[]
    for r in results:
        try: low=int(r["Merge-programmer"])<20
        except (ValueError,TypeError):low=False
        if low:weak.append(r)
    wb=Workbook();wb.remove(wb.active)
    summary=[{"Måling":"Kørt","Værdi":datetime.now().strftime("%Y-%m-%d %H:%M:%S")},{"Måling":"EPGoal kanal-id'er","Værdi":len(set(ep_names)|set(ep_counts))},{"Måling":"Streams analyseret","Værdi":len(results)},{"Måling":"Problemer","Værdi":len(problems)}]+[{"Måling":f"Status: {k}","Værdi":v} for k,v in sorted(statuses.items())]
    add_sheet(wb,"Summary",["Måling","Værdi"],summary);add_sheet(wb,"Problems",h,problems);add_sheet(wb,"Focus Channels",h,focus);add_sheet(wb,"All Streams",h,results);add_sheet(wb,"Groups",gh,groups);add_sheet(wb,"Weak EPG Sources",h,weak)
    xp=a.report_dir/"missing_stream_epg_analysis.xlsx";cp=a.report_dir/"missing_stream_epg_problems.csv";jp=a.report_dir/"missing_stream_epg_analysis.json";wb.save(xp)
    with cp.open("w",encoding="utf-8-sig",newline="") as f:w=csv.DictWriter(f,fieldnames=h,delimiter=";",extrasaction="ignore");w.writeheader();w.writerows(problems)
    jp.write_text(json.dumps({"generated":datetime.now().isoformat(timespec="seconds"),"statuses":dict(statuses),"problems":problems,"groups":groups},ensure_ascii=False,indent=2),encoding="utf-8")
    print("\n=== STREAM/EPG-ANALYSE FÆRDIG ===");print("Status: "+", ".join(f"{k}={v}" for k,v in sorted(statuses.items())));print(f"Excel : {xp}\nCSV   : {cp}\nJSON  : {jp}")
if __name__=="__main__":main()
