#!/usr/bin/env python3

from pathlib import Path
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

INPUT_XML = ROOT / "output_epgshare" / "epgshare_dk1.xml"
OUTPUT_XML = ROOT / "output_epgshare" / "epgshare_filtered.xml"

CHANNEL_PRIORITY = ROOT / "data" / "channel_priority.xlsx"


def load_followed_channels():
    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    kanal_col = headers.index("Kanal")
    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")

    wanted = set()

    for row in ws.iter_rows(min_row=2):

        follow = row[follow_col].value

        if str(follow or "").strip().upper() != "X":
            continue

        kanal = row[kanal_col].value

        if kanal:
            wanted.add(str(kanal).strip().lower())

        aliases = row[aliases_col].value

        if aliases:
            for alias in str(aliases).split(","):
                alias = alias.strip()

                if alias:
                    wanted.add(alias.lower())

    return wanted


def main():

    wanted_channels = load_followed_channels()

    print(f"Whitelist kanaler: {len(wanted_channels)}")

    tree = ET.parse(INPUT_XML)
    root = tree.getroot()

    new_root = ET.Element("tv")

    kept_channel_ids = set()

    for channel in root.findall("channel"):

        channel_id = (channel.get("id") or "").strip()

        display_names = [
            (dn.text or "").strip()
            for dn in channel.findall("display-name")
        ]

        candidates = [channel_id] + display_names

        match = False

        for value in candidates:
            if value.lower() in wanted_channels:
                match = True
                break

        if match:
            new_root.append(channel)
            kept_channel_ids.add(channel_id)

    programme_count = 0

    for programme in root.findall("programme"):

        channel_id = (programme.get("channel") or "").strip()

        if channel_id in kept_channel_ids:
            new_root.append(programme)
            programme_count += 1

    ET.ElementTree(new_root).write(
        OUTPUT_XML,
        encoding="utf-8",
        xml_declaration=True,
    )

    print()
    print("=== EPGShare filter ===")
    print(f"Kanaler beholdt   : {len(kept_channel_ids):,}")
    print(f"Programmer beholdt: {programme_count:,}")
    print(f"Output            : {OUTPUT_XML}")

    print()
    print("Matched kanaler:")
    print("----------------")

    for ch in sorted(kept_channel_ids):
        print(ch)

    print("----------------")


if __name__ == "__main__":
    main()