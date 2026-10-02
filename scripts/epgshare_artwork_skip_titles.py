#!/usr/bin/env python3

from pathlib import Path
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

INPUT_XML = ROOT / "output_epgshare" / "epgshare_filtered.xml"
CHANNEL_PRIORITY = ROOT / "data" / "channel_priority.xlsx"


def load_skip_artwork_channels():

    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")
    skip_col = headers.index("Artwork ikke pririoteret")

    skip_channels = set()

    for row in ws.iter_rows(min_row=2):

        follow = str(row[follow_col].value or "").strip().upper()
        skip = str(row[skip_col].value or "").strip().upper()

        if follow != "X":
            continue

        if skip != "X":
            continue

        aliases = row[aliases_col].value

        if aliases:
            for alias in str(aliases).split(","):
                alias = alias.strip()

                if alias:
                    skip_channels.add(alias)

    return skip_channels


def main():

    tree = ET.parse(INPUT_XML)
    root = tree.getroot()

    skip_channels = load_skip_artwork_channels()

    channel_titles = defaultdict(Counter)

    for programme in root.findall("programme"):

        channel_id = (programme.get("channel") or "").strip()

        if channel_id not in skip_channels:
            continue

        title = programme.findtext("title")

        if not title:
            continue

        channel_titles[channel_id][title.strip()] += 1

    print()
    print("===== Artwork-skip rapport =====")

    for channel in sorted(channel_titles.keys()):

        print()
        print("=" * 80)
        print(channel)
        print("=" * 80)

        total = sum(channel_titles[channel].values())

        print(f"Programmer: {total}")
        print()

        for title, count in channel_titles[channel].most_common(50):
            print(f"{count:4d}  {title}")


if __name__ == "__main__":
    main()