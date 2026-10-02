#!/usr/bin/env python3

from pathlib import Path
import xml.etree.ElementTree as ET
from collections import Counter
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

INPUT_XML = ROOT / "output_epgshare" / "epgshare_filtered.xml"
CHANNEL_PRIORITY = ROOT / "data" / "channel_priority.xlsx"


def load_skip_artwork_channels():

    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb.active

    headers = [cell.value for cell in ws[1]]

    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")
    skip_col = headers.index("Artwork ikke pririoteret")

    skip_channels = set()

    for row in ws.iter_rows(min_row=2):

        follow = str(row[follow_col].value or "").strip().upper()
        skip = str(row[skip_col].value or "").strip().upper()

        if follow != "X":
            continue

        if skip != "X":
            continue

        aliases = row[aliases_col].value

        if aliases:

            for alias in str(aliases).split(","):

                alias = alias.strip()

                if alias:
                    skip_channels.add(alias)

    return skip_channels


def main():

    tree = ET.parse(INPUT_XML)
    root = tree.getroot()

    skip_channels = load_skip_artwork_channels()

    programme_counts = Counter()

    total_programmes = 0
    skipped_programmes = 0

    for programme in root.findall("programme"):

        total_programmes += 1

        channel_id = (programme.get("channel") or "").strip()

        programme_counts[channel_id] += 1

        if channel_id in skip_channels:
            skipped_programmes += 1

    print()
    print("===== Artwork Skip Analyse =====")
    print()

    print(f"Programmer i alt              : {total_programmes:,}")
    print(f"Programmer uden artwork       : {skipped_programmes:,}")
    print(
        f"Programmer med artwork-test   : "
        f"{total_programmes - skipped_programmes:,}"
    )

    print()
    print("Kanaler markeret som artwork ikke prioriteret")
    print("-" * 80)

    for channel in sorted(skip_channels):

        count = programme_counts[channel]

        print(f"{count:4d}  {channel}")


if __name__ == "__main__":
    main()