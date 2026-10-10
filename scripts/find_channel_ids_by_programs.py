#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Find sandsynlige XMLTV channel-id'er ud fra kendte programtitler.

Mål: Discovery Science og TV2 News.
Read-only. Skriver Excel, CSV og JSON til C:\EPGoal\reports.
"""
from __future__ import annotations
import argparse, csv, json, re, unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE=Path(r"C:\EPGoal\output_epgshare")
TARGETS={
 "Discovery Science":[
  "How It's Made: Dream Cars","How the Universe Works",
  "Morgan Freeman's Through The Wormhole","Extreme Engineering",
  "Truth Behind the Moon Landing",
 ],
 "TV2 News":[
  "Morgennyhederne på News","12 News","Nyheder, sport og vejr",
  "News & Co.","18 News",
 ],
}

def txt(s):return " ".join(str(s or "").replace("’","'").split())
def norm(s):
 s=unicodedata.normalize("NFKD",txt(s).lower())
 s="".join(c for c in s if not unicodedata.combining(c))
 s=s.replace("&","and")
 return re.sub(r"[^a-z0-9]+","",s)
def title_match(actual,expected):
 a,b=norm(actual),norm(expected)
 if not a or not b:return 0.0
 if a==b:return 1.0
 if a.startswith(b) or b.startswith(a):return .94
 if a in b or b in a:return .88
 return SequenceMatcher(None,a,b).ratio()
def parse(path,source):
 channels=defaultdict(list); programs=defaultdict(list)
 if not path.exists():return channels,programs
 try:
  for _,e in ET.iterparse(path,events=("end",)):
   tag=e.tag.rsplit("}",1)[-1]
   if tag=="channel":
    cid=txt(e.get("id"))
    if cid:channels[cid].extend(txt(x.text) for x in e.findall("./display-name") if txt(x.text))
    e.clear()
   elif tag=="programme":
    cid=txt(e.get("channel"));te=e.find("./title");title=txt(te.text if te is not None else "")
    if cid and title:programs[cid].append(title)
    e.clear()
 except ET.ParseError as ex:print(f"ADVARSEL: XML-fejl i {path}: {ex}")
 return channels,programs

def analyze_source(source,path,targets):
 channels,programs=parse(path,source);rows=[]
 all_ids=set(channels)|set(programs)
 for target,expected_titles in targets.items():
  for cid in all_ids:
   actual_titles=programs.get(cid,[]);unique=list(dict.fromkeys(actual_titles));hits=[]
   for expected in expected_titles:
    best_title="";best=0.0
    for actual in unique:
     score=title_match(actual,expected)
     if score>best:best,best_title=score,actual
    if best>=.72:hits.append((expected,best_title,best))
   # Navn bruges kun som sekundært signal. Programtræffere er vigtigst.
   name_blob=" ".join([cid]+channels.get(cid,[]))
   name_score=max(title_match(name_blob,target),0)
   exact=sum(1 for _,_,s in hits if s>=.94);partial=len(hits)-exact
   score=exact*30+partial*18+min(len(unique),100)/25+name_score*12
   if hits or name_score>=.55:
    rows.append({
     "target":target,"source":source,"file":path.name,"channel_id":cid,
     "display_names":", ".join(dict.fromkeys(channels.get(cid,[]))),
     "programmes":len(actual_titles),"unique_titles":len(unique),
     "exact_hits":exact,"partial_hits":partial,"title_hits":len(hits),
     "score":round(score,1),"name_score":round(name_score*100,1),
     "matched_expected":" | ".join(x[0] for x in hits),
     "matched_actual":" | ".join(x[1] for x in hits),
     "match_details":" | ".join(f"{x[0]} => {x[1]} ({x[2]*100:.0f}%)" for x in hits),
    })
 return rows,len(all_ids)

def add_sheet(wb,title,headers,rows):
 ws=wb.create_sheet(title);ws.append(headers)
 for r in rows:ws.append([r.get(h,"") for h in headers])
 ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;fill=PatternFill("solid",fgColor="1F4E78")
 for c in ws[1]:c.font=Font(color="FFFFFF",bold=True);c.fill=fill;c.alignment=Alignment(wrap_text=True)
 for col in range(1,ws.max_column+1):
  vals=[len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,250)+1)]
  ws.column_dimensions[get_column_letter(col)].width=min(65,max(11,max(vals,default=9)+2))
 for row in ws.iter_rows(min_row=2):
  for c in row:c.alignment=Alignment(vertical="top",wrap_text=True)
 return ws

def main():
 ap=argparse.ArgumentParser(description="Find korrekte kanal-ID'er via kendte programtitler")
 ap.add_argument("--base",type=Path,default=BASE);ap.add_argument("--report-dir",type=Path)
 a=ap.parse_args();report=a.report_dir or a.base.parent/"reports";report.mkdir(parents=True,exist_ok=True)
 specs=[("EPGShare",a.base/"epgshare_dk1.xml"),("BSS",a.base/"bss_raw"/"bss_epg.xml"),("EPGoal",a.base/"epgoal.xml")]
 specs += [(f"OpenEPG:{p.stem}",p) for p in sorted((a.base/"openepg_raw").glob("*.xml"))]
 rows=[];counts={}
 for source,path in specs:
  print(f"Analyserer {source}: {path}")
  found,count=analyze_source(source,path,TARGETS);rows.extend(found);counts[source]=count
  print(f"  {count:,} kanal-id'er, {len(found):,} relevante kandidater")
 rows.sort(key=lambda r:(r["target"],-r["score"],-r["title_hits"],-r["programmes"],r["source"],r["channel_id"]))
 rankings=[]
 for target in TARGETS:
  rank=0
  for r in [x for x in rows if x["target"]==target]:
   rank+=1;r["rank"]=rank;rankings.append(r)
 best=[]
 for target in TARGETS:
  candidates=[r for r in rankings if r["target"]==target]
  if candidates:best.append(candidates[0])
 headers=["rank","target","source","file","channel_id","display_names","programmes","unique_titles","exact_hits","partial_hits","title_hits","score","name_score","matched_expected","matched_actual","match_details"]
 wb=Workbook();wb.remove(wb.active)
 summary=[{"Måling":"Kørt","Værdi":datetime.now().strftime("%Y-%m-%d %H:%M:%S")},{"Måling":"Formål","Værdi":"Find ID via kendte programtitler"}]
 for source,count in counts.items():summary.append({"Måling":f"{source} kanal-id'er","Værdi":count})
 add_sheet(wb,"Summary",["Måling","Værdi"],summary)
 add_sheet(wb,"Best Candidates",headers,best)
 add_sheet(wb,"All Candidates",headers,rankings)
 for target in TARGETS:add_sheet(wb,target[:31],headers,[r for r in rankings if r["target"]==target])
 xp=report/"channel_id_program_match_analysis.xlsx";cp=report/"channel_id_program_match_candidates.csv";jp=report/"channel_id_program_match_analysis.json"
 wb.save(xp)
 with cp.open("w",encoding="utf-8-sig",newline="") as f:
  w=csv.DictWriter(f,fieldnames=headers,delimiter=";",extrasaction="ignore");w.writeheader();w.writerows(rankings)
 jp.write_text(json.dumps({"generated":datetime.now().isoformat(timespec="seconds"),"targets":TARGETS,"source_counts":counts,"best_candidates":best,"candidates":rankings},ensure_ascii=False,indent=2),encoding="utf-8")
 print("\n=== ANALYSE FÆRDIG ===")
 for r in best:print(f"{r['target']}: {r['channel_id']} via {r['source']} | træffere={r['title_hits']} | programmer={r['programmes']} | score={r['score']}")
 print(f"Excel : {xp}\nCSV   : {cp}\nJSON  : {jp}")
if __name__=="__main__":main()
