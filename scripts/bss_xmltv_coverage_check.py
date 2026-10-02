#!/usr/bin/env python3
"""
bss_xmltv_coverage_check.py

Henter BSS's xmltv.php (samme backend som get.php-M3U'en) og måler, for
HVER kanal i channel_priority.xlsx (Følg=X), hvor mange programmer og hvor
langt EPG-vindue der findes - matchet via Gruppe-nøgle og aliaser (da BSS's
channel-id'er typisk ligner "tv2.dk", "dr1.dk" osv., tæt på Gruppe-nøgle).

Samme outputformat som openepg_coverage_check.py, så de to kilder kan
sammenlignes direkte side om side.

VIGTIGT: URL'en indeholder dit brugernavn/password. Den gemmes IKKE i denne
fil - giv den som kommandolinje-argument.

Brug:
    python scripts/bss_xmltv_coverage_check.py "http://n2ip.tv:2095/xmltv.php?username=...&password=...&type=m3u_plus"
"""

import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CHANNEL_PRIORITY = ROOT / "data" / "channel_priority.xlsx"
OUTPUT_DIR = ROOT / "output_epgshare" / "bss_raw"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def load_followed_keys():
    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]

    kanal_col = headers.index("Kanal")
    key_col = headers.index("Gruppe-nøgle (intern)")
    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")

    rows = []
    for row in ws.iter_rows(min_row=2):
        if str(row[follow_col].value or "").strip().upper() != "X":
            continue
        kanal = row[kanal_col].value
        key = row[key_col].value
        if not kanal or not key:
            continue
        aliases_raw = row[aliases_col].value or ""
        aliases = [a.strip() for a in str(aliases_raw).split(",") if a.strip()]
        rows.append((str(kanal).strip(), str(key).strip().lower(), aliases))

    return rows


def normalize(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def parse_start(start_str):
    try:
        return datetime.strptime(start_str[:14], "%Y%m%d%H%M%S")
    except (ValueError, IndexError):
        return None


def main():
    if len(sys.argv) < 2:
        sys.exit(
            "❌ Mangler URL.\n"
            "Brug: python bss_xmltv_coverage_check.py \"http://n2ip.tv:2095/xmltv.php?...\""
        )

    url = sys.argv[1]
    cache_path = OUTPUT_DIR / "bss_epg.xml"

    if cache_path.exists():
        print(f"📄 Genbruger cachet fil: {cache_path}")
        content = cache_path.read_bytes()
    else:
        print("📥 Henter BSS xmltv ...")
        resp = requests.get(url, timeout=300)
        resp.raise_for_status()
        content = resp.content
        cache_path.write_bytes(content)
        print(f"✅ Downloadet og cachet: {len(content)//1024} KB")

    print("Parser XML ...")
    root = ET.fromstring(content)

    all_channel_ids = {ch.get("id", "") for ch in root.findall("channel")}

    prog_count = Counter()
    first_seen = {}
    last_seen = {}

    for p in root.findall("programme"):
        cid = p.get("channel", "")
        dt = parse_start(p.get("start", ""))
        if dt is None:
            continue
        prog_count[cid] += 1
        if cid not in first_seen or dt < first_seen[cid]:
            first_seen[cid] = dt
        if cid not in last_seen or dt > last_seen[cid]:
            last_seen[cid] = dt

    print(f"\nTotal antal <channel>: {len(all_channel_ids):,}")
    print(f"Total antal <programme>: {sum(prog_count.values()):,}\n")

    print("=== BSS xmltv dækning pr. kanal (Følg=X, matchet via Gruppe-nøgle) ===")
    print("-" * 100)

    followed = load_followed_keys()

    for kanonisk_navn, gruppe_noegle, aliases in followed:
        target_tokens = {normalize(gruppe_noegle)}
        for alias in aliases:
            target_tokens.add(normalize(alias.replace(".dk", "")))
        target_tokens.discard("")

        print(f"\n{kanonisk_navn}")

        matched_ids = []
        for cid in all_channel_ids:
            norm_id = normalize(cid.replace(".dk", ""))
            if not norm_id:
                continue
            if any(norm_id == tok for tok in target_tokens):
                matched_ids.append(cid)

        if not matched_ids:
            print("   (Intet matchende channel-id fundet i BSS)")
            continue

        for cid in sorted(matched_ids):
            count = prog_count.get(cid, 0)
            if count == 0:
                print(f"   [{cid}] -> <channel> findes, men 0 programmer")
                continue
            first_dt = first_seen.get(cid)
            last_dt = last_seen.get(cid)
            days = (last_dt - first_dt).days if first_dt and last_dt else 0
            print(
                f"   [{cid}] -> {count:,} programmer | "
                f"{first_dt} -> {last_dt} (~{days} dage)"
            )

    print("\n" + "-" * 100)
    print("Færdig. Sammenlign med openepg_coverage_check.py's output for samme kanaler.")


if __name__ == "__main__":
    main()
