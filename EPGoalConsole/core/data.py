import json, re
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from openpyxl import load_workbook
from .paths import paths, CONSOLE_ROOT

def load_json(path, default):
    try: return json.loads(path.read_text(encoding='utf-8-sig'))
    except Exception: return default

def split_ids(v):
    return [x.strip() for x in re.split(r'[,;\n]+', str(v or '')) if x.strip()]

def load_channels():
    p=paths()['priority']
    if not p.exists(): return []
    ws=load_workbook(p, read_only=True, data_only=True).worksheets[0]
    rows=list(ws.iter_rows(values_only=True))
    hi=next((i for i,r in enumerate(rows) if r and str(r[0] or '').strip()=='Kanal'), None)
    if hi is None: return []
    h=[str(x or '').strip() for x in rows[hi]]; out=[]
    overrides=load_json(CONSOLE_ROOT/'config'/'source_overrides.json', {})
    for row in rows[hi+1:]:
        if not row or not row[0]: continue
        rec={h[i]:(row[i] if i<len(row) else None) for i in range(len(h))}
        name=str(rec.get('Kanal') or '')
        ov=overrides.get(name,{})
        out.append({
          'channel':name,'follow':str(rec.get('Følg (X)') or ''),
          'streams':str(rec.get('BSS M3U stream-navne') or ''),
          'epgshare':str(rec.get("EPGShare ID'er") or ''),
          'openepg':str(rec.get("OpenEPG ID'er") or ''),
          'bss':str(rec.get("BSS XMLTV ID'er") or ''),
          'output':str(rec.get('Output/UHF tvg-id') or ''),
          'source_mode':ov.get('source','auto') if ov.get('enabled',True) else 'auto',
          'source_id':ov.get('source_channel_id',''),
        })
    return out

def epgoal_stats():
    p=paths()['epgoal']; counts=Counter(); names={}
    if not p.exists(): return {'channels':0,'programmes':0,'empty':0,'counts':{},'names':{}}
    try:
      for _,e in ET.iterparse(p, events=('end',)):
        tag=e.tag.rsplit('}',1)[-1]
        if tag=='channel':
          cid=e.get('id',''); names[cid]=next(((x.text or '').strip() for x in e.findall('./display-name') if (x.text or '').strip()),cid); e.clear()
        elif tag=='programme': counts[e.get('channel','')]+=1; e.clear()
    except ET.ParseError: pass
    return {'channels':len(names),'programmes':sum(counts.values()),'empty':sum(1 for x in names if counts[x]==0),'counts':dict(counts),'names':names}

def merge_stats():
    d=load_json(paths()['merge_log'], {})
    kept=d.get('kept_as_is',[]); supplied=d.get('supplemented',[]); missing=d.get('no_fallback_available',[])
    sources=Counter(['epgshare']*len(kept))
    for x in supplied:sources[x.get('kilde_brugt','ukendt')]+=1
    return {'kept':len(kept),'supplied':len(supplied),'missing':len(missing),'sources':dict(sources),'raw':d}

def current_snapshot():
    e=epgoal_stats();m=merge_stats()
    return {'created_at':datetime.now().isoformat(timespec='seconds'),'channels':e['channels'],'programmes':e['programmes'],'empty_channels':e['empty'],'kept_epgshare':m['kept'],'fallback_channels':m['supplied'],'missing_channels':m['missing'],'source_counts':m['sources']}
