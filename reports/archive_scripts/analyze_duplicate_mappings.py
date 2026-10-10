# analyze_duplicate_mappings.py

from openpyxl import load_workbook
from pathlib import Path
from collections import defaultdict
import re

XLSX = Path(r"C:\EPGoal\data\channel_priority_v2.xlsx")

def normalize(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

def split_ids(value):
    if value is None:
        return []

    return [
        x.strip()
        for x in re.split(r"[,;\n]+", str(value))
        if x.strip()
    ]

wb = load_workbook(XLSX, data_only=True)

if "Kanal-prioritering v2" in wb.sheetnames:
    ws = wb["Kanal-prioritering v2"]
else:
    ws = wb.active

headers = [c.value for c in ws[1]]

cols = {
    h: headers.index(h)
    for h in headers
    if h
}

entries = defaultdict(list)

for row_no, row in enumerate(
    ws.iter_rows(min_row=2),
    start=2
):

    kanal = str(
        row[cols["Kanal"]].value or ""
    ).strip()

    output = str(
        row[cols["Output/UHF tvg-id"]].value or ""
    ).strip()

    if not output:
        continue

    for source_col in (
        "EPGShare ID'er",
        "OpenEPG ID'er",
        "BSS XMLTV ID'er"
    ):

        raw_value = row[cols[source_col]].value

        for source_id in split_ids(raw_value):

            key = normalize(source_id)

            entries[key].append({
                "row": row_no,
                "kanal": kanal,
                "output": output,
                "source_col": source_col,
                "source_id": source_id,
            })

print()
print("=" * 80)
print("DUPLIKATE NORMALISEREDE NØGLER")
print("=" * 80)
print()

duplicates = 0

for key in sorted(entries):

    values = entries[key]

    outputs = {
        x["output"]
        for x in values
    }

    if len(outputs) <= 1:
        continue

    duplicates += 1

    print()
    print(f"NØGLE: {key}")
    print("-" * 80)

    for item in values:

        print(
            f"Række {item['row'\]:<4} | "
            f"Kanal={item['kanal']} | "
            f"Kilde={item['source_col']} | "
            f"ID={item['source_id']} | "
            f"Output={item['output']}"
        )

if duplicates == 0:
    print("Ingen dubletter fundet.")
else:
    print()
    print(f"Dubletter fundet: {duplicates}")