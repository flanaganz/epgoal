#!/usr/bin/env python3
"""
opdater_epgoal.py — SAMLET produktionsscript for EPGOAL.

Erstatter den spredte kæde af enkelt-scripts (epgshare_download.py,
epgshare_filter.py, epgshare_merge_all.py, normalize_uhf_channel_ids.py,
enrich_epg_epgshare.py, danish_backdrops_epgshare.py) med ÉT script der
kører alle trin i rækkefølge og ender med epgoal.xml committet og pushet
til GitHub.

Kilde-prioritet pr. Følg(X)-kanal (fra data/channel_priority_v2.xlsx):
    1. EPGShare01   - primær kilde. Bruges uændret hvis vinduet er
                       >= MIN_WINDOW_HOURS.
    2. OpenEPG      - fallback nr. 1 (6 kilder fra config.json).
    3. BSS XMLTV    - fallback nr. 2 (ultratv.one, kræver BSS_XMLTV_URL i .env).
Den kilde med længst vindue / flest programmer vinder, hvis EPGShare ikke
selv er god nok.

Herefter:
    4. Normalisering - alle channel-id'er omskrives til den autoritative
                        "Output/UHF tvg-id" fra channel_priority_v2.xlsx,
                        så normal/HD/FHD-streams i UHF peger på samme kanal.
    5. Sport-artwork  - samme logik/data som produktionen
                        (sport_channels.json, sport_categories.json osv.),
                        baseret på enrich_epg.py / enrich_epg_epgshare.py.
    6. Danske TMDb-backdrops - 1:1 portering af danish_backdrops.py /
                        danish_backdrops_epgshare.py: samme cache med
                        differentieret levetid, samme genopfrisk-liste,
                        samme godkendelsesfil-logik, samme manuelle
                        overrides, skriver KUN <icon> (ikke <backdrop>).
    7. Git commit + push (med samme fetch/rebase-retry-logik som
                        originalerne) af hele repoet (inkl.
                        output_epgshare/epgoal.xml).

TRIN DER IKKE ER MED HER (bevidst - kør separat ved behov):
    - build_master_channel_map_v2.py (genererer/opdaterer channel_priority_v2.xlsx)
    - analyze_channel_variants_v2.py (fejlsøgning af enkeltkanaler)
Disse er analyseværktøjer, ikke en del af den daglige produktionskørsel.

Kør:
    python scripts/opdater_epgoal.py
"""

from __future__ import annotations

import copy
import difflib
import gzip
import json
import os
import re
import subprocess
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path
from urllib.parse import quote
from xml.etree import ElementTree as ET

import requests
from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Stier og konfiguration
# ---------------------------------------------------------------------------

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"
OUTPUT_DIR = ROOT / "output_epgshare"
CONFIG_FILE = ROOT / "config.json"
CHANNEL_PRIORITY_V2 = DATA_DIR / "channel_priority_v2.xlsx"

EPGSHARE_URL = "https://epgshare01.online/epgshare01/epg_ripper_DK1.xml.gz"
EPGSHARE_RAW_FILE = OUTPUT_DIR / "epgshare_dk1.xml"
EPGSHARE_FILTERED_FILE = OUTPUT_DIR / "epgshare_filtered.xml"
EPGSHARE_MERGED_FILE = OUTPUT_DIR / "epgshare_merged.xml"
EPGOAL_FILE = OUTPUT_DIR / "epgoal.xml"

MERGE_LOG_FILE = DATA_DIR / "epgshare_merge_log.json"
NORMALIZE_LOG_FILE = DATA_DIR / "normalize_uhf_channel_ids_log.json"

CACHE_FILE = DATA_DIR / "cache.json"
TMDB_OVERRIDES_FILE = DATA_DIR / "overrides.json"
SPORT_CHANNELS_FILE = DATA_DIR / "sport_channels.json"
SPORT_CATEGORIES_FILE = DATA_DIR / "sport_categories.json"
SPORT_PROGRAM_OVERRIDES_FILE = DATA_DIR / "sport_program_overrides.json"
SPORT_SKIP_TITLES_FILE = DATA_DIR / "sport_skip_titles.json"
SPORT_PREFER_TMDB_TITLES_FILE = DATA_DIR / "sport_prefer_tmdb_titles.json"

# -- Danske TMDb-backdrops (1:1 med danish_backdrops.py / danish_backdrops_epgshare.py) --
DANISH_ARTWORK_CACHE_FILE = DATA_DIR / "danish_artwork_cache.json"
DANISH_ARTWORK_REVIEW_FILE = DATA_DIR / "danish_artwork_review.xlsx"
DANISH_BACKDROPS_RUN_LOG_FILE = DATA_DIR / "danish_backdrops_run_log.json"
MANUAL_ARTWORK_OVERRIDES_FILE = DATA_DIR / "manual_artwork_overrides.xlsx"
GENOPFRISK_TITLER_FILE = DATA_DIR / "genopfrisk_titler.txt"
MAX_RUN_LOG_ENTRIES = 200
NOT_FOUND_CACHE_MAX_AGE_DAYS_DEFAULT = 10

# Hvor mange timers EPG-vindue EPGShare MINDST skal have for en kanal, for at
# vi IKKE leder videre efter en bedre kilde (OpenEPG/BSS). Juster ved behov.
MIN_WINDOW_HOURS = 24

SPORT_IMAGE_BASE_URL = "https://raw.githubusercontent.com/flanaganz/epgoal/main/Sport/"
TMDB_BACKDROP_SIZE = "w1280"
CACHE_MAX_AGE_DAYS = 30

GIT_ENABLED = True
GIT_COMMIT_PREFIX = "Auto-opdater EPGOAL"

# ---------------------------------------------------------------------------
# .env
# ---------------------------------------------------------------------------

ENV_FILE = ROOT / ".env"


def load_env() -> None:
    if not ENV_FILE.exists():
        return
    for line in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))


load_env()
BSS_XMLTV_URL = os.environ.get("BSS_XMLTV_URL")
TMDB_API_KEY = os.environ.get("TMDB_API_KEY")

TMDB_BASE = "https://api.themoviedb.org/3"
IMAGE_BASE = "https://image.tmdb.org/t/p"
MATCH_SIMILARITY_MIN = 0.55
REQUEST_SLEEP_SECONDS = 0.05

INVISIBLE_CHARS_PATTERN = re.compile(r"[\u200B\u200C\u200D\u2060\uFEFF\u00AD]")
COLON_PATTERN = re.compile(r"\s*:\s*")

SESSION = requests.Session()


# ---------------------------------------------------------------------------
# Hjælpefunktioner (fælles)
# ---------------------------------------------------------------------------

def load_json(path: Path, default):
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            print(f"ADVARSEL: Kunne ikke læse {path}, bruger default", file=sys.stderr)
    return default


def save_json(path: Path, data) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True),
        encoding="utf-8",
    )


def normalize_id(value: str) -> str:
    """Normaliserer et kanal-id/alias til ren a-z0-9 - gør fx OpenEPG's
    'Kanal 4.dk', BSS's 'kanal4.dk' og EPGShare's 'Kanal.4.dk' sammenlignelige."""
    return re.sub(r"[^a-z0-9]", "", (value or "").lower())


def normalize_title(title: str) -> str:
    t = title.strip()
    t = unicodedata.normalize("NFKC", t)
    t = INVISIBLE_CHARS_PATTERN.sub(" ", t)
    t = re.sub(r"\s+", " ", t)
    return t.strip().lower()


def strip_colons(normalized_title: str) -> str:
    t = COLON_PATTERN.sub(" ", normalized_title)
    t = re.sub(r"\s+", " ", t)
    return t.strip()


def parse_start(start_str: str):
    try:
        return datetime.strptime(start_str[:14], "%Y%m%d%H%M%S")
    except (ValueError, IndexError, TypeError):
        return None


def is_full_url(value: str) -> bool:
    return value.strip().lower().startswith(("http://", "https://"))


def section(title: str) -> None:
    print()
    print("=" * 70)
    print(title)
    print("=" * 70)


# ---------------------------------------------------------------------------
# TRIN 1: Download EPGShare01
# ---------------------------------------------------------------------------

def step1_download_epgshare() -> None:
    section("TRIN 1/7: Henter EPGShare01 DK1")

    print(f"Henter {EPGSHARE_URL} ...")
    resp = SESSION.get(EPGSHARE_URL, timeout=300)
    resp.raise_for_status()
    print(f"Downloadet {len(resp.content) // 1024} KB (gzip)")

    xml_data = gzip.decompress(resp.content)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    EPGSHARE_RAW_FILE.write_bytes(xml_data)

    print(f"Gemt: {EPGSHARE_RAW_FILE} ({len(xml_data) // 1024} KB udpakket)")


# ---------------------------------------------------------------------------
# TRIN 2: Filtrer til Følg(X)-kanaler
# ---------------------------------------------------------------------------

def load_followed_channels_for_filter() -> set[str]:
    wb = load_workbook(CHANNEL_PRIORITY_V2, data_only=True)
    ws = wb["Kanal-prioritering v2"] if "Kanal-prioritering v2" in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]

    kanal_col = headers.index("Kanal")
    aliases_col = headers.index("Sammenlagte kanal-ID'er (gammel reference)")
    follow_col = headers.index("Følg (X)")

    wanted = set()
    for row in ws.iter_rows(min_row=2):
        if str(row[follow_col].value or "").strip().upper() != "X":
            continue

        kanal = row[kanal_col].value
        if kanal:
            wanted.add(str(kanal).strip().lower())

        aliases = row[aliases_col].value
        if aliases:
            for alias in str(aliases).split(","):
                alias = alias.strip()
                if alias:
                    wanted.add(alias.lower())

    return wanted


def step2_filter_epgshare() -> None:
    section("TRIN 2/7: Filtrerer EPGShare til dine kanaler")

    wanted_channels = load_followed_channels_for_filter()
    print(f"Whitelist kanaler: {len(wanted_channels)}")

    tree = ET.parse(EPGSHARE_RAW_FILE)
    root = tree.getroot()

    new_root = ET.Element("tv")
    kept_channel_ids: set[str] = set()

    for channel in root.findall("channel"):
        channel_id = (channel.get("id") or "").strip()
        display_names = [
            (dn.text or "").strip() for dn in channel.findall("display-name")
        ]
        candidates = [channel_id] + display_names

        if any(value.lower() in wanted_channels for value in candidates if value):
            new_root.append(channel)
            kept_channel_ids.add(channel_id)

    programme_count = 0
    for programme in root.findall("programme"):
        channel_id = (programme.get("channel") or "").strip()
        if channel_id in kept_channel_ids:
            for tag in ("icon", "backdrop"):
                for old in programme.findall(tag):
                    programme.remove(old)
            new_root.append(programme)
            programme_count += 1

    ET.ElementTree(new_root).write(
        EPGSHARE_FILTERED_FILE, encoding="utf-8", xml_declaration=True
    )

    print(f"Kanaler beholdt   : {len(kept_channel_ids):,}")
    print(f"Programmer beholdt: {programme_count:,}")
    print(f"Output            : {EPGSHARE_FILTERED_FILE}")


# ---------------------------------------------------------------------------
# TRIN 3: Merge EPGShare + OpenEPG + BSS (fallback pr. kanal)
# ---------------------------------------------------------------------------

def load_followed_rows_for_merge() -> list[tuple[str, set[str]]]:
    wb = load_workbook(CHANNEL_PRIORITY_V2, data_only=True)
    ws = wb["Kanal-prioritering v2"] if "Kanal-prioritering v2" in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]

    kanal_col = headers.index("Kanal")
    key_col = headers.index("Gruppe-nøgle (intern)")
    aliases_col = headers.index("Sammenlagte kanal-ID'er (gammel reference)")
    follow_col = headers.index("Følg (X)")

    rows = []
    for row in ws.iter_rows(min_row=2):
        if str(row[follow_col].value or "").strip().upper() != "X":
            continue
        kanal = row[kanal_col].value
        if not kanal:
            continue

        tokens = set()
        key = row[key_col].value
        if key:
            tokens.add(normalize_id(str(key)))

        aliases_raw = row[aliases_col].value or ""
        for alias in str(aliases_raw).split(","):
            alias = alias.strip()
            if alias:
                tokens.add(normalize_id(alias))

        tokens.add(normalize_id(str(kanal)))
        tokens.discard("")
        rows.append((str(kanal).strip(), tokens))

    return rows


def build_channel_stats(programmes_by_channel: dict) -> dict:
    stats = {}
    for cid, plist in programmes_by_channel.items():
        dates = [parse_start(p.get("start", "")) for p in plist]
        dates = [d for d in dates if d is not None]
        count = len(dates)
        span = (max(dates) - min(dates)).total_seconds() / 3600 if dates else 0
        stats[cid] = {"count": count, "span_hours": span}
    return stats


def fetch_openepg() -> dict:
    config = load_json(CONFIG_FILE, {})
    sources = config.get("sources", [])
    if not sources:
        print("ADVARSEL: Ingen kilder fundet i config.json - springer OpenEPG over.")
        return {}

    programmes_by_channel: dict[str, list] = {}
    for source in sources:
        name = source["name"]
        url = source["url"]
        print(f"Henter OpenEPG-kilde {name} ...")
        try:
            resp = SESSION.get(url, timeout=90)
            resp.raise_for_status()
        except requests.RequestException as exc:
            print(f"ADVARSEL: Kunne ikke hente {name}: {exc}", file=sys.stderr)
            continue
        root = ET.fromstring(resp.content)
        for p in root.findall("programme"):
            cid = p.get("channel", "")
            programmes_by_channel.setdefault(cid, []).append(p)

    total = sum(len(v) for v in programmes_by_channel.values())
    print(f"OpenEPG: {total:,} programmer hentet i alt.")
    return programmes_by_channel


def fetch_bss() -> dict:
    if not BSS_XMLTV_URL:
        print("INFO: BSS_XMLTV_URL er ikke sat i .env - springer BSS over.")
        return {}

    print("Henter BSS (ultratv.one) ...")
    try:
        resp = SESSION.get(BSS_XMLTV_URL, timeout=300)
        resp.raise_for_status()
    except requests.RequestException as exc:
        print(f"ADVARSEL: Kunne ikke hente BSS: {exc}", file=sys.stderr)
        return {}

    root = ET.fromstring(resp.content)
    programmes_by_channel: dict[str, list] = {}
    for p in root.findall("programme"):
        cid = p.get("channel", "")
        programmes_by_channel.setdefault(cid, []).append(p)

    total = sum(len(v) for v in programmes_by_channel.values())
    print(f"BSS: {total:,} programmer hentet i alt.")
    return programmes_by_channel


def find_best_candidate(tokens: set, programmes_by_channel: dict, stats: dict):
    best_id = None
    best_stat = {"count": 0, "span_hours": 0}
    for cid in programmes_by_channel:
        if normalize_id(cid) in tokens:
            cand_stat = stats.get(cid, {"count": 0, "span_hours": 0})
            if (cand_stat["count"], cand_stat["span_hours"]) > (best_stat["count"], best_stat["span_hours"]):
                best_id = cid
                best_stat = cand_stat
    return best_id, best_stat


def strip_artwork(programme_el) -> None:
    for tag in ("icon", "backdrop"):
        for old in programme_el.findall(tag):
            programme_el.remove(old)


def step3_merge_sources() -> None:
    section("TRIN 3/7: Samler EPGShare + OpenEPG + BSS (fallback pr. kanal)")

    tree = ET.parse(EPGSHARE_FILTERED_FILE)
    root = tree.getroot()

    epgshare_channel_ids = {ch.get("id", "") for ch in root.findall("channel")}
    epgshare_programmes_by_channel: dict[str, list] = {}
    for p in root.findall("programme"):
        cid = p.get("channel", "")
        epgshare_programmes_by_channel.setdefault(cid, []).append(p)
    epgshare_stats = build_channel_stats(epgshare_programmes_by_channel)

    openepg_programmes_by_channel = fetch_openepg()
    openepg_stats = build_channel_stats(openepg_programmes_by_channel)

    bss_programmes_by_channel = fetch_bss()
    bss_stats = build_channel_stats(bss_programmes_by_channel)

    followed_rows = load_followed_rows_for_merge()

    report = {"kept_as_is": [], "supplemented": [], "no_fallback_available": []}

    for kanonisk_navn, tokens in followed_rows:
        matched_epgshare_ids = [
            cid for cid in epgshare_channel_ids if normalize_id(cid) in tokens
        ]
        epgshare_count = sum(
            epgshare_stats.get(c, {}).get("count", 0) for c in matched_epgshare_ids
        )
        epgshare_span = max(
            (epgshare_stats.get(c, {}).get("span_hours", 0) for c in matched_epgshare_ids),
            default=0,
        )

        if epgshare_count > 0 and epgshare_span >= MIN_WINDOW_HOURS:
            report["kept_as_is"].append({
                "kanal": kanonisk_navn, "kilde": "epgshare",
                "programmer": epgshare_count, "vindue_timer": round(epgshare_span, 1),
            })
            continue

        openepg_id, openepg_stat = find_best_candidate(tokens, openepg_programmes_by_channel, openepg_stats)
        bss_id, bss_stat = find_best_candidate(tokens, bss_programmes_by_channel, bss_stats)

        candidates = []
        if openepg_id and openepg_stat["count"] > 0:
            candidates.append(("openepg", openepg_id, openepg_stat, openepg_programmes_by_channel))
        if bss_id and bss_stat["count"] > 0:
            candidates.append(("bss", bss_id, bss_stat, bss_programmes_by_channel))

        if not candidates:
            report["no_fallback_available"].append({
                "kanal": kanonisk_navn, "epgshare_programmer": epgshare_count,
                "epgshare_vindue_timer": round(epgshare_span, 1),
            })
            continue

        candidates.sort(key=lambda c: (c[2]["span_hours"], c[2]["count"]), reverse=True)
        best_source, best_id, best_stat, best_programmes_map = candidates[0]

        if epgshare_span >= best_stat["span_hours"] and epgshare_count > 0:
            report["kept_as_is"].append({
                "kanal": kanonisk_navn,
                "kilde": "epgshare (bedst tilgængelige, selvom < min. vindue)",
                "programmer": epgshare_count, "vindue_timer": round(epgshare_span, 1),
            })
            continue

        if matched_epgshare_ids:
            primary_id = max(
                matched_epgshare_ids, key=lambda c: epgshare_stats.get(c, {}).get("count", 0)
            )
        else:
            primary_id = f"{normalize_id(kanonisk_navn)}.dk"
            new_channel_el = ET.Element("channel", {"id": primary_id})
            dn = ET.SubElement(new_channel_el, "display-name")
            dn.text = kanonisk_navn
            root.append(new_channel_el)

        for cid in matched_epgshare_ids:
            for p in epgshare_programmes_by_channel.get(cid, []):
                root.remove(p)

        inserted = 0
        for p in best_programmes_map.get(best_id, []):
            clone = copy.deepcopy(p)
            clone.set("channel", primary_id)
            strip_artwork(clone)
            root.append(clone)
            inserted += 1

        report["supplemented"].append({
            "kanal": kanonisk_navn,
            "epgshare_programmer_foer": epgshare_count,
            "epgshare_vindue_timer_foer": round(epgshare_span, 1),
            "kilde_brugt": best_source,
            "kilde_channel_id": best_id,
            "programmer_indsat": inserted,
            "vindue_timer_efter": round(best_stat["span_hours"], 1),
            "primaer_channel_id": primary_id,
        })

    tree.write(EPGSHARE_MERGED_FILE, encoding="utf-8", xml_declaration=True)
    save_json(MERGE_LOG_FILE, report)

    print(f"Beholdt EPGShare uændret : {len(report['kept_as_is']):,} kanaler")
    print(f"Suppleret med OpenEPG/BSS: {len(report['supplemented']):,} kanaler")
    print(f"Intet fallback muligt    : {len(report['no_fallback_available']):,} kanaler")
    print(f"Output: {EPGSHARE_MERGED_FILE}")


# ---------------------------------------------------------------------------
# TRIN 4: Normaliser til Output/UHF tvg-id
# ---------------------------------------------------------------------------

def load_uhf_mapping() -> dict[str, str]:
    wb = load_workbook(CHANNEL_PRIORITY_V2, data_only=True)
    ws = wb["Kanal-prioritering v2"] if "Kanal-prioritering v2" in wb.sheetnames else wb.active
    headers = [cell.value for cell in ws[1]]

    epgshare_col = headers.index("EPGShare ID'er")
    openepg_col = headers.index("OpenEPG ID'er")
    bss_col = headers.index("BSS XMLTV ID'er")
    output_col = headers.index("Output/UHF tvg-id")

    mapping: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2):
        output_id = str(row[output_col].value or "").strip()
        if not output_id:
            continue

        source_values = []
        for idx in (epgshare_col, openepg_col, bss_col):
            raw = row[idx].value
            if not raw:
                continue
            source_values.extend(
                part.strip() for part in str(raw).split(",") if part.strip()
            )

        for src in source_values:
            mapping[normalize_id(src)] = output_id

    return mapping


def step4_normalize_uhf_ids() -> None:
    section("TRIN 4/7: Normaliserer channel-id'er til Output/UHF tvg-id")

    mapping = load_uhf_mapping()

    tree = ET.parse(EPGSHARE_MERGED_FILE)
    root = tree.getroot()

    remapped_channels: dict[str, str] = {}
    new_channels: dict[str, ET.Element] = {}
    new_programmes = []

    for channel in root.findall("channel"):
        old_id = channel.get("id") or ""
        new_id = mapping.get(normalize_id(old_id), old_id)
        channel.set("id", new_id)

        if new_id not in new_channels:
            new_channels[new_id] = channel

        remapped_channels[old_id] = new_id

    root[:] = [c for c in root if c.tag != "channel"]
    for channel in new_channels.values():
        root.append(channel)

    seen_prog = set()
    for programme in root.findall("programme"):
        old_channel = programme.get("channel") or ""
        new_channel = mapping.get(
            normalize_id(old_channel),
            remapped_channels.get(old_channel, old_channel),
        )
        programme.set("channel", new_channel)

        key = (new_channel, programme.get("start", ""), programme.findtext("title", ""))
        if key in seen_prog:
            continue
        seen_prog.add(key)
        new_programmes.append(programme)

    root[:] = [n for n in root if n.tag != "programme"]
    for programme in new_programmes:
        root.append(programme)

    tree.write(EPGOAL_FILE, encoding="utf-8", xml_declaration=True)
    save_json(NORMALIZE_LOG_FILE, remapped_channels)

    print(f"Mappings  : {len(mapping):,}")
    print(f"Kanaler   : {len(new_channels):,}")
    print(f"Programmer: {len(new_programmes):,}")
    print(f"Output    : {EPGOAL_FILE}")


# ---------------------------------------------------------------------------
# TRIN 5: Sport-berigelse (på epgoal.xml, in-place)
# ---------------------------------------------------------------------------

class SportMatcher:
    def __init__(self, image_base_url: str):
        self.image_base_url = image_base_url.rstrip("/") + "/"

        channels_raw = load_json(SPORT_CHANNELS_FILE, {})
        if isinstance(channels_raw, list):
            self.exclude_patterns: list[str] = []
            self.channels = channels_raw
        else:
            self.exclude_patterns = [e["match"].lower() for e in channels_raw.get("exclude", [])]
            self.channels = channels_raw.get("channels", [])

        self.categories = load_json(SPORT_CATEGORIES_FILE, [])
        self.program_overrides = load_json(SPORT_PROGRAM_OVERRIDES_FILE, {})
        self.skip_titles = set(load_json(SPORT_SKIP_TITLES_FILE, []))

        self.program_overrides_colon_stripped: dict[str, dict] = {}
        for key, value in self.program_overrides.items():
            stripped = strip_colons(key.strip().lower())
            if stripped not in self.program_overrides_colon_stripped:
                self.program_overrides_colon_stripped[stripped] = value

        prefer_tmdb_raw = load_json(SPORT_PREFER_TMDB_TITLES_FILE, {"titles": []})
        self.prefer_tmdb_titles = set(
            t.strip().lower() for t in prefer_tmdb_raw.get("titles", [])
        )

        self.prefix_lookup: dict[str, dict] = {}
        for cat in self.categories:
            for prefix in cat.get("prefix", []) or []:
                self.prefix_lookup[prefix.strip().lower()] = cat

        priority_keywords: list[tuple[str, dict]] = []
        normal_keywords: list[tuple[str, dict]] = []
        for cat in self.categories:
            target = priority_keywords if cat.get("priority") else normal_keywords
            for keyword in cat.get("keywords", []) or []:
                target.append((keyword.strip().lower(), cat))
        normal_keywords.sort(key=lambda pair: len(pair[0]), reverse=True)
        self.keyword_lookup: list[tuple[str, dict]] = priority_keywords + normal_keywords

    def match_channel(self, channel_id: str) -> dict | None:
        low = (channel_id or "").lower()
        low_normalized = low.replace(".", " ").replace("_", " ")
        low_normalized = re.sub(r"\s+", " ", low_normalized).strip()

        for pattern in self.exclude_patterns:
            if pattern in low or pattern in low_normalized:
                return None
        for entry in self.channels:
            pattern = entry["match"].lower()
            if pattern in low or pattern in low_normalized:
                return entry
        return None

    def _image_urls(self, backdrop_value, poster_value) -> dict:
        def build(value):
            if not value:
                return None
            if is_full_url(value):
                return value.strip()
            return self.image_base_url + quote(value)
        return {"backdrop": build(backdrop_value), "poster": build(poster_value)}

    def _real_match(self, backdrop_value, poster_value):
        if not backdrop_value and not poster_value:
            return None
        return self._image_urls(backdrop_value, poster_value)

    def prefers_tmdb(self, raw_title: str) -> bool:
        return normalize_title(raw_title) in self.prefer_tmdb_titles

    def _lookup_override(self, norm: str):
        override = self.program_overrides.get(norm)
        if override:
            match = self._real_match(override.get("backdrop"), override.get("poster"))
            if match:
                return match
        stripped = strip_colons(norm)
        override = self.program_overrides_colon_stripped.get(stripped)
        if override:
            match = self._real_match(override.get("backdrop"), override.get("poster"))
            if match:
                return match
        return None

    def resolve_local(self, raw_title: str):
        norm = normalize_title(raw_title)
        if norm in self.skip_titles:
            return {"skip": True}
        if norm in self.prefer_tmdb_titles:
            return None
        override_match = self._lookup_override(norm)
        if override_match:
            return override_match
        for keyword, cat in self.keyword_lookup:
            if keyword in norm:
                match = self._real_match(cat.get("backdrop"), cat.get("poster"))
                if match:
                    return match
        prefix = norm.split(":", 1)[0].strip()
        cat = self.prefix_lookup.get(prefix)
        if cat:
            match = self._real_match(cat.get("backdrop"), cat.get("poster"))
            if match:
                return match
        return None


def sport_tmdb_search(title: str):
    resp = SESSION.get(
        f"{TMDB_BASE}/search/multi",
        params={"api_key": TMDB_API_KEY, "query": title, "language": "da-DK", "include_adult": "false"},
        timeout=15,
    )
    resp.raise_for_status()
    results = [r for r in resp.json().get("results", []) if r.get("media_type") in ("tv", "movie")]
    if not results:
        return None

    def score(r):
        name = r.get("name") or r.get("title") or ""
        similarity = difflib.SequenceMatcher(None, name.lower(), title.lower()).ratio()
        popularity_bonus = min(r.get("popularity", 0), 50) / 50
        return similarity * 2 + popularity_bonus

    results.sort(key=score, reverse=True)
    best = results[0]
    best_name = best.get("name") or best.get("title") or ""
    similarity = difflib.SequenceMatcher(None, best_name.lower(), title.lower()).ratio()
    if similarity < MATCH_SIMILARITY_MIN:
        return None
    return best["media_type"], best["id"]


def sport_tmdb_images(media_type: str, tmdb_id: int):
    """Sport-fallback: bredere sprogvalg (da,en,null) end den danske
    backdrop-logik i trin 6, da sport ofte ikke har dansksprogede billeder."""
    resp = SESSION.get(
        f"{TMDB_BASE}/{media_type}/{tmdb_id}/images",
        params={"api_key": TMDB_API_KEY, "include_image_language": "da,en,null"},
        timeout=15,
    )
    resp.raise_for_status()
    data = resp.json()
    backdrops = data.get("backdrops", [])
    if not backdrops:
        return None
    lang_rank = {"da": 0, "en": 1, None: 2}
    backdrops = sorted(
        backdrops, key=lambda b: (lang_rank.get(b.get("iso_639_1"), 3), -b.get("vote_average", 0))
    )
    return backdrops[0]["file_path"]


def resolve_sport_tmdb_artwork(raw_title: str, overrides: dict, cache: dict, cache_max_age_days: int,
                                backdrop_size: str):
    key = normalize_title(raw_title)
    override = overrides.get(key)
    if override:
        if override.get("backdrop_url") or override.get("poster_url"):
            return {"backdrop": override.get("backdrop_url")}, False
        if override.get("tmdb_id") and override.get("media_type"):
            try:
                b_path = sport_tmdb_images(override["media_type"], override["tmdb_id"])
                return {"backdrop": f"{IMAGE_BASE}/{backdrop_size}{b_path}" if b_path else None}, False
            except requests.RequestException as exc:
                print(f"TMDb-fejl (override) for '{raw_title}': {exc}", file=sys.stderr)

    cached = cache.get(key)
    if cached is not None and (time.time() - cached.get("ts", 0)) / 86400 < cache_max_age_days:
        return {"backdrop": cached.get("backdrop")}, True

    backdrop_url = None
    try:
        match = sport_tmdb_search(raw_title)
        if match:
            media_type, tmdb_id = match
            b_path = sport_tmdb_images(media_type, tmdb_id)
            if b_path:
                backdrop_url = f"{IMAGE_BASE}/{backdrop_size}{b_path}"
    except requests.RequestException as exc:
        print(f"TMDb-fejl for '{raw_title}': {exc}", file=sys.stderr)

    cache[key] = {"backdrop": backdrop_url, "ts": time.time()}
    return {"backdrop": backdrop_url}, False


def set_artwork(programme, backdrop_url, poster_url):
    for old in programme.findall("icon"):
        programme.remove(old)
    for old in programme.findall("backdrop"):
        programme.remove(old)
    if poster_url:
        ET.SubElement(programme, "icon").set("src", poster_url)
    if backdrop_url:
        ET.SubElement(programme, "backdrop").set("src", backdrop_url)


def set_backdrop_only(programme, backdrop_url) -> None:
    set_artwork(programme, backdrop_url, backdrop_url)


def clear_artwork(programme) -> None:
    for old in programme.findall("icon"):
        programme.remove(old)
    for old in programme.findall("backdrop"):
        programme.remove(old)


def step5_sport_enrichment() -> None:
    section("TRIN 5/7: Sport-berigelse (epgoal.xml)")

    matcher = SportMatcher(SPORT_IMAGE_BASE_URL)
    tmdb_overrides = load_json(TMDB_OVERRIDES_FILE, {})
    cache = load_json(CACHE_FILE, {})
    cache_before = len(cache)

    sport_tmdb_fallback_enabled = bool(TMDB_API_KEY)

    tree = ET.parse(EPGOAL_FILE)
    root = tree.getroot()

    channel_role: dict[str, dict | None] = {}
    for ch in root.findall("channel"):
        cid = ch.get("id", "")
        channel_role[cid] = matcher.match_channel(cid)

    stats = {
        "programmes": 0, "sport_matched": 0, "sport_tmdb_matched": 0,
        "sport_tmdb_cache_hit": 0, "sport_tmdb_fresh_call": 0,
        "sport_defaulted": 0, "sport_skipped": 0, "sport_no_image_yet": 0,
    }
    sport_cache_this_run: dict[str, dict | None] = {}
    tmdb_cache_this_run: dict[str, tuple[dict, bool]] = {}

    for programme in root.findall("programme"):
        stats["programmes"] += 1
        title_el = programme.find("title")
        if title_el is None or not title_el.text:
            continue
        title = title_el.text
        chan_id = programme.get("channel", "")

        role_entry = channel_role.get(chan_id)
        if role_entry is None:
            continue

        role = role_entry.get("role")

        if title not in sport_cache_this_run:
            sport_cache_this_run[title] = matcher.resolve_local(title)
        result = sport_cache_this_run[title]

        if result and result.get("skip"):
            clear_artwork(programme)
            stats["sport_skipped"] += 1
            continue

        if result:
            set_artwork(programme, result.get("backdrop"), result.get("poster"))
            stats["sport_matched"] += 1
            continue

        should_try_tmdb = sport_tmdb_fallback_enabled and (
            role == "always_sport" or (role == "partial_sport" and matcher.prefers_tmdb(title))
        )

        if should_try_tmdb:
            if title not in tmdb_cache_this_run:
                art, from_cache = resolve_sport_tmdb_artwork(
                    title, tmdb_overrides, cache, CACHE_MAX_AGE_DAYS, TMDB_BACKDROP_SIZE
                )
                tmdb_cache_this_run[title] = (art, from_cache)
                if not from_cache:
                    time.sleep(REQUEST_SLEEP_SECONDS)
            art, from_cache = tmdb_cache_this_run[title]

            if art.get("backdrop"):
                set_backdrop_only(programme, art.get("backdrop"))
                stats["sport_tmdb_matched"] += 1
                if from_cache:
                    stats["sport_tmdb_cache_hit"] += 1
                else:
                    stats["sport_tmdb_fresh_call"] += 1
                continue

        if role == "always_sport":
            fallback_backdrop = role_entry.get("default_backdrop")
            fallback_poster = role_entry.get("default_poster")
            if fallback_backdrop or fallback_poster:
                urls = matcher._image_urls(fallback_backdrop, fallback_poster)
                set_artwork(programme, urls["backdrop"], urls["poster"])
                stats["sport_defaulted"] += 1
            else:
                stats["sport_no_image_yet"] += 1

    tree.write(EPGOAL_FILE, encoding="utf-8", xml_declaration=True)
    save_json(CACHE_FILE, cache)

    print(f"Programmer i alt        : {stats['programmes']:,}")
    print(f"Sport - specifikt match : {stats['sport_matched']:,}")
    print(
        f"Sport - TMDb-match      : {stats['sport_tmdb_matched']:,} "
        f"(cache: {stats['sport_tmdb_cache_hit']:,} / friske: {stats['sport_tmdb_fresh_call']:,})"
    )
    print(f"Sport - kanal-fallback  : {stats['sport_defaulted']:,}")
    print(f"Sport - sprunget over   : {stats['sport_skipped']:,}")
    print(f"Sport - mangler billede : {stats['sport_no_image_yet']:,}")
    print(f"Cache voksede fra {cache_before:,} til {len(cache):,} unikke titler")


# ---------------------------------------------------------------------------
# TRIN 6: Danske TMDb-backdrops — 1:1 portering af danish_backdrops.py /
#         danish_backdrops_epgshare.py, kørt direkte på epgoal.xml.
# ---------------------------------------------------------------------------

def load_refresh_titles(path: Path) -> list[str]:
    if not path.exists():
        return []
    titles: list[str] = []
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        titles.append(line)
    return titles


def save_refresh_titles(path: Path, titles: list[str]) -> None:
    if not titles:
        path.write_text("", encoding="utf-8")
        return
    path.write_text("\n".join(titles) + "\n", encoding="utf-8")


def load_manual_overrides(path: Path):
    """Returnerer (index, display_titles).
    index: normaliseret titel -> liste af {'channel', 'backdrop_url'}."""
    index: dict[str, list[dict]] = {}
    display_titles: dict[str, str] = {}
    if not path.exists():
        return index, display_titles

    wb = load_workbook(path, data_only=True)
    ws = wb["Manuelle overrides"] if "Manuelle overrides" in wb.sheetnames else wb.active
    headers = [c.value for c in ws[1]]
    try:
        title_col = headers.index("Titel (som i EPG)")
        channel_col = headers.index("Kanal (valgfri)")
        url_col = headers.index("Backdrop URL")
    except ValueError:
        print(f"ADVARSEL: {path.name} mangler forventede kolonner - ingen manuelle overrides indlæst.", file=sys.stderr)
        return index, display_titles

    for row in ws.iter_rows(min_row=2):
        title_val = row[title_col].value
        url_val = row[url_col].value
        if not title_val or not url_val:
            continue
        title = str(title_val).strip()
        url = str(url_val).strip()
        if not title or not url or title.upper().startswith("EKSEMPEL"):
            continue
        norm = normalize_title(title)
        display_titles[norm] = title
        channel_val = row[channel_col].value
        channel = str(channel_val).strip().lower() if channel_val else ""
        index.setdefault(norm, []).append({"channel": channel, "backdrop_url": url})

    return index, display_titles


def resolve_manual_override(title: str, channel_id: str, manual_index: dict) -> str | None:
    norm = normalize_title(title)
    entries = manual_index.get(norm)
    if not entries:
        return None
    channel_low = (channel_id or "").lower()
    for e in entries:
        if e["channel"] and e["channel"] in channel_low:
            return e["backdrop_url"]
    for e in entries:
        if not e["channel"]:
            return e["backdrop_url"]
    return None


def load_approved_keys(review_path: Path) -> set[str] | None:
    """Titler markeret 'Godkendt (X)' - respekterer 'Ignorer (X)' som
    sikkerhedsspærre (ignorer vinder altid over godkendt)."""
    if not review_path.exists():
        return None
    wb = load_workbook(review_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    try:
        key_col = headers.index("Nøgle (intern)")
        godkendt_col = headers.index("Godkendt (X)")
    except ValueError:
        return set()
    ignorer_col = headers.index("Ignorer (X)") if "Ignorer (X)" in headers else None

    approved: set[str] = set()
    for row in ws.iter_rows(min_row=2):
        key_val = row[key_col].value
        godkendt_val = row[godkendt_col].value
        ignorer_val = row[ignorer_col].value if ignorer_col is not None else None
        if str(ignorer_val or "").strip().upper() == "X":
            continue
        if key_val and str(godkendt_val).strip().upper() == "X":
            approved.add(str(key_val).strip())
    return approved


def load_undecided_review_keys(review_path: Path) -> dict[str, str]:
    """Rækker uden X i både Godkendt og Ignorer; opfriskes frisk."""
    if not review_path.exists():
        return {}
    wb = load_workbook(review_path, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    try:
        key_col = headers.index("Nøgle (intern)")
        approved_col = headers.index("Godkendt (X)")
    except ValueError:
        return {}
    ignored_col = headers.index("Ignorer (X)") if "Ignorer (X)" in headers else None
    title_col = headers.index("Titel") if "Titel" in headers else None

    result = {}
    for row in ws.iter_rows(min_row=2):
        key = row[key_col].value
        if not key:
            continue
        approved = str(row[approved_col].value or "").strip().upper() == "X"
        ignored = (
            str(row[ignored_col].value or "").strip().upper() == "X"
            if ignored_col is not None else False
        )
        if approved or ignored:
            continue
        norm = normalize_title(str(key))
        title = row[title_col].value if title_col is not None else key
        result[norm] = str(title).strip()
    return result


def danish_tmdb_search(title: str):
    resp = SESSION.get(
        f"{TMDB_BASE}/search/multi",
        params={"api_key": TMDB_API_KEY, "query": title, "language": "da-DK", "include_adult": "false"},
        timeout=15,
    )
    resp.raise_for_status()
    results = [r for r in resp.json().get("results", []) if r.get("media_type") in ("tv", "movie")]
    if not results:
        return None

    def score(r):
        name = r.get("name") or r.get("title") or ""
        similarity = difflib.SequenceMatcher(None, name.lower(), title.lower()).ratio()
        popularity_bonus = min(r.get("popularity", 0), 50) / 50
        return similarity * 2 + popularity_bonus

    results.sort(key=score, reverse=True)
    best = results[0]
    best_name = best.get("name") or best.get("title") or ""
    similarity = difflib.SequenceMatcher(None, best_name.lower(), title.lower()).ratio()
    if similarity < MATCH_SIMILARITY_MIN:
        return None
    return best["media_type"], best["id"]


def tmdb_danish_backdrop(media_type: str, tmdb_id: int) -> str | None:
    """Henter bedste DANSKE backdrop (kun include_image_language=da) - IKKE
    samme bredde som sport-fallback'en i trin 5. Posters håndteres ikke."""
    resp = SESSION.get(
        f"{TMDB_BASE}/{media_type}/{tmdb_id}/images",
        params={"api_key": TMDB_API_KEY, "include_image_language": "da"},
        timeout=15,
    )
    resp.raise_for_status()
    data = resp.json()
    backdrops = data.get("backdrops", [])
    if not backdrops:
        return None
    backdrops = sorted(backdrops, key=lambda b: -b.get("vote_average", 0))
    return backdrops[0]["file_path"]


def resolve_danish_artwork(raw_title: str, cache: dict, cache_max_age_days: int,
                            not_found_cache_max_age_days: int,
                            backdrop_size: str, force_refresh: bool = False):
    """Returnerer (backdrop_url eller None, from_cache). Differentieret
    cache-levetid: FUND holder cache_max_age_days, 'IKKE fundet' holder kun
    not_found_cache_max_age_days (så nye TMDb-billeder opdages hurtigt)."""
    key = normalize_title(raw_title)
    if not force_refresh:
        cached = cache.get(key)
        if cached is not None:
            age_days = (time.time() - cached.get("ts", 0)) / 86400
            had_backdrop = bool(cached.get("backdrop"))
            max_age = cache_max_age_days if had_backdrop else not_found_cache_max_age_days
            if age_days < max_age:
                return cached.get("backdrop"), True

    backdrop_url = None
    try:
        match = danish_tmdb_search(raw_title)
        if match:
            media_type, tmdb_id = match
            b_path = tmdb_danish_backdrop(media_type, tmdb_id)
            if b_path:
                backdrop_url = f"{IMAGE_BASE}/{backdrop_size}{b_path}"
    except requests.RequestException as exc:
        print(f"   TMDb-fejl for '{raw_title}': {exc}", file=sys.stderr)
        return None, False

    cache[key] = {"title": raw_title, "backdrop": backdrop_url, "ts": time.time()}
    return backdrop_url, False


def append_run_log(log_path: Path, stats: dict, cache_size_before: int, cache_size_after: int,
                    approved_count: int, review_exists: bool, unique_found: int, unique_pending: int,
                    manual_defined_count: int, manual_titles_matched_display: list[str],
                    manual_titles_unmatched_display: list[str]) -> None:
    history = load_json(log_path, [])
    if not isinstance(history, list):
        history = []
    history.append({
        "timestamp": time.time(),
        "date_str": time.strftime("%Y-%m-%d %H:%M:%S"),
        "totals": stats,
        "cache_size_before": cache_size_before,
        "cache_size_after": cache_size_after,
        "approved_count": approved_count,
        "unique_found": unique_found,
        "unique_pending": unique_pending,
        "review_file_existed": review_exists,
        "manual_defined_count": manual_defined_count,
        "manual_titles_matched": manual_titles_matched_display,
        "manual_titles_unmatched": manual_titles_unmatched_display,
    })
    history = history[-MAX_RUN_LOG_ENTRIES:]
    save_json(log_path, history)


def step6_danish_backdrops() -> None:
    section("TRIN 6/7: Danske TMDb-backdrops (epgoal.xml) - skrives som <icon>")

    config = load_json(CONFIG_FILE, {})
    cache_max_age_days = config.get("cache_max_age_days", CACHE_MAX_AGE_DAYS)
    not_found_cache_max_age_days = config.get(
        "not_found_cache_max_age_days", NOT_FOUND_CACHE_MAX_AGE_DAYS_DEFAULT
    )

    cache = load_json(DANISH_ARTWORK_CACHE_FILE, {})
    cache_size_before = len(cache)

    approved_keys = load_approved_keys(DANISH_ARTWORK_REVIEW_FILE)
    review_exists = DANISH_ARTWORK_REVIEW_FILE.exists()

    manual_index, manual_display_titles = load_manual_overrides(MANUAL_ARTWORK_OVERRIDES_FILE)
    manual_titles_matched: set[str] = set()

    refresh_titles_raw = load_refresh_titles(GENOPFRISK_TITLER_FILE)
    undecided_review_keys = load_undecided_review_keys(DANISH_ARTWORK_REVIEW_FILE)
    refresh_titles_normalized = (
        {normalize_title(t) for t in refresh_titles_raw} | set(undecided_review_keys)
    )
    refresh_seen: set[str] = set()
    force_refreshed_this_run: set[str] = set()

    print(f"Cache indeholder {cache_size_before:,} tidligere opslag "
          f"(levetid: {cache_max_age_days} dage for fund, {not_found_cache_max_age_days} dage for 'ikke fundet')")

    if MANUAL_ARTWORK_OVERRIDES_FILE.exists():
        print(f"Manuelle overrides indlæst: {len(manual_index):,} unikke titler")
    else:
        print(f"INFO: {MANUAL_ARTWORK_OVERRIDES_FILE.name} findes ikke - ingen manuelle overrides denne gang.")

    if refresh_titles_raw:
        print(f"Tvangsopfrisker {len(refresh_titles_raw):,} titel(r) fra {GENOPFRISK_TITLER_FILE.name} ...")

    if undecided_review_keys:
        print(f"{len(undecided_review_keys):,} uafklarede titel(r) i {DANISH_ARTWORK_REVIEW_FILE.name} "
              "får friske TMDb-opslag ...")

    if approved_keys is None:
        print(f"ADVARSEL: {DANISH_ARTWORK_REVIEW_FILE.name} findes IKKE endnu. Ingen TMDb-fund injiceres.")
        approved_keys = set()
    else:
        print(f"Godkendelsesfil fundet: {len(approved_keys):,} unikke titler markeret med X.")

    all_found_keys: set[str] = set()
    resolved_this_run: dict[str, str | None] = {}

    stats = {
        "programmes": 0, "already_had_artwork": 0, "checked": 0,
        "danish_found": 0, "danish_not_found": 0, "danish_injected": 0,
        "cache_hits": 0, "fresh_calls": 0, "manual_override_injected": 0,
        "rechecked_after_not_found": 0, "force_refreshed": 0,
    }

    tree = ET.parse(EPGOAL_FILE)
    root = tree.getroot()

    for programme in root.findall("programme"):
        stats["programmes"] += 1

        # Allerede sat artwork i trin 5 (sport) -> rør IKKE, bevar som den er.
        if programme.find("icon") is not None or programme.find("backdrop") is not None:
            stats["already_had_artwork"] += 1
            continue

        title_el = programme.find("title")
        if title_el is None or not title_el.text:
            continue
        title = title_el.text
        chan_id = programme.get("channel", "")

        manual_url = resolve_manual_override(title, chan_id, manual_index)
        if manual_url:
            ET.SubElement(programme, "icon").set("src", manual_url)
            stats["manual_override_injected"] += 1
            manual_titles_matched.add(normalize_title(title))
            continue

        norm = normalize_title(title)
        if norm in refresh_titles_normalized:
            refresh_seen.add(norm)

        if norm not in resolved_this_run:
            force_refresh = norm in refresh_titles_normalized and norm not in force_refreshed_this_run
            was_cached_not_found = (
                norm in cache
                and not cache[norm].get("backdrop")
                and (time.time() - cache[norm].get("ts", 0)) / 86400 >= not_found_cache_max_age_days
                and (time.time() - cache[norm].get("ts", 0)) / 86400 < cache_max_age_days
            )
            backdrop_url, from_cache = resolve_danish_artwork(
                title, cache, cache_max_age_days, not_found_cache_max_age_days,
                TMDB_BACKDROP_SIZE, force_refresh=force_refresh,
            )
            resolved_this_run[norm] = backdrop_url
            stats["checked"] += 1

            if force_refresh:
                force_refreshed_this_run.add(norm)
                stats["force_refreshed"] += 1
            if from_cache:
                stats["cache_hits"] += 1
            else:
                stats["fresh_calls"] += 1
                if was_cached_not_found:
                    stats["rechecked_after_not_found"] += 1
                time.sleep(REQUEST_SLEEP_SECONDS)

            if backdrop_url:
                stats["danish_found"] += 1
                all_found_keys.add(norm)
            else:
                stats["danish_not_found"] += 1

        backdrop_url = resolved_this_run[norm]
        if backdrop_url and norm in approved_keys:
            ET.SubElement(programme, "icon").set("src", backdrop_url)
            stats["danish_injected"] += 1

    tree.write(EPGOAL_FILE, encoding="utf-8", xml_declaration=True)
    save_json(DANISH_ARTWORK_CACHE_FILE, cache)
    cache_size_after = len(cache)

    unique_found = len(all_found_keys)
    unique_approved_and_found = len(all_found_keys & approved_keys)
    unique_pending = unique_found - unique_approved_and_found

    all_manual_keys = set(manual_index.keys())
    unmatched_manual_keys = all_manual_keys - manual_titles_matched
    manual_matched_display = sorted(manual_display_titles.get(k, k) for k in manual_titles_matched)
    manual_unmatched_display = sorted(manual_display_titles.get(k, k) for k in unmatched_manual_keys)

    append_run_log(
        DANISH_BACKDROPS_RUN_LOG_FILE, stats, cache_size_before, cache_size_after,
        len(approved_keys), review_exists, unique_found, unique_pending,
        len(manual_index), manual_matched_display, manual_unmatched_display,
    )

    print(f"Programmer i alt          : {stats['programmes']:,}")
    print(f"Sprunget over (sport)     : {stats['already_had_artwork']:,}")
    if manual_index:
        print(f"Manuelle overrides indsat : {stats['manual_override_injected']:,} "
              f"({len(manual_titles_matched):,} ud af {len(manual_index):,} fundet i denne kørsel)")
    if manual_unmatched_display:
        print(f"INFO: {len(manual_unmatched_display):,} manual_artwork_overrides blev IKKE fundet i dagens EPG-vindue:")
        for t in manual_unmatched_display:
            print(f"     - {t}")
    print(f"Unikke titler med dansk backdrop: {unique_found:,} "
          f"(godkendt: {unique_approved_and_found:,}, afventer: {unique_pending:,})")
    if stats["rechecked_after_not_found"]:
        print(f"Titler gen-tjekket efter tidligere 'ikke fundet': {stats['rechecked_after_not_found']:,}")
    print(f"Cache voksede fra {cache_size_before:,} til {cache_size_after:,}")

    if refresh_titles_raw:
        found_now: list[str] = []
        still_pending: list[str] = []
        not_seen: list[str] = []
        for raw_title in refresh_titles_raw:
            norm = normalize_title(raw_title)
            cached_entry = cache.get(norm)
            has_backdrop = bool(cached_entry and cached_entry.get("backdrop"))
            if has_backdrop:
                found_now.append(raw_title)
            elif norm in refresh_seen:
                still_pending.append(raw_title)
            else:
                not_seen.append(raw_title)

        print(f"\nStatus for {GENOPFRISK_TITLER_FILE.name}:")
        print(f"   Fandt nu et dansk backdrop : {len(found_now):,} (fjernet fra listen)")
        for t in found_now:
            print(f"     - {t}")
        print(f"   Stadig intet fundet        : {len(still_pending):,} (bevares på listen)")
        if not_seen:
            print(f"   Ikke set i denne kørsel    : {len(not_seen):,} (bevares på listen)")
            for t in not_seen:
                print(f"     - {t}")

        remaining_titles = still_pending + not_seen
        save_refresh_titles(GENOPFRISK_TITLER_FILE, remaining_titles)
        print(f"   {GENOPFRISK_TITLER_FILE.name} opdateret ({len(remaining_titles):,} titel(r) tilbage).")


# ---------------------------------------------------------------------------
# TRIN 7: Git commit + push (samme fetch/rebase-retry-logik som originalerne)
# ---------------------------------------------------------------------------

def git_push(repo_dir: Path, commit_message: str) -> None:
    print("Committer og pusher til GitHub ...")
    try:
        subprocess.run(["git", "add", "-A"], cwd=repo_dir, check=False)
        result = subprocess.run(
            ["git", "commit", "-m", commit_message], cwd=repo_dir, check=False,
            capture_output=True, text=True,
        )
        if "nothing to commit" in (result.stdout + result.stderr).lower():
            print("Ingen ændringer at committe.")
            return

        subprocess.run(["git", "fetch", "origin"], cwd=repo_dir, check=False, capture_output=True, text=True)
        push_result = subprocess.run(
            ["git", "push", "origin", "main"], cwd=repo_dir, check=False, capture_output=True, text=True
        )
        if push_result.returncode == 0:
            print("Git push lykkedes.")
            return

        combined_output = (push_result.stdout + push_result.stderr).lower()
        if "rejected" in combined_output or "fetch first" in combined_output or "non-fast-forward" in combined_output:
            print("Push afvist (fjernrepo har nyere commits) - forsøger 'git pull --rebase' ...")
            rebase_result = subprocess.run(
                ["git", "pull", "--rebase", "origin", "main"], cwd=repo_dir, check=False,
                capture_output=True, text=True,
            )
            if rebase_result.returncode != 0:
                print("FEJL: 'git pull --rebase' fejlede - løs konflikten manuelt:", file=sys.stderr)
                print(rebase_result.stdout + rebase_result.stderr, file=sys.stderr)
                return
            retry_result = subprocess.run(
                ["git", "push", "origin", "main"], cwd=repo_dir, check=False, capture_output=True, text=True
            )
            if retry_result.returncode == 0:
                print("Git push lykkedes efter rebase.")
            else:
                print("FEJL: Git push fejlede STADIG efter rebase - tjek manuelt:", file=sys.stderr)
                print(retry_result.stdout + retry_result.stderr, file=sys.stderr)
        else:
            print("FEJL: Git push fejlede af en anden årsag:", file=sys.stderr)
            print(push_result.stdout + push_result.stderr, file=sys.stderr)
    except FileNotFoundError:
        print("ADVARSEL: git blev ikke fundet i PATH - springer commit/push over.", file=sys.stderr)


def step7_git_push() -> None:
    section("TRIN 7/7: Git commit + push")

    if not GIT_ENABLED:
        print("Git er deaktiveret (GIT_ENABLED = False) - springer over.")
        return

    message = f"{GIT_COMMIT_PREFIX} {time.strftime('%Y-%m-%d %H:%M:%S')}"
    git_push(ROOT, message)


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main() -> None:
    if not CHANNEL_PRIORITY_V2.exists():
        sys.exit(
            f"FEJL: Mangler {CHANNEL_PRIORITY_V2}.\n"
            "Kør build_master_channel_map_v2.py først for at generere denne fil."
        )
    if not TMDB_API_KEY:
        sys.exit("FEJL: TMDB_API_KEY er ikke sat i .env.")

    print("=== OPDATER EPGOAL (samlet produktionsscript) ===")
    print(f"Tidspunkt: {time.strftime('%Y-%m-%d %H:%M:%S')}")

    step1_download_epgshare()
    step2_filter_epgshare()
    step3_merge_sources()
    step4_normalize_uhf_ids()
    step5_sport_enrichment()
    step6_danish_backdrops()
    step7_git_push()

    section("FÆRDIG")
    print(f"epgoal.xml er opdateret: {EPGOAL_FILE}")
    print("UHF-URL (indsæt som eneste kilde i UHF):")
    print("  https://raw.githubusercontent.com/flanaganz/epgoal/main/output_epgshare/epgoal.xml")
    print()
    print(f"Se {MERGE_LOG_FILE.name} for hvilken kilde der blev brugt pr. kanal.")
    print(f"Se {NORMALIZE_LOG_FILE.name} for hvordan channel-id'er blev normaliseret.")
    print(f"Se {DANISH_BACKDROPS_RUN_LOG_FILE.name} for historik over danske backdrop-kørsler.")


if __name__ == "__main__":
    main()
