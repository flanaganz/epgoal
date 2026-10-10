#!/usr/bin/env python3
"""
bss_m3u_tvgid_check.py

Henter BSS's M3U-kanalliste (get.php?...&output=ts) og undersøger om
normal/HD/FHD-varianter af samme fysiske kanal deler samme tvg-id.

VIGTIGT: URL'en indeholder dit brugernavn/password. Den gemmes IKKE i denne
fil - du giver den som kommandolinje-argument, så den aldrig havner i et
scriptfil der evt. committes til GitHub.

Brug:
    python scripts/bss_m3u_tvgid_check.py "http://n2ip.tv:2095/get.php?username=...&password=...&type=m3u_plus&output=ts"
"""

import re
import sys
from collections import defaultdict
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
CHANNEL_PRIORITY = ROOT / "data" / "channel_priority.xlsx"

EXTINF_ATTR_PATTERN = re.compile(r'([\w-]+)="([^"]*)"')


def load_followed_keys():
    """Returnerer liste af (kanonisk_navn, gruppe_nøgle, [aliases])."""
    wb = load_workbook(CHANNEL_PRIORITY, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]

    kanal_col = headers.index("Kanal")
    key_col = headers.index("Gruppe-nøgle (intern)")
    aliases_col = headers.index("Sammenlagte kanal-ID'er")
    follow_col = headers.index("Følg (X)")

    rows = []
    for row in ws.iter_rows(min_row=2):
        if str(row[follow_col].value or "").strip().upper() != "X":
            continue
        kanal = row[kanal_col].value
        key = row[key_col].value
        if not kanal or not key:
            continue
        aliases_raw = row[aliases_col].value or ""
        aliases = [a.strip() for a in str(aliases_raw).split(",") if a.strip()]
        rows.append((str(kanal).strip(), str(key).strip().lower(), aliases))

    return rows


def normalize(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def main():
    if len(sys.argv) < 2:
        sys.exit(
            "❌ Mangler URL.\n"
            "Brug: python bss_m3u_tvgid_check.py \"http://n2ip.tv:2095/get.php?...&output=ts\""
        )

    url = sys.argv[1]

    print("📥 Henter M3U fra BSS ...")
    resp = requests.get(url, timeout=120)
    resp.raise_for_status()
    text = resp.text
    print(f"✅ Downloadet {len(text)//1024} KB")

    lines = text.splitlines()

    # tvg_id -> set af (tvg_name, group_title)
    tvgid_groups: dict[str, set] = defaultdict(set)
    total_entries = 0

    for i, line in enumerate(lines):
        if not line.startswith("#EXTINF"):
            continue
        total_entries += 1

        attrs = dict(EXTINF_ATTR_PATTERN.findall(line))
        tvg_id = attrs.get("tvg-id", "").strip()
        tvg_name = attrs.get("tvg-name", "").strip()
        group_title = attrs.get("group-title", "").strip()

        display_name = ""
        if "," in line:
            display_name = line.rsplit(",", 1)[-1].strip()

        key = tvg_id if tvg_id else f"(intet tvg-id: {display_name})"
        tvgid_groups[key].add((tvg_name or display_name, group_title))

    print(f"\n=== OVERBLIK ===")
    print(f"Total #EXTINF entries : {total_entries:,}")
    print(f"Unikke tvg-id         : {len(tvgid_groups):,}")

    # Hvor mange tvg-id'er har MERE end én navn-variant (dvs. dedup i praksis)?
    multi = {k: v for k, v in tvgid_groups.items() if len(v) > 1}
    print(f"tvg-id'er med >1 navn-variant (normal/HD/FHD dedup): {len(multi):,}")

    print("\n=== DIN FØLG-LISTE: tvg-id-gruppering pr. kanal ===")
    print("-" * 100)

    followed = load_followed_keys()

    for kanonisk_navn, gruppe_noegle, aliases in followed:
        print(f"\n{kanonisk_navn}  (gruppe-nøgle: {gruppe_noegle})")

        target_tokens = {normalize(gruppe_noegle)}
        for alias in aliases:
            target_tokens.add(normalize(alias.replace(".dk", "")))

        found_any = False
        for tvg_id, variants in tvgid_groups.items():
            norm_id = normalize(tvg_id.replace(".dk", ""))
            if not norm_id:
                continue
            if any(tok and tok in norm_id or norm_id in tok for tok in target_tokens if tok):
                found_any = True
                print(f"   tvg-id='{tvg_id}'  ->  {len(variants)} navn-variant(er):")
                for name, group in sorted(variants):
                    print(f"        - {name!r}  (group-title={group!r})")

        if not found_any:
            print("   (Intet tvg-id fundet der matcher denne kanal)")

    print("\n" + "-" * 100)
    print("Færdig. Hvis en kanal viser FLERE navne (fx 'DR1', 'DR1 HD', 'DR1 FHD')")
    print("under SAMME tvg-id, betyder det at BSS's xmltv.php sandsynligvis")
    print("allerede leverer ÉT samlet sæt programmer, der dækker alle varianter.")


if __name__ == "__main__":
    main()
