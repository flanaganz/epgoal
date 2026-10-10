#!/usr/bin/env python3

from pathlib import Path
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
import json
import re

ROOT = Path(__file__).resolve().parent.parent

PRIORITY_FILE = ROOT / "data" / "channel_priority_v2.xlsx"

INPUT_XML = ROOT / "output_epgshare" / "epgshare_merged.xml"
OUTPUT_XML = ROOT / "output_epgshare" / "epgoal.xml"

LOG_FILE = ROOT / "data" / "normalize_uhf_channel_ids_log.json"


def normalize(value):
    return re.sub(
        r"[^a-z0-9]",
        "",
        str(value or "").lower()
    )


def load_mapping():

    wb = load_workbook(
        PRIORITY_FILE,
        data_only=True
    )

    ws = wb["Kanal-prioritering v2"]

    headers = [
        c.value for c in ws[1]
    ]

    epgshare_col = headers.index(
        "EPGShare ID'er"
    )

    openepg_col = headers.index(
        "OpenEPG ID'er"
    )

    bss_col = headers.index(
        "BSS XMLTV ID'er"
    )

    output_col = headers.index(
        "Output/UHF tvg-id"
    )

    mapping = {}

    for row in ws.iter_rows(min_row=2):

        output_id = (
            str(
                row[output_col].value or ""
            )
            .strip()
        )

        if not output_id:
            continue

        source_values = []

        for idx in [
            epgshare_col,
            openepg_col,
            bss_col
        ]:

            raw = row[idx].value

            if not raw:
                continue

            parts = [
                x.strip()
                for x in str(raw).split(",")
                if x.strip()
            ]

            source_values.extend(parts)

        for src in source_values:

            mapping[
                normalize(src)
            ] = output_id

    return mapping


def main():

    if not INPUT_XML.exists():

        raise SystemExit(
            f"Mangler: {INPUT_XML}"
        )

    mapping = load_mapping()

    tree = ET.parse(INPUT_XML)
    root = tree.getroot()

    remapped_channels = {}

    new_channels = {}
    new_programmes = []

    for channel in root.findall("channel"):

        old_id = (
            channel.get("id") or ""
        )

        new_id = mapping.get(
            normalize(old_id),
            old_id
        )

        channel.set(
            "id",
            new_id
        )

        if new_id not in new_channels:

            new_channels[new_id] = channel

        remapped_channels[old_id] = new_id

    root[:] = [
        c
        for c in root
        if c.tag != "channel"
    ]

    for channel in new_channels.values():

        root.append(channel)

    seen_prog = set()

    for programme in root.findall("programme"):

        old_channel = (
            programme.get("channel")
            or ""
        )

        new_channel = mapping.get(
            normalize(old_channel),
            remapped_channels.get(
                old_channel,
                old_channel
            )
        )

        programme.set(
            "channel",
            new_channel
        )

        key = (
            new_channel,
            programme.get("start", ""),
            programme.findtext("title", "")
        )

        if key in seen_prog:
            continue

        seen_prog.add(key)

        new_programmes.append(
            programme
        )

    root[:] = [
        n
        for n in root
        if n.tag != "programme"
    ]

    for programme in new_programmes:

        root.append(programme)

    tree.write(
        OUTPUT_XML,
        encoding="utf-8",
        xml_declaration=True
    )

    with open(
        LOG_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            remapped_channels,
            f,
            ensure_ascii=False,
            indent=2
        )

    print()
    print("=== normalize_uhf_channel_ids ===")
    print()
    print(
        f"Mappings: {len(mapping):,}"
    )

    print(
        f"Kanaler : {len(new_channels):,}"
    )

    print(
        f"Programmer : {len(new_programmes):,}"
    )

    print(
        f"Output : {OUTPUT_XML}"
    )

    print(
       f"Log : {LOG_FILE}"
    )


if __name__ == "__main__":
    main()