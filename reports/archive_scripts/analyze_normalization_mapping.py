#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from openpyxl import load_workbook
from pathlib import Path
import re

XLSX = Path(r"C:\EPGoal\data\channel_priority_v2.xlsx")

WATCH = [
    "tv3",
    "tv3+",
]

def normalize(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

def split_ids(value):
    if value is None:
        return []

    return [
        x.strip()
        for x in re.split(r"[,;\n]+", str(value))
        if x.strip()
    ]

wb = load_workbook(XLSX, data_only=True)

if "Kanal-prioritering v2" in wb.sheetnames:
    ws = wb["Kanal-prioritering v2"]
else:
    ws = wb.active

headers = [c.value for c in ws[1]]

cols = {
    h: headers.index(h)
    for h in headers
    if h
}

print()
print("=== NORMALIZATION MAPPING ANALYSIS ===")
print()

for row in ws.iter_rows(min_row=2):

    kanal = row[cols["Kanal"]].value

    if not kanal:
        continue

    values = []

    for col in (
        "Kanal",
        "EPGShare ID'er",
        "OpenEPG ID'er",
        "BSS XMLTV ID'er",
        "Output/UHF tvg-id"
    ):
        if col in cols:
            values.append(str(row[cols[col]].value or ""))

    search_blob = " ".join(values).lower()

    if not any(x in search_blob for x in WATCH):
        continue

    print("=" * 70)
    print(f"Kanal: {kanal}")
    print("-" * 70)

    for col in (
        "Kanal",
        "EPGShare ID'er",
        "OpenEPG ID'er",
        "BSS XMLTV ID'er",
        "Output/UHF tvg-id",
        "Sammenlagte kanal-ID'er (gammel reference)"
    ):
        if col not in cols:
            continue

        value = row[cols[col]].value

        print(f"{col}:")
        print(value)
        print()

print()
print("=" * 70)
print("NORMALISERET OPSLAGSTABEL")
print("=" * 70)

mapping = {}

for row in ws.iter_rows(min_row=2):

    output = ""

    if "Output/UHF tvg-id" in cols:
        output = str(
            row[cols["Output/UHF tvg-id"]].value or ""
        ).strip()

    if not output:
        continue

    for source_col in (
        "EPGShare ID'er",
        "OpenEPG ID'er",
        "BSS XMLTV ID'er"
    ):

        if source_col not in cols:
            continue

        value = row[cols[source_col]].value

        for sid in split_ids(value):

            mapping[normalize(sid)] = output

for k in sorted(mapping):

    if "tv3" in k:

        print(f"{k:25} -> {mapping[k]}")