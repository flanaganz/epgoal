#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""EPGoal variantanalyse v3.

Analyserer Normal/HD/FHD/UHD-streamgrupper fra channel_priority_v2.xlsx.
Validerer tvg-id'er mod epgoal.xml og rå EPG-kilder. Ændrer ingen kilder.
"""
from __future__ import annotations
import argparse, csv, json, re, unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("FEJL: openpyxl mangler. Kør: python -m pip install openpyxl")

BASE=Path(r"C:\EPGoal\output_epgshare")
QUALITY_ORDER={"NORMAL":0,"SD":1,"HD":2,"FHD":3,"UHD":4,"4K":4,"UKENDT":9}
IGNORE={"dk","denmark","danmark","multi","audio","sub","only","on","devices"}

def ascii_text(s):return "".join(c for c in unicodedata.normalize("NFKD",str(s or "")) if not unicodedata.combining(c))
def split_values(v):return [x.strip() for x in re.split(r"[,;\n]+",str(v or "")) if x.strip() and x.strip().upper()!="X"]
def clean_id(s):return re.sub(r"[^a-z0-9+]","",ascii_text(s).lower())
def canonical_name(s):
    s=ascii_text(s).lower();s=re.sub(r"\[[^]]*]"," ",s);s=re.sub(r"\([^)]*\)"," ",s)
    toks=[x for x in re.findall(r"[a-z0-9+]+",s) if x not in IGNORE and x not in {"sd","hd","fhd","uhd","4k"}]
    return "".join(toks)
def quality(name):
    u=str(name or "").upper()
    if re.search(r"\b(4K|UHD)\b",u):return "UHD"
    if re.search(r"\bFHD\b",u):return "FHD"
    if re.search(r"\bHD\b",u):return "HD"
    if re.search(r"\bSD\b",u):return "SD"
    return "NORMAL"
def parse_xml(path):
    channels={};counts=Counter();names=defaultdict(list)
    if not path.exists():return channels,counts,names
    for _,e in ET.iterparse(path,events=("end",)):
        tag=e.tag.rsplit("}",1)[-1]
        if tag=="channel":
            cid=(e.get("id") or "").strip()
            if cid:
                channels[cid]=True;names[cid]=[(x.text or "").strip() for x in e.findall("./display-name") if (x.text or "").strip()]
            e.clear()
        elif tag=="programme":
            cid=(e.get("channel") or "").strip()
            if cid:counts[cid]+=1
            e.clear()
    return channels,counts,names
def priority_path(base,explicit):
    opts=[explicit,base.parent/"data"/"channel_priority_v2.xlsx",base.parent/"channel_priority_v2.xlsx",Path.cwd()/"channel_priority_v2.xlsx"]
    return next((p for p in opts if p and p.exists()),None)
def load_priority(path):
    ws=load_workbook(path,read_only=True,data_only=True).worksheets[0];rows=list(ws.iter_rows(values_only=True))
    hi=next(i for i,r in enumerate(rows) if r and str(r[0]).strip()=="Kanal");hdr=[str(x or "").strip() for x in rows[hi]]
    return [{hdr[i]:(r[i] if i<len(r) else None) for i in range(len(hdr))} for r in rows[hi+1:] if r and r[0]]
def pair_streams(ids,names):
    if not names:return []
    if len(ids)==len(names):return list(zip(ids,names))
    if len(ids)==1:return [(ids[0],n) for n in names]
    # Bedste forsøg: match id til streamnavn; ellers blankt id.
    out=[]
    for n in names:
        cn=canonical_name(n);best="";bestlen=0
        for i in ids:
            ci=clean_id(i)
            if ci and (ci in cn or cn in ci) and len(ci)>bestlen:best=i;bestlen=len(ci)
        out.append((best,n))
    return out

def source_ids(rec):
    ids=set()
    for col in ("EPGShare ID'er","OpenEPG ID'er","BSS XMLTV ID'er","Output/UHF tvg-id","Sammenlagte kanal-ID'er (gammel reference)"):
        ids.update(split_values(rec.get(col)))
    return ids
def status_for(group,stream_id,stream_name,out_id,known_ids,epgoal_counts):
    reasons=[];status="OK"
    if not stream_id:
        return "MANGLER TVG-ID","Streamen har intet tvg-id"
    sid=clean_id(stream_id);oid=clean_id(out_id);known={clean_id(x) for x in known_ids}
    if oid and sid==oid: pass
    elif sid in known: reasons.append("kendt alternativt kilde-ID")
    elif sid in {clean_id(x) for x in epgoal_counts}: reasons.append("findes som andet EPGoal-ID");status="KONTROLLÉR"
    else: reasons.append("ukendt tvg-id");status="KONTROLLÉR"
    # Bevar TV3 og TV3+ helt adskilt.
    g=canonical_name(group);n=canonical_name(stream_name);raw=str(stream_name).lower()
    if g=="tv3" and ("tv3+" in raw or "tv3plus" in n):return "FORKERT KANAL","TV3-gruppen indeholder en TV3+ stream"
    if (g in {"tv3+","tv3plus"} or "+" in str(group)) and re.search(r"\btv3\b",raw) and "+" not in raw:return "FORKERT KANAL","TV3+-gruppen indeholder en TV3 stream"
    if oid and epgoal_counts.get(out_id,0)==0:reasons.append("output-ID har 0 programmer");status="MANGLER EPG"
    return status,", ".join(reasons)
def add_sheet(wb,title,headers,rows):
    ws=wb.create_sheet(title);ws.append(headers)
    for r in rows:ws.append(r)
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;head=PatternFill("solid",fgColor="1F4E78")
    colors={"OK":"C6EFCE","KONTROLLÉR":"FFEB9C","MANGLER TVG-ID":"FFC7CE","MANGLER EPG":"FFC7CE","FORKERT KANAL":"F4B084"}
    for c in ws[1]:c.font=Font(color="FFFFFF",bold=True);c.fill=head;c.alignment=Alignment(wrap_text=True)
    for r in range(2,ws.max_row+1):
        st=str(ws.cell(r,1).value or "")
        if st in colors:ws.cell(r,1).fill=PatternFill("solid",fgColor=colors[st])
    for col in range(1,ws.max_column+1):
        vals=[len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,300)+1)]
        ws.column_dimensions[get_column_letter(col)].width=min(55,max(10,max(vals,default=8)+2))
    return ws

def main():
    ap=argparse.ArgumentParser(description="Analyser Normal/HD/FHD/UHD tvg-id grupper")
    ap.add_argument("--base",type=Path,default=BASE);ap.add_argument("--priority",type=Path);ap.add_argument("--report-dir",type=Path)
    a=ap.parse_args();pf=priority_path(a.base,a.priority)
    if not pf:raise SystemExit("FEJL: channel_priority_v2.xlsx blev ikke fundet. Brug --priority STi")
    report=a.report_dir or a.base.parent/"reports";report.mkdir(parents=True,exist_ok=True)
    print(f"Prioritetsfil: {pf}");rows=load_priority(pf)
    print("Indlæser EPGoal og kilde-ID'er...")
    epch,epcounts,epnames=parse_xml(a.base/"epgoal.xml")
    source_sets={}
    for src,path in {"EPGShare":a.base/"epgshare_dk1.xml","BSS":a.base/"bss_raw"/"bss_epg.xml"}.items():
        ch,counts,names=parse_xml(path);source_sets[src]=(set(ch),counts);print(f"  {src}: {len(ch):,} ID'er")
    op_ids=set();op_counts=Counter()
    for p in sorted((a.base/"openepg_raw").glob("*.xml")):
        ch,c,_=parse_xml(p);op_ids.update(ch);op_counts.update(c)
    source_sets["OpenEPG"]=(op_ids,op_counts);print(f"  OpenEPG: {len(op_ids):,} ID'er")
    detail=[];group_rows=[];status_count=Counter()
    for rec in rows:
        group=str(rec.get("Kanal") or "");out_id=str(rec.get("Output/UHF tvg-id") or "").strip();ids=split_values(rec.get("BSS M3U tvg-id'er"));names=split_values(rec.get("BSS M3U stream-navne"));pairs=pair_streams(ids,names);known=source_ids(rec)
        group_status=[];quals=Counter();used=Counter()
        for sid,sname in pairs:
            q=quality(sname);quals[q]+=1;used[sid]+=1
            st,note=status_for(group,sid,sname,out_id,known,epcounts);status_count[st]+=1;group_status.append(st)
            src_hits=[]
            for src,(sids,counts) in source_sets.items():
                exact=next((x for x in sids if clean_id(x)==clean_id(sid)),None)
                if exact:src_hits.append(f"{src}:{exact}({counts.get(exact,0)})")
            detail.append([st,note,group,q,sname,sid,out_id,epcounts.get(out_id,0),", ".join(sorted(known)),"; ".join(src_hits)])
        if not pairs:gstatus="INGEN STREAMS"
        elif "FORKERT KANAL" in group_status:gstatus="FORKERT KANAL"
        elif "MANGLER EPG" in group_status:gstatus="MANGLER EPG"
        elif "MANGLER TVG-ID" in group_status:gstatus="MANGLER TVG-ID"
        elif "KONTROLLÉR" in group_status:gstatus="KONTROLLÉR"
        else:gstatus="OK"
        duplicate_ids=", ".join(f"{k} x{v}" for k,v in used.items() if k and v>1)
        group_rows.append([gstatus,group,out_id,epcounts.get(out_id,0),len(pairs),quals.get("NORMAL",0),quals.get("HD",0),quals.get("FHD",0),quals.get("UHD",0),len(set(x for x,_ in pairs if x)),duplicate_ids,", ".join(names)])
    headers=["Status","Bemærkning","Kanalgruppe","Kvalitet","Streamnavn","Stream tvg-id","Output/UHF tvg-id","Output programmer","Kendte ID'er fra priority","Fundet i kilder"]
    gh=["Status","Kanalgruppe","Output/UHF tvg-id","Output programmer","Streams","Normal","HD","FHD","UHD/4K","Unikke tvg-id'er","Genbrugte tvg-id'er","Streamnavne"]
    wb=Workbook();wb.remove(wb.active)
    summary=[["Kørt",datetime.now().strftime("%Y-%m-%d %H:%M:%S")],["Prioritetsfil",str(pf)],["EPGoal kanal-id'er",len(epch)],["Grupper",len(group_rows)],["Streams",len(detail)]]+[[f"Streamstatus: {k}",v] for k,v in sorted(status_count.items())]
    add_sheet(wb,"Summary",["Måling","Værdi"],summary);add_sheet(wb,"Variant Groups",gh,group_rows);add_sheet(wb,"Stream Details",headers,detail);add_sheet(wb,"Problems",headers,[r for r in detail if r[0]!="OK"]);add_sheet(wb,"TV3-TV3plus",headers,[r for r in detail if canonical_name(r[2]) in {"tv3","tv3+","tv3plus"} or "tv3" in str(r[4]).lower()])
    xp=report/"variant_groups_v3.xlsx";cp=report/"variant_group_problems_v3.csv";jp=report/"variant_groups_v3.json";wb.save(xp)
    with cp.open("w",encoding="utf-8-sig",newline="") as f:w=csv.writer(f,delimiter=";");w.writerow(headers);w.writerows(r for r in detail if r[0]!="OK")
    jp.write_text(json.dumps({"generated":datetime.now().isoformat(timespec="seconds"),"statuses":dict(status_count),"groups":[dict(zip(gh,r)) for r in group_rows],"problems":[dict(zip(headers,r)) for r in detail if r[0]!="OK"]},ensure_ascii=False,indent=2),encoding="utf-8")
    print("\n=== VARIANTANALYSE V3 FÆRDIG ===");print(f"Excel : {xp}\nCSV   : {cp}\nJSON  : {jp}")
if __name__=="__main__":main()
