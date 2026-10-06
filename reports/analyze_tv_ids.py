#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""EPGoal TV-ID analyseværktøj.

Læser epgoal.xml, EPGShare, BSS, OpenEPG og channel_priority_v2.xlsx.
Skriver en samlet Excel-rapport samt CSV og JSON. Kilder ændres aldrig.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Set, Tuple
import xml.etree.ElementTree as ET

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("FEJL: openpyxl mangler. Kør: python -m pip install openpyxl")
    raise SystemExit(2)

DEFAULT_BASE = Path(r"C:\EPGoal\output_epgshare")
DEFAULT_PRIORITY = Path(r"C:\EPGoal\data\channel_priority_v2.xlsx")

QUALITY_WORDS = {"hd", "fhd", "uhd", "4k", "sd"}
NOISE_WORDS = {
    "dk", "denmark", "danmark", "da", "d", "t", "tv", "channel", "kanal",
    "nordic", "nordics", "multi", "audio", "sub", "digital"
}

@dataclass
class Programme:
    start: str = ""
    stop: str = ""
    title: str = ""
    subtitle: str = ""
    desc: str = ""

@dataclass
class Channel:
    source: str
    source_file: str
    channel_id: str
    names: List[str] = field(default_factory=list)
    icon: str = ""
    programmes: List[Programme] = field(default_factory=list)

    @property
    def display_name(self) -> str:
        return next((x for x in self.names if x.strip()), self.channel_id)

    @property
    def count(self) -> int:
        return len(self.programmes)

    @property
    def first_start(self) -> str:
        vals = sorted(p.start for p in self.programmes if p.start)
        return vals[0] if vals else ""

    @property
    def last_stop(self) -> str:
        vals = sorted(p.stop for p in self.programmes if p.stop)
        return vals[-1] if vals else ""

    @property
    def title_set(self) -> Set[str]:
        return {normalize_title(p.title) for p in self.programmes if normalize_title(p.title)}


def text_of(parent, tag: str) -> str:
    el = parent.find(tag)
    return (el.text or "").strip() if el is not None else ""


def strip_accents(value: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", value) if not unicodedata.combining(c))


def normalize_name(value: str, remove_quality: bool = True) -> str:
    value = strip_accents(value or "").lower().replace("&amp;", "and").replace("&", "and")
    tokens = re.findall(r"[a-z0-9]+", value)
    out = []
    for token in tokens:
        if remove_quality and token in QUALITY_WORDS:
            continue
        if token in NOISE_WORDS:
            continue
        out.append(token)
    return "".join(out)


def normalize_title(value: str) -> str:
    value = strip_accents(value or "").lower()
    return re.sub(r"[^a-z0-9]+", "", value)


def parse_xmltv(path: Path, source: str) -> Dict[str, Channel]:
    channels: Dict[str, Channel] = {}
    if not path.exists():
        return channels
    try:
        for event, elem in ET.iterparse(path, events=("end",)):
            tag = elem.tag.rsplit("}", 1)[-1]
            if tag == "channel":
                cid = (elem.get("id") or "").strip()
                if cid:
                    names = [(n.text or "").strip() for n in elem.findall("./display-name") if (n.text or "").strip()]
                    icon_el = elem.find("./icon")
                    icon = (icon_el.get("src") or "") if icon_el is not None else ""
                    channels[cid] = Channel(source, str(path), cid, names, icon)
                elem.clear()
            elif tag == "programme":
                cid = (elem.get("channel") or "").strip()
                if cid:
                    ch = channels.setdefault(cid, Channel(source, str(path), cid, [cid]))
                    ch.programmes.append(Programme(
                        start=(elem.get("start") or "").strip(),
                        stop=(elem.get("stop") or "").strip(),
                        title=text_of(elem, "./title"),
                        subtitle=text_of(elem, "./sub-title"),
                        desc=text_of(elem, "./desc"),
                    ))
                elem.clear()
    except ET.ParseError as exc:
        print(f"ADVARSEL: Kunne ikke parse {path}: {exc}")
    return channels


def load_sources(base: Path) -> Tuple[Dict[str, Dict[str, Channel]], List[str]]:
    specs = [
        ("EPGoal", base / "epgoal.xml"),
        ("EPGShare", base / "epgshare_dk1.xml"),
        ("BSS", base / "bss_raw" / "bss_epg.xml"),
    ]
    sources: Dict[str, Dict[str, Channel]] = {}
    warnings: List[str] = []
    for name, path in specs:
        print(f"Indlæser {name}: {path}")
        if not path.exists():
            warnings.append(f"Mangler: {path}")
        sources[name] = parse_xmltv(path, name)
        print(f"  {len(sources[name]):,} kanal-id'er")

    open_dir = base / "openepg_raw"
    open_files = sorted(open_dir.glob("*.xml")) if open_dir.exists() else []
    if not open_files:
        warnings.append(f"Ingen XML-filer fundet i: {open_dir}")
    merged: Dict[str, Channel] = {}
    for path in open_files:
        print(f"Indlæser OpenEPG: {path.name}")
        parsed = parse_xmltv(path, "OpenEPG")
        for cid, ch in parsed.items():
            if cid not in merged:
                merged[cid] = ch
            else:
                merged[cid].programmes.extend(ch.programmes)
                merged[cid].names.extend(n for n in ch.names if n not in merged[cid].names)
                merged[cid].source_file += "; " + str(path)
    sources["OpenEPG"] = merged
    print(f"  {len(merged):,} unikke OpenEPG kanal-id'er")
    return sources, warnings


def find_priority_file(base: Path, explicit: Optional[Path]) -> Optional[Path]:
    candidates = [explicit] if explicit else []
    candidates += [
        DEFAULT_PRIORITY,
        base.parent / "data" / "channel_priority_v2.xlsx",
        base.parent / "channel_priority_v2.xlsx",
        Path.cwd() / "channel_priority_v2.xlsx",
    ]
    return next((p for p in candidates if p and p.exists()), None)


def split_ids(value) -> List[str]:
    if value is None:
        return []
    return [x.strip() for x in re.split(r"[,;\n]+", str(value)) if x.strip() and x.strip().upper() != "X"]


def load_priority(path: Optional[Path]) -> Tuple[List[dict], Dict[str, List[dict]]]:
    if not path:
        return [], defaultdict(list)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next((i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Kanal"), None)
    if header_idx is None:
        return [], defaultdict(list)
    headers = [str(x).strip() if x is not None else "" for x in rows[header_idx]]
    data, index = [], defaultdict(list)
    for row in rows[header_idx + 1:]:
        if not row or not row[0]:
            continue
        rec = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        data.append(rec)
        keys = set()
        for col in ("Kanal", "EPGShare ID'er", "OpenEPG ID'er", "BSS XMLTV ID'er", "Output/UHF tvg-id", "Sammenlagte kanal-ID'er (gammel reference)"):
            keys.update(split_ids(rec.get(col)))
        for key in keys:
            index[normalize_name(key)].append(rec)
    return data, index


def name_similarity(a: Channel, b: Channel) -> float:
    variants_a = [a.channel_id, a.display_name] + a.names
    variants_b = [b.channel_id, b.display_name] + b.names
    best = 0.0
    for x in variants_a:
        nx = normalize_name(x)
        if not nx:
            continue
        for y in variants_b:
            ny = normalize_name(y)
            if not ny:
                continue
            score = SequenceMatcher(None, nx, ny).ratio()
            if nx == ny:
                score = 1.0
            elif nx in ny or ny in nx:
                score = max(score, 0.88)
            best = max(best, score)
    return best


def title_overlap(a: Channel, b: Channel) -> Tuple[int, float]:
    sa, sb = a.title_set, b.title_set
    if not sa or not sb:
        return 0, 0.0
    common = len(sa & sb)
    return common, common / max(1, min(len(sa), len(sb)))


def candidate_score(target: Channel, cand: Channel, priority_ids: Set[str]) -> Tuple[float, float, int, float, str]:
    ns = name_similarity(target, cand)
    common, overlap = title_overlap(target, cand)
    exact_id = normalize_name(target.channel_id) == normalize_name(cand.channel_id)
    mapped = cand.channel_id in priority_ids or normalize_name(cand.channel_id) in {normalize_name(x) for x in priority_ids}
    score = ns * 65 + overlap * 25 + (8 if exact_id else 0) + (12 if mapped else 0)
    if cand.count == 0:
        score -= 8
    reasons = []
    if mapped: reasons.append("priority-map")
    if exact_id: reasons.append("normaliseret ID")
    if ns >= .80: reasons.append("navnelighed")
    if overlap >= .50: reasons.append("programoverlap")
    return round(max(0, min(100, score)), 1), round(ns * 100, 1), common, round(overlap * 100, 1), ", ".join(reasons)


def priority_for(ch: Channel, index: Dict[str, List[dict]]) -> Optional[dict]:
    keys = [normalize_name(ch.channel_id), normalize_name(ch.display_name)]
    for key in keys:
        if index.get(key):
            return index[key][0]
    return None


def expected_ids(rec: Optional[dict], source: str) -> Set[str]:
    if not rec:
        return set()
    col = {"EPGShare": "EPGShare ID'er", "BSS": "BSS XMLTV ID'er", "OpenEPG": "OpenEPG ID'er"}.get(source)
    return set(split_ids(rec.get(col))) if col else set()


def best_candidates(target: Channel, pool: Dict[str, Channel], mapped_ids: Set[str], limit=5) -> List[dict]:
    result = []
    for cand in pool.values():
        score, ns, common, overlap, reason = candidate_score(target, cand, mapped_ids)
        if score >= 42 or cand.channel_id in mapped_ids:
            result.append({"channel": cand, "score": score, "name": ns, "common": common, "overlap": overlap, "reason": reason})
    return sorted(result, key=lambda x: (-x["score"], -x["channel"].count, x["channel"].channel_id.lower()))[:limit]


def status_for(ch: Channel, rows_by_source: Dict[str, List[dict]]) -> Tuple[str, str]:
    if ch.count == 0:
        best = max((r["score"] for rows in rows_by_source.values() for r in rows), default=0)
        return ("MANGLER EPG", "Kandidat fundet i en rå kilde" if best >= 65 else "Ingen sikker kandidat")
    high_other = [r for rows in rows_by_source.values() for r in rows if r["score"] >= 78 and r["overlap"] < 15 and ch.count >= 5]
    if high_other:
        return "KONTROLLÉR INDHOLD", "Stærkt navnematch, men lav programoverlap"
    exact_sources = sum(any(r["score"] >= 70 for r in rows) for rows in rows_by_source.values())
    if exact_sources == 0:
        return "ID-KONFLIKT", "Ingen tydelig rå kilde matcher outputkanalen"
    return "OK", ""


def add_sheet(wb, title: str, headers: List[str], rows: Iterable[Iterable], freeze="A2", autofilter=True):
    ws = wb.create_sheet(title[:31])
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    ws.freeze_panes = freeze
    if autofilter and ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col in range(1, ws.max_column + 1):
        values = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 250) + 1)]
        width = min(55, max(10, max((len(v) for v in values), default=10) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def main() -> int:
    ap = argparse.ArgumentParser(description="Analyser TV-ID'er på tværs af EPGoal, EPGShare, BSS og OpenEPG")
    ap.add_argument("--base", type=Path, default=DEFAULT_BASE, help="Mappe med epgoal.xml og rå kilder")
    ap.add_argument("--priority", type=Path, help="Sti til channel_priority_v2.xlsx")
    ap.add_argument("--report-dir", type=Path, help="Outputmappe, standard: C:\\EPGoal\\reports")
    ap.add_argument("--top", type=int, default=5, help="Antal kandidater pr. kilde")
    args = ap.parse_args()

    base = args.base
    report_dir = args.report_dir or base.parent / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    priority_path = find_priority_file(base, args.priority)
    sources, warnings = load_sources(base)
    priority_rows, priority_index = load_priority(priority_path)
    if not priority_path:
        warnings.append("channel_priority_v2.xlsx blev ikke fundet; analysen fortsatte uden manuelle mappings")
    else:
        print(f"Prioritetsfil: {priority_path}")

    epgoal = sources["EPGoal"]
    all_rows, candidate_rows, variant_rows, missing_rows = [], [], [], []
    statuses = Counter()

    for ch in sorted(epgoal.values(), key=lambda x: (x.display_name.lower(), x.channel_id.lower())):
        rec = priority_for(ch, priority_index)
        by_source = {}
        for source in ("EPGShare", "BSS", "OpenEPG"):
            mapped = expected_ids(rec, source)
            rows = best_candidates(ch, sources[source], mapped, args.top)
            by_source[source] = rows
            for rank, item in enumerate(rows, 1):
                c = item["channel"]
                candidate_rows.append([
                    ch.display_name, ch.channel_id, ch.count, source, rank, c.display_name, c.channel_id,
                    c.count, c.first_start, c.last_stop, item["score"], item["name"], item["common"],
                    item["overlap"], item["reason"], Path(c.source_file.split("; ")[0]).name,
                ])
        status, note = status_for(ch, by_source)
        statuses[status] += 1
        best = {s: (rows[0] if rows else None) for s, rows in by_source.items()}
        all_rows.append([
            status, note, ch.display_name, ch.channel_id, ch.count, ch.first_start, ch.last_stop,
            rec.get("Kanal", "") if rec else "", rec.get("Output/UHF tvg-id", "") if rec else "",
            best["EPGShare"]["channel"].channel_id if best["EPGShare"] else "",
            best["EPGShare"]["channel"].count if best["EPGShare"] else 0,
            best["EPGShare"]["score"] if best["EPGShare"] else 0,
            best["BSS"]["channel"].channel_id if best["BSS"] else "",
            best["BSS"]["channel"].count if best["BSS"] else 0,
            best["BSS"]["score"] if best["BSS"] else 0,
            best["OpenEPG"]["channel"].channel_id if best["OpenEPG"] else "",
            best["OpenEPG"]["channel"].count if best["OpenEPG"] else 0,
            best["OpenEPG"]["score"] if best["OpenEPG"] else 0,
        ])
        if ch.count == 0:
            missing_rows.append(all_rows[-1])

    # Variantanalyse fra priority-arket
    for rec in priority_rows:
        stream_ids = split_ids(rec.get("BSS M3U tvg-id'er"))
        stream_names = split_ids(rec.get("BSS M3U stream-navne"))
        qualities = sorted({q.upper() for n in stream_names for q in QUALITY_WORDS if re.search(rf"\b{q}\b", n, re.I)})
        variant_rows.append([
            rec.get("Kanal", ""), rec.get("Følg (X)", ""), rec.get("Output/UHF tvg-id", ""),
            ", ".join(stream_ids), ", ".join(stream_names), ", ".join(qualities),
            len(stream_ids), len(stream_names),
            "KONTROLLÉR" if len(stream_ids) > 1 or (stream_names and not stream_ids) else "OK",
        ])

    wb = Workbook()
    wb.remove(wb.active)
    summary_rows = [
        ["Kørt", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Basismappe", str(base)], ["Prioritetsfil", str(priority_path or "IKKE FUNDET")],
        ["EPGoal kanal-id'er", len(sources["EPGoal"])], ["EPGShare kanal-id'er", len(sources["EPGShare"])],
        ["BSS kanal-id'er", len(sources["BSS"])], ["OpenEPG kanal-id'er", len(sources["OpenEPG"])],
    ] + [[f"Status: {k}", v] for k, v in sorted(statuses.items())] + [["Advarsel", w] for w in warnings]
    add_sheet(wb, "Summary", ["Måling", "Værdi"], summary_rows, autofilter=False)
    headers = ["Status", "Bemærkning", "Display name", "EPGoal ID", "EPGoal programmer", "Første start", "Sidste stop",
               "Priority kanal", "Output/UHF tvg-id", "Bedste EPGShare ID", "EPGShare programmer", "EPGShare score",
               "Bedste BSS ID", "BSS programmer", "BSS score", "Bedste OpenEPG ID", "OpenEPG programmer", "OpenEPG score"]
    ws_all = add_sheet(wb, "Channel Sources", headers, all_rows)
    ws_missing = add_sheet(wb, "Missing EPG", headers, missing_rows)
    cand_headers = ["EPGoal navn", "EPGoal ID", "EPGoal programmer", "Kilde", "Rang", "Kandidatnavn", "Kandidat-ID",
                    "Kandidat programmer", "Første start", "Sidste stop", "Samlet score", "Navnescore", "Fælles titler",
                    "Titeloverlap pct", "Matchgrund", "Kildefil"]
    add_sheet(wb, "Potential Matches", cand_headers, candidate_rows)
    wrong = [r for r in all_rows if r[0] in ("KONTROLLÉR INDHOLD", "ID-KONFLIKT")]
    add_sheet(wb, "Wrong Mappings", headers, wrong)
    add_sheet(wb, "HD-FHD Variants", ["Kanal", "Følg", "Output ID", "M3U tvg-id'er", "Streamnavne", "Varianter", "ID-antal", "Stream-antal", "Status"], variant_rows)

    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    green = PatternFill("solid", fgColor="C6EFCE")
    for ws in (ws_all, ws_missing):
        for row in range(2, ws.max_row + 1):
            status = str(ws.cell(row, 1).value or "")
            fill = green if status == "OK" else red if status == "MANGLER EPG" else yellow
            ws.cell(row, 1).fill = fill

    xlsx_path = report_dir / "tv_id_analysis.xlsx"
    wb.save(xlsx_path)

    csv_path = report_dir / "tv_id_candidates.csv"
    import csv
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(cand_headers)
        w.writerows(candidate_rows)

    json_path = report_dir / "tv_id_analysis.json"
    json_path.write_text(json.dumps({
        "generated": datetime.now().isoformat(timespec="seconds"), "base": str(base), "warnings": warnings,
        "counts": {k: len(v) for k, v in sources.items()}, "statuses": dict(statuses),
        "channels": [dict(zip(headers, row)) for row in all_rows],
    }, ensure_ascii=False, indent=2), encoding="utf-8")

    print("\n=== TV-ID ANALYSE FÆRDIG ===")
    print(f"Excel : {xlsx_path}")
    print(f"CSV   : {csv_path}")
    print(f"JSON  : {json_path}")
    print("Status: " + ", ".join(f"{k}={v}" for k, v in sorted(statuses.items())))
    if warnings:
        print("Advarsler:")
        for w in warnings: print(f"  - {w}")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
