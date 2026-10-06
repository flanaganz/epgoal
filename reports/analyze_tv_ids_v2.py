#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""EPGoal TV-ID analyse v2: indekserer først, beregner kun programoverlap på shortlisten."""
from __future__ import annotations
import argparse, csv, json, re, unicodedata
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher, get_close_matches
from pathlib import Path
from typing import Dict, List, Set
try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    raise SystemExit("FEJL: openpyxl mangler. Kør: python -m pip install openpyxl")

BASE = Path(r"C:\EPGoal\output_epgshare")
NOISE = {"dk","denmark","danmark","da","d","t","tv","channel","kanal","nordic","nordics","multi","audio","sub","digital"}
QUALITY = {"sd","hd","fhd","uhd","4k"}

@dataclass
class Programme:
    start:str=""; stop:str=""; title:str=""
@dataclass
class Channel:
    source:str; source_file:str; channel_id:str; names:List[str]=field(default_factory=list); programmes:List[Programme]=field(default_factory=list)
    @property
    def name(self): return next((x for x in self.names if x.strip()), self.channel_id)
    @property
    def count(self): return len(self.programmes)
    @property
    def first(self): return min((p.start for p in self.programmes if p.start), default="")
    @property
    def last(self): return max((p.stop for p in self.programmes if p.stop), default="")
    @property
    def titles(self): return {norm_title(p.title) for p in self.programmes if norm_title(p.title)}

def ascii_text(s): return "".join(c for c in unicodedata.normalize("NFKD", s or "") if not unicodedata.combining(c))
def tokens(s):
    ts=set(re.findall(r"[a-z0-9]+", ascii_text(s).lower().replace("&amp;","and").replace("&","and")))
    return {x for x in ts if x not in NOISE|QUALITY and len(x)>=2}
def norm(s): return "".join(sorted(tokens(s)))
def norm_title(s): return re.sub(r"[^a-z0-9]+","",ascii_text(s).lower())
def child_text(e, tag):
    x=e.find(tag); return (x.text or "").strip() if x is not None else ""

def parse_xml(path:Path, source:str)->Dict[str,Channel]:
    out={}
    if not path.exists(): return out
    try:
        for _,e in ET.iterparse(path,events=("end",)):
            tag=e.tag.rsplit("}",1)[-1]
            if tag=="channel":
                cid=(e.get("id") or "").strip()
                if cid:
                    names=[(x.text or "").strip() for x in e.findall("./display-name") if (x.text or "").strip()]
                    old=out.get(cid)
                    if old: old.names.extend(x for x in names if x not in old.names)
                    else: out[cid]=Channel(source,str(path),cid,names)
                e.clear()
            elif tag=="programme":
                cid=(e.get("channel") or "").strip()
                if cid:
                    ch=out.setdefault(cid,Channel(source,str(path),cid,[cid]))
                    ch.programmes.append(Programme((e.get("start") or "").strip(),(e.get("stop") or "").strip(),child_text(e,"./title")))
                e.clear()
    except ET.ParseError as ex: print(f"ADVARSEL: XML-fejl i {path}: {ex}")
    return out

def split_ids(v): return [x.strip() for x in re.split(r"[,;\n]+",str(v or "")) if x.strip() and x.strip().upper()!="X"]
def priority_file(base, explicit):
    choices=[explicit,base.parent/"data"/"channel_priority_v2.xlsx",base.parent/"channel_priority_v2.xlsx",Path.cwd()/"channel_priority_v2.xlsx"]
    return next((x for x in choices if x and x.exists()),None)
def load_priority(path):
    if not path:return [],defaultdict(list)
    ws=load_workbook(path,read_only=True,data_only=True).worksheets[0]; rows=list(ws.iter_rows(values_only=True))
    hi=next((i for i,r in enumerate(rows) if r and str(r[0]).strip()=="Kanal"),None)
    if hi is None:return [],defaultdict(list)
    h=[str(x or "").strip() for x in rows[hi]]; data=[]; idx=defaultdict(list)
    for row in rows[hi+1:]:
        if not row or not row[0]:continue
        rec={h[i]:(row[i] if i<len(row) else None) for i in range(len(h))}; data.append(rec)
        for col in ("Kanal","EPGShare ID'er","OpenEPG ID'er","BSS XMLTV ID'er","Output/UHF tvg-id","Sammenlagte kanal-ID'er (gammel reference)"):
            for x in split_ids(rec.get(col)): idx[norm(x)].append(rec)
    return data,idx

class SourceIndex:
    def __init__(self,channels):
        self.channels=channels; self.by_norm=defaultdict(set); self.by_token=defaultdict(set); self.key_ids=defaultdict(set)
        for cid,ch in channels.items():
            for v in {cid,ch.name,*ch.names}:
                n=norm(v)
                if n:self.by_norm[n].add(cid);self.key_ids[n].add(cid)
                for t in tokens(v):self.by_token[t].add(cid)
        self.keys=list(self.key_ids)
    def shortlist(self,target,mapped,limit):
        score=Counter(); variants={target.channel_id,target.name,*target.names}; ns={norm(x) for x in variants if norm(x)}; ts=set().union(*(tokens(x) for x in variants))
        for mid in mapped:
            if mid in self.channels:score[mid]+=1200
            for cid in self.by_norm.get(norm(mid),()):score[cid]+=1000
        for n in ns:
            for cid in self.by_norm.get(n,()):score[cid]+=900
            for close in get_close_matches(n,self.keys,n=12,cutoff=.48):
                r=SequenceMatcher(None,n,close).ratio()
                for cid in self.key_ids[close]:score[cid]+=int(r*100)
        for t in ts:
            for cid in self.by_token.get(t,()):score[cid]+=min(90,15+len(t)*5)
        return [self.channels[cid] for cid,_ in score.most_common(limit)]

def name_score(a,b):
    best=0
    for x in {a.channel_id,a.name,*a.names}:
        nx=norm(x)
        for y in {b.channel_id,b.name,*b.names}:
            ny=norm(y)
            if nx and ny:
                r=1 if nx==ny else max(SequenceMatcher(None,nx,ny).ratio(),.88 if nx in ny or ny in nx else 0)
                best=max(best,r)
    return best
def overlap(a,b):
    x,y=a.titles,b.titles
    if not x or not y:return 0,0
    same=len(x&y);return same,same/max(1,min(len(x),len(y)))
def candidates(target,index,mapped,top,shortlist):
    mapped_norm={norm(x) for x in mapped}; out=[]
    for cand in index.shortlist(target,mapped,shortlist):
        ns=name_score(target,cand); common,ov=overlap(target,cand); ismap=cand.channel_id in mapped or norm(cand.channel_id) in mapped_norm; exact=norm(target.channel_id)==norm(cand.channel_id)
        score=min(100,max(0,ns*65+ov*25+(12 if ismap else 0)+(8 if exact else 0)-(8 if cand.count==0 else 0)))
        reason=", ".join(x for x,v in (("priority-map",ismap),("normaliseret ID",exact),("navnelighed",ns>=.8),("programoverlap",ov>=.5)) if v)
        if score>=42 or ismap:out.append((score,ns,common,ov,reason,cand))
    return sorted(out,key=lambda x:(-x[0],-x[5].count,x[5].channel_id.lower()))[:top]
def matched_priority(ch,idx):
    for k in (norm(ch.channel_id),norm(ch.name)):
        if idx.get(k):return idx[k][0]
def mapped_ids(rec,source):
    col={"EPGShare":"EPGShare ID'er","BSS":"BSS XMLTV ID'er","OpenEPG":"OpenEPG ID'er"}[source]
    return set(split_ids(rec.get(col))) if rec else set()

def sheet(wb,title,headers,rows):
    ws=wb.create_sheet(title);ws.append(headers)
    for r in rows:ws.append(r)
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions;fill=PatternFill("solid",fgColor="1F4E78")
    for c in ws[1]:c.font=Font(color="FFFFFF",bold=True);c.fill=fill;c.alignment=Alignment(wrap_text=True)
    for col in range(1,ws.max_column+1):
        vals=[len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,200)+1)];ws.column_dimensions[get_column_letter(col)].width=min(50,max(10,max(vals,default=8)+2))
    return ws

def main():
    ap=argparse.ArgumentParser(description="EPGoal TV-ID analyse v2 med hurtig indeksering")
    ap.add_argument("--base",type=Path,default=BASE);ap.add_argument("--priority",type=Path);ap.add_argument("--report-dir",type=Path);ap.add_argument("--top",type=int,default=5);ap.add_argument("--shortlist",type=int,default=35)
    a=ap.parse_args();report=a.report_dir or a.base.parent/"reports";report.mkdir(parents=True,exist_ok=True);warnings=[]
    specs={"EPGoal":a.base/"epgoal.xml","EPGShare":a.base/"epgshare_dk1.xml","BSS":a.base/"bss_raw"/"bss_epg.xml"}; sources={}
    for name,path in specs.items():
        print(f"Indlæser {name}: {path}");sources[name]=parse_xml(path,name);print(f"  {len(sources[name]):,} kanal-id'er")
        if not path.exists():warnings.append(f"Mangler: {path}")
    merged={};od=a.base/"openepg_raw"
    for path in sorted(od.glob("*.xml")):
        print(f"Indlæser OpenEPG: {path.name}")
        for cid,ch in parse_xml(path,"OpenEPG").items():
            if cid not in merged:merged[cid]=ch
            else:merged[cid].programmes.extend(ch.programmes);merged[cid].names.extend(x for x in ch.names if x not in merged[cid].names)
    sources["OpenEPG"]=merged;print(f"  {len(merged):,} unikke OpenEPG kanal-id'er")
    pf=priority_file(a.base,a.priority);prows,pidx=load_priority(pf);print(f"Prioritetsfil: {pf or 'IKKE FUNDET'}")
    print("Bygger kanalindekser...");indexes={}
    for src in ("EPGShare","BSS","OpenEPG"):
        indexes[src]=SourceIndex(sources[src]);print(f"  {src}: {len(indexes[src].keys):,} normaliserede nøgler")
    print("Analyserer 59/output-kanaler mod shortlistede kandidater...")
    headers=["Status","Bemærkning","Display name","EPGoal ID","EPGoal programmer","Første start","Sidste stop","Priority kanal","Output/UHF tvg-id","Bedste EPGShare ID","EPGShare programmer","EPGShare score","Bedste BSS ID","BSS programmer","BSS score","Bedste OpenEPG ID","OpenEPG programmer","OpenEPG score"]
    cand_h=["EPGoal navn","EPGoal ID","EPGoal programmer","Kilde","Rang","Kandidatnavn","Kandidat-ID","Kandidat programmer","Første start","Sidste stop","Samlet score","Navnescore","Fælles titler","Titeloverlap pct","Matchgrund","Kildefil"]
    rows=[];crows=[];status_count=Counter()
    total=len(sources["EPGoal"])
    for no,ch in enumerate(sorted(sources["EPGoal"].values(),key=lambda x:x.name.lower()),1):
        print(f"  [{no:02d}/{total:02d}] {ch.name}")
        rec=matched_priority(ch,pidx);best={}; all_best=[]
        for src in ("EPGShare","BSS","OpenEPG"):
            found=candidates(ch,indexes[src],mapped_ids(rec,src),a.top,a.shortlist);best[src]=found[0] if found else None;all_best += found
            for rank,item in enumerate(found,1):
                sc,ns,common,ov,reason,c=item;crows.append([ch.name,ch.channel_id,ch.count,src,rank,c.name,c.channel_id,c.count,c.first,c.last,round(sc,1),round(ns*100,1),common,round(ov*100,1),reason,Path(c.source_file).name])
        if ch.count==0:status,note="MANGLER EPG",("Kandidat fundet" if max((x[0] for x in all_best),default=0)>=65 else "Ingen sikker kandidat")
        elif any(x[0]>=78 and x[3]<.15 for x in all_best):status,note="KONTROLLÉR INDHOLD","Stærkt navnematch, men lavt programoverlap"
        elif not any(x[0]>=70 for x in all_best):status,note="ID-KONFLIKT","Ingen tydelig rå kilde matcher"
        else:status,note="OK",""
        status_count[status]+=1
        def cols(src):
            x=best[src];return [x[5].channel_id,x[5].count,round(x[0],1)] if x else ["",0,0]
        rows.append([status,note,ch.name,ch.channel_id,ch.count,ch.first,ch.last,rec.get("Kanal","") if rec else "",rec.get("Output/UHF tvg-id","") if rec else "",*cols("EPGShare"),*cols("BSS"),*cols("OpenEPG")])
    wb=Workbook();wb.remove(wb.active)
    summary=[["Kørt",datetime.now().strftime("%Y-%m-%d %H:%M:%S")],["Basismappe",str(a.base)],["Prioritetsfil",str(pf or "IKKE FUNDET")],*[ [f"{k} kanal-id'er",len(v)] for k,v in sources.items()],*[ [f"Status: {k}",v] for k,v in sorted(status_count.items())]]
    sheet(wb,"Summary",["Måling","Værdi"],summary);sheet(wb,"Channel Sources",headers,rows);sheet(wb,"Missing EPG",headers,[r for r in rows if r[0]=="MANGLER EPG"]);sheet(wb,"Potential Matches",cand_h,crows);sheet(wb,"Wrong Mappings",headers,[r for r in rows if r[0] in ("KONTROLLÉR INDHOLD","ID-KONFLIKT")])
    vr=[]
    for r in prows:
        names=split_ids(r.get("BSS M3U stream-navne"));ids=split_ids(r.get("BSS M3U tvg-id'er"));q=sorted({x.upper() for n in names for x in QUALITY if re.search(rf"\b{x}\b",n,re.I)})
        vr.append([r.get("Kanal",""),r.get("Følg (X)",""),r.get("Output/UHF tvg-id",""),", ".join(ids),", ".join(names),", ".join(q),len(ids),len(names),"KONTROLLÉR" if len(ids)>1 or(names and not ids) else "OK"])
    sheet(wb,"HD-FHD Variants",["Kanal","Følg","Output ID","M3U tvg-id'er","Streamnavne","Varianter","ID-antal","Stream-antal","Status"],vr)
    xp=report/"tv_id_analysis_v2.xlsx";cp=report/"tv_id_candidates_v2.csv";jp=report/"tv_id_analysis_v2.json";wb.save(xp)
    with cp.open("w",encoding="utf-8-sig",newline="") as f:w=csv.writer(f,delimiter=";");w.writerow(cand_h);w.writerows(crows)
    jp.write_text(json.dumps({"generated":datetime.now().isoformat(timespec="seconds"),"statuses":dict(status_count),"channels":[dict(zip(headers,r)) for r in rows]},ensure_ascii=False,indent=2),encoding="utf-8")
    print("\n=== TV-ID ANALYSE V2 FÆRDIG ===");print(f"Excel : {xp}\nCSV   : {cp}\nJSON  : {jp}")
if __name__=="__main__":main()
