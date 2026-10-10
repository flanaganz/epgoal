#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Installer EPGoal Console MVP i C:\EPGoal\EPGoalConsole."""
from pathlib import Path
import json

ROOT = Path(r"C:\EPGoal\EPGoalConsole")
FILES = {}

FILES['requirements.txt'] = '''nicegui>=2.0,<4
openpyxl>=3.1
'''

FILES['start_console.bat'] = r'''@echo off
cd /d C:\EPGoal\EPGoalConsole
if not exist .venv\Scripts\python.exe (
  py -m venv .venv
  .venv\Scripts\python.exe -m pip install --upgrade pip
  .venv\Scripts\python.exe -m pip install -r requirements.txt
)
start "EPGoal Console" http://127.0.0.1:8088
.venv\Scripts\python.exe app.py
pause
'''

FILES['config/gui_settings.json'] = json.dumps({
    "epgoal_root": r"C:\EPGoal",
    "host": "127.0.0.1",
    "port": 8088,
    "history_keep_runs": 1000
}, ensure_ascii=False, indent=2)

FILES['config/source_overrides.json'] = json.dumps({
    "Discovery Science.dk": {
        "mode": "forced",
        "source": "openepg",
        "source_channel_id": "Discovery Science.dk",
        "output_id": "discoveryscience.dk",
        "enabled": True
    }
}, ensure_ascii=False, indent=2)

FILES['config/title_rules.json'] = json.dumps({
    "rules": [
        {
            "name": "Episodeangivelse efter kolon",
            "pattern": r"^(.+?):\s*\(\d+\s*:\s*\d+\)",
            "replacement": r"\1",
            "enabled": True,
            "preserve_original_title": True
        }
    ]
}, ensure_ascii=False, indent=2)

FILES['core/__init__.py'] = ''
FILES['core/paths.py'] = r'''from pathlib import Path
import json
CONSOLE_ROOT = Path(__file__).resolve().parent.parent
SETTINGS_FILE = CONSOLE_ROOT / 'config' / 'gui_settings.json'
def settings():
    return json.loads(SETTINGS_FILE.read_text(encoding='utf-8'))
def paths():
    root = Path(settings()['epgoal_root'])
    return {
        'root': root,
        'scripts': root / 'scripts',
        'data': root / 'data',
        'output': root / 'output_epgshare',
        'priority': root / 'data' / 'channel_priority_v2.xlsx',
        'merge_log': root / 'data' / 'epgshare_merge_log.json',
        'normalize_log': root / 'data' / 'normalize_uhf_channel_ids_log.json',
        'danish_log': root / 'data' / 'danish_backdrops_run_log.json',
        'epgoal': root / 'output_epgshare' / 'epgoal.xml',
        'pipeline': root / 'scripts' / 'opdater_epgoal.py',
        'save_choices': root / 'scripts' / 'gem_epgoal_valg.py',
    }
'''

FILES['core/data.py'] = r'''import json, re
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime
from openpyxl import load_workbook
from .paths import paths, CONSOLE_ROOT

def load_json(path, default):
    try: return json.loads(path.read_text(encoding='utf-8-sig'))
    except Exception: return default

def split_ids(v):
    return [x.strip() for x in re.split(r'[,;\n]+', str(v or '')) if x.strip()]

def load_channels():
    p=paths()['priority']
    if not p.exists(): return []
    ws=load_workbook(p, read_only=True, data_only=True).worksheets[0]
    rows=list(ws.iter_rows(values_only=True))
    hi=next((i for i,r in enumerate(rows) if r and str(r[0] or '').strip()=='Kanal'), None)
    if hi is None: return []
    h=[str(x or '').strip() for x in rows[hi]]; out=[]
    overrides=load_json(CONSOLE_ROOT/'config'/'source_overrides.json', {})
    for row in rows[hi+1:]:
        if not row or not row[0]: continue
        rec={h[i]:(row[i] if i<len(row) else None) for i in range(len(h))}
        name=str(rec.get('Kanal') or '')
        ov=overrides.get(name,{})
        out.append({
          'channel':name,'follow':str(rec.get('Følg (X)') or ''),
          'streams':str(rec.get('BSS M3U stream-navne') or ''),
          'epgshare':str(rec.get("EPGShare ID'er") or ''),
          'openepg':str(rec.get("OpenEPG ID'er") or ''),
          'bss':str(rec.get("BSS XMLTV ID'er") or ''),
          'output':str(rec.get('Output/UHF tvg-id') or ''),
          'source_mode':ov.get('source','auto') if ov.get('enabled',True) else 'auto',
          'source_id':ov.get('source_channel_id',''),
        })
    return out

def epgoal_stats():
    p=paths()['epgoal']; counts=Counter(); names={}
    if not p.exists(): return {'channels':0,'programmes':0,'empty':0,'counts':{},'names':{}}
    try:
      for _,e in ET.iterparse(p, events=('end',)):
        tag=e.tag.rsplit('}',1)[-1]
        if tag=='channel':
          cid=e.get('id',''); names[cid]=next(((x.text or '').strip() for x in e.findall('./display-name') if (x.text or '').strip()),cid); e.clear()
        elif tag=='programme': counts[e.get('channel','')]+=1; e.clear()
    except ET.ParseError: pass
    return {'channels':len(names),'programmes':sum(counts.values()),'empty':sum(1 for x in names if counts[x]==0),'counts':dict(counts),'names':names}

def merge_stats():
    d=load_json(paths()['merge_log'], {})
    kept=d.get('kept_as_is',[]); supplied=d.get('supplemented',[]); missing=d.get('no_fallback_available',[])
    sources=Counter(['epgshare']*len(kept))
    for x in supplied:sources[x.get('kilde_brugt','ukendt')]+=1
    return {'kept':len(kept),'supplied':len(supplied),'missing':len(missing),'sources':dict(sources),'raw':d}

def current_snapshot():
    e=epgoal_stats();m=merge_stats()
    return {'created_at':datetime.now().isoformat(timespec='seconds'),'channels':e['channels'],'programmes':e['programmes'],'empty_channels':e['empty'],'kept_epgshare':m['kept'],'fallback_channels':m['supplied'],'missing_channels':m['missing'],'source_counts':m['sources']}
'''

FILES['core/history.py'] = r'''import sqlite3, json
from .paths import CONSOLE_ROOT
DB=CONSOLE_ROOT/'database'/'epgoal_history.sqlite'
def init_db():
 DB.parent.mkdir(parents=True,exist_ok=True)
 with sqlite3.connect(DB) as c:
  c.execute("CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, created_at TEXT, status TEXT, channels INTEGER, programmes INTEGER, empty_channels INTEGER, kept_epgshare INTEGER, fallback_channels INTEGER, missing_channels INTEGER, source_counts TEXT, log_file TEXT)")
def add_run(snapshot,status='success',log_file=''):
 init_db()
 with sqlite3.connect(DB) as c:
  c.execute('INSERT INTO runs(created_at,status,channels,programmes,empty_channels,kept_epgshare,fallback_channels,missing_channels,source_counts,log_file) VALUES(?,?,?,?,?,?,?,?,?,?)',(snapshot['created_at'],status,snapshot['channels'],snapshot['programmes'],snapshot['empty_channels'],snapshot['kept_epgshare'],snapshot['fallback_channels'],snapshot['missing_channels'],json.dumps(snapshot['source_counts']),str(log_file)))
def runs(limit=100):
 init_db()
 with sqlite3.connect(DB) as c:
  c.row_factory=sqlite3.Row
  return [dict(x) for x in c.execute('SELECT * FROM runs ORDER BY id DESC LIMIT ?',(limit,))]
'''

FILES['core/rules.py'] = r'''import json,re
from .paths import CONSOLE_ROOT
FILE=CONSOLE_ROOT/'config'/'title_rules.json'
def load_rules(): return json.loads(FILE.read_text(encoding='utf-8')).get('rules',[])
def derive_title(title):
 for rule in load_rules():
  if not rule.get('enabled',True):continue
  if re.search(rule['pattern'],title,re.I):return re.sub(rule['pattern'],rule['replacement'],title,flags=re.I).strip(),rule['name']
 return title,'Ingen regel'
'''

FILES['core/runner.py'] = r'''import subprocess,sys,threading
from datetime import datetime
from .paths import paths,CONSOLE_ROOT
from .data import current_snapshot
from .history import add_run

def run_script(which,on_line,on_done):
 target=paths()[which]
 log_dir=CONSOLE_ROOT/'logs';log_dir.mkdir(exist_ok=True)
 log=log_dir/f"{which}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
 def work():
  rc=1
  try:
   with log.open('w',encoding='utf-8') as f:
    p=subprocess.Popen([sys.executable,str(target)],cwd=paths()['root'],stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,encoding='utf-8',errors='replace',bufsize=1)
    for line in p.stdout:
     f.write(line);f.flush();on_line(line.rstrip())
    rc=p.wait()
  except Exception as e:on_line('FEJL: '+str(e))
  snap=current_snapshot();add_run(snap,'success' if rc==0 else 'failed',log);on_done(rc,log)
 threading.Thread(target=work,daemon=True).start()
'''

FILES['app.py'] = r'''from nicegui import ui
import json
from pathlib import Path
from core.paths import paths,CONSOLE_ROOT,settings
from core.data import load_channels,epgoal_stats,merge_stats,current_snapshot
from core.history import init_db,runs,add_run
from core.rules import derive_title
from core.runner import run_script

init_db(); dark=ui.dark_mode(); dark.enable()
COLORS={'ok':'positive','warn':'warning','bad':'negative'}

def metric(title,value,sub=''):
 with ui.card().classes('w-56'):
  ui.label(title).classes('text-sm text-grey-5');ui.label(str(value)).classes('text-3xl font-bold');ui.label(sub).classes('text-xs text-grey-5')

def dashboard():
 ui.label('EPGoal Console').classes('text-3xl font-bold')
 e=epgoal_stats();m=merge_stats()
 with ui.row().classes('gap-4 wrap'):
  metric('Kanaler',e['channels']);metric('Programmer',f"{e['programmes']:,}");metric('Kanaler uden EPG',e['empty']);metric('Fallback-kanaler',m['supplied']);metric('Uden fallback',m['missing'])
 ui.separator();ui.label('Kildefordeling').classes('text-xl')
 data=[{'name':k,'value':v} for k,v in m['sources'].items()]
 ui.echart({'tooltip':{'trigger':'item'},'series':[{'type':'pie','radius':['45%','72%'],'data':data}]}).classes('w-full h-80')
 hist=list(reversed(runs(60)))
 if hist:
  ui.label('Historik').classes('text-xl')
  ui.echart({'tooltip':{'trigger':'axis'},'legend':{'data':['Programmer','Problemkanaler']},'xAxis':{'type':'category','data':[x['created_at'][5:16] for x in hist]},'yAxis':{'type':'value'},'series':[{'name':'Programmer','type':'line','data':[x['programmes'] for x in hist]},{'name':'Problemkanaler','type':'line','data':[x['empty_channels']+x['missing_channels'] for x in hist]}]}).classes('w-full h-96')

def channels_page():
 ui.label('Kanaler').classes('text-3xl font-bold');chs=load_channels();e=epgoal_stats();overfile=CONSOLE_ROOT/'config'/'source_overrides.json';over=json.loads(overfile.read_text(encoding='utf-8'))
 rows=[]
 for i,c in enumerate(chs):
  count=e['counts'].get(c['output'],0);rows.append({'id':i,**c,'programmes':count,'status':'OK' if count else 'INGEN EPG'})
 cols=[{'name':'channel','label':'Kanal','field':'channel','sortable':True},{'name':'source_mode','label':'Kilde','field':'source_mode','sortable':True},{'name':'source_id','label':'Kilde-ID','field':'source_id'},{'name':'output','label':'Output-ID','field':'output'},{'name':'programmes','label':'Programmer','field':'programmes','sortable':True},{'name':'status','label':'Status','field':'status'}]
 table=ui.table(columns=cols,rows=rows,row_key='id',pagination=20).classes('w-full').props('dense')
 with table.add_slot('body-cell-source_mode'):
  ui.html("""<q-td :props=\"props\"><q-select dense borderless emit-value map-options :options=\"['auto','epgshare','openepg','bss','exclude']\" v-model=\"props.row.source_mode\" @update:model-value=\"() => $parent.$emit('source_change', props.row)\" /></q-td>""")
 def changed(msg):
  r=msg.args;name=r['channel'];mode=r['source_mode']
  if mode=='auto':over.pop(name,None)
  else:over[name]={'mode':'forced','source':mode,'source_channel_id':r.get('source_id',''),'output_id':r.get('output',''),'enabled':True}
  overfile.write_text(json.dumps(over,ensure_ascii=False,indent=2),encoding='utf-8');ui.notify(f'Gemt: {name} → {mode}',type='positive')
 table.on('source_change',changed)

def run_page():
 ui.label('Kørsel').classes('text-3xl font-bold');state={'running':False};log=ui.log(max_lines=2500).classes('w-full h-[620px] bg-black text-green-4')
 def start(which):
  if state['running']:ui.notify('En kørsel er allerede i gang',type='warning');return
  state['running']=True;log.clear();log.push('Starter '+which+' ...')
  def line(x):ui.timer(0,lambda:log.push(x),once=True)
  def done(rc,path):
   def finish():state['running']=False;log.push(f'FÆRDIG, exit={rc}, log={path}');ui.notify('Kørsel færdig' if rc==0 else 'Kørsel fejlede',type='positive' if rc==0 else 'negative')
   ui.timer(0,finish,once=True)
  run_script(which,line,done)
 with ui.row():
  ui.button('Byg frisk EPG',on_click=lambda:start('pipeline'),icon='play_arrow').props('color=primary')
  ui.button('Gem artwork-valg',on_click=lambda:start('save_choices'),icon='save').props('color=secondary')
  ui.button('Registrér snapshot',on_click=lambda:(add_run(current_snapshot(),'manual',''),ui.notify('Snapshot gemt',type='positive')),icon='camera')

def rules_page():
 ui.label('Titelregler').classes('text-3xl font-bold');title=ui.input('Testtitel',value='Folk og Fæ: (6:7) Frank fister Margrethe med piskeris.').classes('w-full');result=ui.label().classes('text-xl')
 def test():
  cleaned,rule=derive_title(title.value);result.set_text(f'TMDb-søgetitel: {cleaned} | Regel: {rule}')
 ui.button('Test titelrensning',on_click=test);test()
 ui.markdown('Originaltitlen i XML ændres ikke. Reglerne bruges kun til TMDb-opslag.')

def history_page():
 ui.label('Kørselshistorik').classes('text-3xl font-bold');r=runs(500)
 cols=[{'name':x,'label':x.replace('_',' ').title(),'field':x,'sortable':True} for x in ['created_at','status','channels','programmes','empty_channels','fallback_channels','missing_channels','log_file']]
 ui.table(columns=cols,rows=r,row_key='id',pagination=25).classes('w-full').props('dense')

with ui.header().classes('items-center justify-between'):
 ui.label('EPGoal Console').classes('text-xl font-bold');ui.label(str(paths()['root'])).classes('text-xs')
with ui.left_drawer(value=True).classes('bg-grey-10'):
 ui.button('Dashboard',on_click=lambda:ui.navigate.to('/'),icon='dashboard').props('flat').classes('w-full')
 ui.button('Kanaler',on_click=lambda:ui.navigate.to('/channels'),icon='tv').props('flat').classes('w-full')
 ui.button('Kørsel',on_click=lambda:ui.navigate.to('/run'),icon='play_circle').props('flat').classes('w-full')
 ui.button('Titelregler',on_click=lambda:ui.navigate.to('/rules'),icon='auto_fix_high').props('flat').classes('w-full')
 ui.button('Historik',on_click=lambda:ui.navigate.to('/history'),icon='insights').props('flat').classes('w-full')
@ui.page('/')
def p1():dashboard()
@ui.page('/channels')
def p2():channels_page()
@ui.page('/run')
def p3():run_page()
@ui.page('/rules')
def p4():rules_page()
@ui.page('/history')
def p5():history_page()

s=settings();ui.run(title='EPGoal Console',host=s['host'],port=s['port'],reload=False,show=False,favicon='📺')
'''

FILES['README.txt'] = '''EPGoal Console MVP

Start:
  Dobbeltklik start_console.bat

Første start opretter et lokalt Python-miljø og installerer NiceGUI.
Åbn derefter http://127.0.0.1:8088

Eksisterende EPGoal-filer ændres ikke af installationen.
Kildevalg gemmes i config/source_overrides.json.
Historik gemmes i database/epgoal_history.sqlite.
'''

def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    for rel, content in FILES.items():
        path=ROOT/rel;path.parent.mkdir(parents=True,exist_ok=True)
        if path.exists() and rel.startswith('config/'):
            print(f'BEVARET: {path}')
            continue
        path.write_text(content,encoding='utf-8')
        print(f'SKREV: {path}')
    print('\nEPGoal Console er installeret.')
    print(r'Start med: C:\EPGoal\EPGoalConsole\start_console.bat')
if __name__=='__main__':main()
